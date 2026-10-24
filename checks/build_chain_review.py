"""Builds checks/chain_review.xlsx: every part in dependency order (inputs first),
with its sources, so J.B. and Claude can review the chain node by node.

Sheets:
  Order    one row per part, in build order (level 0 = depends on nothing)
  Sources  one row per part × source, the thing a source check confirms
  Key      what the columns and statuses mean
Review columns (left blank here) are filled during the review; this script
keeps any review entries already in the file, so it can be re-run safely.
"""
import json, os, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
d = json.load(open(os.path.join(ROOT, 'register', 'diagram_data.json')))
reg = {q['id']: q for q in json.load(open(os.path.join(ROOT, 'register', 'formula_register.json')))['quantities']}
SRC = d['sources']
N = {x['id']: x for x in d['nodes']}
SN = {s['id']: s['name'] for s in d['stages']}
ORDER = ['input', 'size', 'place', 'characterise', 'prospect', 'opportunity', 'test', 'decide']

dep = {i: [j for j in x['inputs'] if j in N] for i, x in N.items()}
lvl = {}
def level(i):
    if i not in lvl: lvl[i] = 0 if not dep[i] else 1 + max(level(j) for j in dep[i])
    return lvl[i]
for i in N: level(i)
used_by = {i: [] for i in N}
for i, ds in dep.items():
    for j in ds: used_by[j].append(i)
seq = sorted(N, key=lambda i: (lvl[i], ORDER.index(N[i]['stage']), -len(used_by[i]), i))

OUT = os.path.join(HERE, 'chain_review.xlsx')
REVIEW_O = ['Source check', 'Math check', 'Display check', 'J.B. decision', 'Status', 'Notes']
REVIEW_S = ['Agent verdict', 'Page / table / quote', 'Opened (date)', 'J.B. OK', 'Notes']
keep_o, keep_s = {}, {}
if os.path.exists(OUT):  # keep review work already entered
    old = load_workbook(OUT)
    o = old['Order']; hdr = [c.value for c in o[1]]
    for r in o.iter_rows(min_row=2, values_only=True):
        keep_o[r[hdr.index('id')]] = {h: r[hdr.index(h)] for h in REVIEW_O if h in hdr}
    s = old['Sources']; hdr = [c.value for c in s[1]]
    for r in s.iter_rows(min_row=2, values_only=True):
        keep_s[(r[hdr.index('Part id')], r[hdr.index('Source id')])] = {h: r[hdr.index(h)] for h in REVIEW_S if h in hdr}

wb = Workbook(); o = wb.active; o.title = 'Order'
H = ['#', 'Level', 'id', 'Part', 'Stage', 'Kind', 'Depends on', 'Used by (count)', 'Formula', 'Coefficients', 'Unit',
     'Sources (count)', 'Source text (inputs)', 'Trust', 'Chain trust'] + REVIEW_O
o.append(H)
for n, i in enumerate(seq, 1):
    x = N[i]
    srcs = x.get('sources') or []
    row = [n, lvl[i], i, x['name'], SN[x['stage']], x['kind'], ', '.join(N[j]['name'] for j in dep[i]), len(used_by[i]),
           x.get('formula', ''), x.get('coef', ''), x.get('unit', ''), len(srcs), x.get('source_text', '') if not srcs else '',
           x.get('conf'), x.get('upstream')]
    k = keep_o.get(i, {}); row += [k.get(h) or ('To do' if h == 'Status' else None) for h in REVIEW_O]
    o.append(row)

s = wb.create_sheet('Sources')
HS = ['Part #', 'Part id', 'Part', 'Source id', 'Used for', 'Role', 'Register trust', 'Register says verified', 'Source (as listed)'] + REVIEW_S
s.append(HS)
for n, i in enumerate(seq, 1):
    x = N[i]; roles = {r['id']: r for r in (reg.get(i, {}).get('source_roles') or [])}
    ids = list(x.get('sources') or [])
    if not ids and x.get('source_text'):
        ids = ['(text only)']
    for sid in ids:
        rr = roles.get(sid, {})
        txt = SRC.get(sid, x.get('source_text', '') if sid == '(text only)' else '(missing from source list)')
        row = [n, i, x['name'], sid, rr.get('term', ''), rr.get('role', ''), rr.get('conf', ''),
               {True: 'yes', False: 'no'}.get(rr.get('verified'), ''), txt]
        k = keep_s.get((i, sid), {}); row += [k.get(h) for h in REVIEW_S]
        s.append(row)

k = wb.create_sheet('Key')
for line in [
    ['How to read this workbook'],
    ['Order', 'Every part in build order. Level 0 depends on nothing; level 1 depends only on level 0; and so on. Review top to bottom.'],
    ['Sources', 'One row per part and source. A source check opens the source and confirms the figure, page and date.'],
    ['Status', 'To do · In review · OK · Fix needed · Decision needed · Unverified (kept, marked) · Retired'],
    ['Agent verdict', 'Confirmed · Corrected (figure differs) · Unverified (could not open) · Wrong source · Popular source (replace)'],
    ['Trust', '100 arithmetic · 90 rule or standard · 80 government data · 70 peer-reviewed, transferred · 60 derived · 50 claim · 40 proxy · 20 judged · 0 gap'],
    ['Chain trust', 'The lowest trust among the part and everything it depends on.'],
]: k.append(line)

for ws, widths in ((o, {'A': 5, 'B': 6, 'C': 18, 'D': 34, 'E': 18, 'F': 9, 'G': 34, 'H': 8, 'I': 50, 'J': 60, 'K': 8, 'L': 8, 'M': 40, 'N': 7, 'O': 7, 'P': 16, 'Q': 16, 'R': 16, 'S': 24, 'T': 16, 'U': 40}),
                   (s, {'A': 6, 'B': 18, 'C': 30, 'D': 16, 'E': 40, 'F': 12, 'G': 8, 'H': 9, 'I': 70, 'J': 16, 'K': 40, 'L': 12, 'M': 10, 'N': 40}),
                   (k, {'A': 16, 'B': 120})):
    for c, w in widths.items(): ws.column_dimensions[c].width = w
    for c in ws[1]: c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for r in ws.iter_rows(min_row=2):
        for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')

dv = DataValidation(type='list', formula1='"To do,In review,OK,Fix needed,Decision needed,Unverified,Retired"', allow_blank=True)
o.add_data_validation(dv); dv.add(f'T2:T{o.max_row}')
dv2 = DataValidation(type='list', formula1='"Confirmed,Corrected,Unverified,Wrong source,Popular source"', allow_blank=True)
s.add_data_validation(dv2); dv2.add(f'J2:J{s.max_row}')
o.auto_filter.ref = o.dimensions; s.auto_filter.ref = s.dimensions
wb.save(OUT)
print('parts', len(seq), 'source rows', s.max_row - 1, 'levels', max(lvl.values()) + 1)
