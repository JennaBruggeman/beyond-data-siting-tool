"""Builds checks/size_check.xlsx: an independent hand calculation of the Size stage
(How big it can be) with live Excel formulas, for the five golden lots.

The formulas are written from the register text and the engine's documented rules,
not copied from engine values. Engine values are pasted below them only to compare.
Run:  python3 build_size_check.py   then   python3 compare_size_check.py
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
LOTS = ['V09', 'E05', 'CLARK', 'E123', 'HAST']
G = {k: json.load(open(os.path.join(HERE, 'golden', k + '.json'))) for k in LOTS}

wb = Workbook()
c = wb.active; c.title = 'Constants'
CONST = [  # name, value, source (register / engine R), trust
    ('ftf_ground', 6.0, 'register hall_area / engine R.ftf.ground: ground floor 6.0 m (I-2/I-3 minimum; Burlington)', 90),
    ('ftf_up_lo', 5.5, 'R.ftf.up[0]: data-hall floor 5.5 m (HK TPB A/KC/510B) -> more storeys', 50),
    ('ftf_up_hi', 6.0, 'R.ftf.up[1]: data-hall floor 6.0 m (Burlington, HK TPB A/KC/480) -> fewer storeys', 50),
    ('hall_share_lo', 0.48, 'register hall_area: Burlington 48 % of gross', 50),
    ('hall_share_hi', 0.72, 'register hall_area: Burlington 72 % of gross', 50),
    ('m2_per_mw_lo', 290, 'register it_load_outside: densest built Canadian hall (eStruxture VAN-1, claim)', 50),
    ('m2_per_mw_hi', 630, 'register it_load_outside: VAN5 ≈627 (claim)', 50),
    ('gross_per_mw_lo', 505, 'engine R.grossPerMW: min of 11 benchmark buildings', 50),
    ('gross_per_mw_hi', 1109, 'engine R.grossPerMW: max of 11 benchmark buildings', 50),
    ('mw_cap', 40, 'engine toMW(): Math.min(40, …) — NOT in the register (finding S-03)', 0),
    ('w_per_server', 600, 'register servers: 600 W per 2-socket server (LBNL 2024)', 80),
    ('kw_per_rack', 8, 'engine R.kWPerRack: Uptime 2024 average (survey)', 50),
]
c.append(['name', 'value', 'source', 'trust'])
for row in CONST: c.append(list(row))
for i, (n, *_r) in enumerate(CONST, start=2):
    wb.defined_names[n] = __import__('openpyxl').workbook.defined_name.DefinedName(n, attr_text=f"Constants!$B${i}")
c.column_dimensions['A'].width = 18; c.column_dimensions['C'].width = 90

s = wb.create_sheet('Size')
s['A1'] = 'Size stage — hand calculation (live formulas) vs engine'; s['A1'].font = Font(bold=True, size=13)
s['A2'] = 'Inputs in blue are read from the site (lot, zone rules, footprint from the site setup). Everything else is a formula.'
for j, k in enumerate(LOTS):
    s.cell(3, 2 + j, k).font = Font(bold=True)

rows = {}
def row(key, label, vals=None, formula=None, inp=False):
    r = s.max_row + 1; rows[key] = r; s.cell(r, 1, label)
    for j, k in enumerate(LOTS):
        col = L(2 + j)
        if inp:
            cell = s.cell(r, 2 + j, vals[j]); cell.font = Font(color='1F4E9E')
        elif formula:
            s.cell(r, 2 + j, '=' + formula(col))
    return r
R = lambda key, col: f'{col}{rows[key]}'

def facts(k): return G[k]['facts']
def ans(k): return G[k]['answers']
env = lambda k: facts(k).get('envU') or {}

s.append(['INPUTS'])
row('lot', 'Lot area (m²)', [facts(k)['lot'] for k in LOTS], inp=True)
row('fsr_lo', 'Floor-space ratio, low', [ (env(k).get('fsr') or [0, 0])[0] for k in LOTS], inp=True)
row('fsr_hi', 'Floor-space ratio, high', [ (env(k).get('fsr') or [0, 0])[1] for k in LOTS], inp=True)
row('h', 'Height limit (m)', [facts(k).get('envH') or 0 for k in LOTS], inp=True)
row('fp', 'Footprint from site setup (m²)', [facts(k).get('fp') or 0 for k in LOTS], inp=True)
row('ind', 'Indoor plant floor (m²)', [facts(k).get('indoorPlant') or 0 for k in LOTS], inp=True)
row('ground', 'No data storage at grade (1 = yes)', [1 if env(k).get('ground') else 0 for k in LOTS], inp=True)
row('new', 'New building (1) or existing (0)', [1 if ans(k)['built'] == 'new' else 0 for k in LOTS], inp=True)
row('whole', 'Whole building (1) or part (0)', [1 if ans(k)['extent'] == 'whole' else 0 for k in LOTS], inp=True)
row('plate', 'Largest existing footprint (m²)', [facts(k).get('plate') or 0 for k in LOTS], inp=True)

s.append(['CALCULATION'])
row('nf_lo', 'Storeys, fewer (6.0 m floors)', formula=lambda c: f'1+MAX(0,ROUNDDOWN(({R("h",c)}-ftf_ground)/ftf_up_hi,0))')
row('nf_hi', 'Storeys, more (5.5 m floors)', formula=lambda c: f'1+MAX(0,ROUNDDOWN(({R("h",c)}-ftf_ground)/ftf_up_lo,0))')
row('g_lo', 'Zoning floor area, low = lot × FSR', formula=lambda c: f'{R("lot",c)}*{R("fsr_lo",c)}')
row('g_hi', 'Zoning floor area, high', formula=lambda c: f'{R("lot",c)}*{R("fsr_hi",c)}')
row('G_lo', 'Buildable, fewer storeys = min(zoning, footprint × storeys) − indoor plant', formula=lambda c: f'MAX(0,MIN({R("g_lo",c)},{R("fp",c)}*{R("nf_lo",c)})-{R("ind",c)})')
row('G_hi', 'Buildable, more storeys', formula=lambda c: f'MAX(0,MIN({R("g_hi",c)},{R("fp",c)}*{R("nf_hi",c)})-{R("ind",c)})')
row('gross_lo', 'Gross floor area, low', formula=lambda c: f'IF({R("new",c)}=1,IF({R("whole",c)}=1,MIN({R("G_lo",c)},{R("G_hi",c)}),MIN({R("G_lo",c)}/{R("nf_lo",c)},{R("G_hi",c)}/{R("nf_hi",c)})),IF({R("whole",c)}=1,"see existing",{R("plate",c)}))')
row('gross_hi', 'Gross floor area, high', formula=lambda c: f'IF({R("new",c)}=1,IF({R("whole",c)}=1,MAX({R("G_lo",c)},{R("G_hi",c)}),MAX({R("G_lo",c)}/{R("nf_lo",c)},{R("G_hi",c)}/{R("nf_hi",c)})),IF({R("whole",c)}=1,"see existing",{R("plate",c)}))')
gr = lambda c: f'AND({R("new",c)}=1,{R("whole",c)}=1,{R("ground",c)}=1)'
row('up_lo', 'Upper floors, fewer-storey case = G − G ÷ storeys', formula=lambda c: f'IF({gr(c)},{R("G_lo",c)}-{R("G_lo",c)}/{R("nf_lo",c)},"")')
row('up_hi', 'Upper floors, more-storey case', formula=lambda c: f'IF({gr(c)},{R("G_hi",c)}-{R("G_hi",c)}/{R("nf_hi",c)},"")')
row('hg_lo', 'Floor the hall may use, low', formula=lambda c: f'IF({gr(c)},MIN({R("up_lo",c)},{R("up_hi",c)}),{R("gross_lo",c)})')
row('hg_hi', 'Floor the hall may use, high', formula=lambda c: f'IF({gr(c)},MAX({R("up_lo",c)},{R("up_hi",c)}),{R("gross_hi",c)})')
row('hs_lo', 'Hall, strict reading, low = × 0.48', formula=lambda c: f'IF({gr(c)},MIN({R("up_lo",c)}*hall_share_lo,{R("up_hi",c)}*hall_share_lo),{R("hg_lo",c)}*hall_share_lo)')
row('hs_hi', 'Hall, strict reading, high = × 0.72', formula=lambda c: f'IF({gr(c)},MAX({R("up_lo",c)}*hall_share_hi,{R("up_hi",c)}*hall_share_hi),{R("hg_hi",c)}*hall_share_hi)')
row('hl_lo', 'Hall, lenient reading, low = min(upper, G × 0.48)', formula=lambda c: f'IF({gr(c)},MIN(MIN({R("up_lo",c)},{R("G_lo",c)}*hall_share_lo),MIN({R("up_hi",c)},{R("G_hi",c)}*hall_share_lo)),"")')
row('hl_hi', 'Hall, lenient reading, high = min(upper, G × 0.72)', formula=lambda c: f'IF({gr(c)},MAX(MIN({R("up_lo",c)},{R("G_lo",c)}*hall_share_hi),MIN({R("up_hi",c)},{R("G_hi",c)}*hall_share_hi)),"")')
row('mw_lo', 'Largest that fits, low (MW) = hall ÷ 630', formula=lambda c: f'MIN(mw_cap,{R("hs_lo",c)}/m2_per_mw_hi)')
row('mw_hi', 'Largest that fits, high (MW) = hall ÷ 290', formula=lambda c: f'MAX(MIN(mw_cap,{R("hs_hi",c)}/m2_per_mw_lo),IF({gr(c)},MIN(mw_cap,{R("hl_hi",c)}/m2_per_mw_lo),0))')
row('mw_hi_nocap', 'High, without the 40 MW cap', formula=lambda c: f'MAX({R("hs_hi",c)},IF({gr(c)},{R("hl_hi",c)},0))/m2_per_mw_lo')
row('plantfor', 'Plant-fit size from the engine search (input; checked in component_space)', [facts(k).get('plantFor') or '' for k in LOTS], inp=True)
row('mw_hi_final', 'Largest that fits, high, after plant cap = min(high, plant-fit size)', formula=lambda c: f'IF({R("plantfor",c)}="",{R("mw_hi",c)},IF(AND({R("mw_hi",c)}>{R("plantfor",c)},{R("plantfor",c)}<mw_cap),{R("plantfor",c)},{R("mw_hi",c)}))')
row('chk_lo', 'Whole-building check, low = gross ÷ 1,109', formula=lambda c: f'IF(AND({R("new",c)}=1,{R("whole",c)}=1),{R("gross_lo",c)}/gross_per_mw_hi,"")')
row('chk_hi', 'Whole-building check, high = gross ÷ 505', formula=lambda c: f'IF(AND({R("new",c)}=1,{R("whole",c)}=1),{R("gross_hi",c)}/gross_per_mw_lo,"")')
row('servers', 'Servers per MW of IT (register: P_IT ÷ 600 W)', formula=lambda c: '1000/(w_per_server/1000)')
row('racks', 'Racks per MW at 8 kW', formula=lambda c: '1000/kw_per_rack')

s.append(['ENGINE (pasted from golden runs, for comparison only)'])
def ev(k, key, i):
    f = facts(k); v = f.get(key)
    return v[i] if isinstance(v, list) and len(v) > i and v[i] is not None else None
pairs = [('gross_lo', 'gross', 0), ('gross_hi', 'gross', 1), ('hg_lo', 'hallGross', 0), ('hg_hi', 'hallGross', 1),
         ('hs_lo', 'hallS', 0), ('hs_hi', 'hallS', 1), ('hl_lo', 'hallL', 0), ('hl_hi', 'hallL', 1),
         ('mw_lo', 'mw', 0), ('mw_hi_final', 'mw', 1), ('chk_lo', 'mwCheck', 0), ('chk_hi', 'mwCheck', 1)]
for key, ek, i in pairs:
    row('E_' + key, f'engine {ek}[{i}]', [ev(k, ek, i) for k in LOTS], inp=True)
s.append(['DIFFERENCE (blank = not applicable; 0 = match)'])
for key, ek, i in pairs:
    row('D_' + key, f'Δ {key}', formula=lambda c, key=key: f'IF(OR({R("E_"+key,c)}="",ISTEXT({R(key,c)}),{R(key,c)}=""),"",ROUND(ABS({R(key,c)}-{R("E_"+key,c)}),3))')

s.column_dimensions['A'].width = 58
for j in range(len(LOTS)): s.column_dimensions[L(2 + j)].width = 14
s.freeze_panes = 'B4'
out = os.path.join(HERE, 'size_check.xlsx'); wb.save(out); print('wrote', out)
json.dump(rows, open(os.path.join(HERE, 'size_check_rows.json'), 'w'))
