"""Builds checks/emit_check/emit_check.xlsx and emit_check.json: an INDEPENDENT hand calculation of the
"What it gives off" (characterise) stage of Beyond Data, for the same eight lots as checks/place_check.

Written from the register only (register/formula_register.json: the 53 quantities with stage "characterise",
their formula / coef / inputs / sources / source_roles, and the new_sources / old_sources texts), raw data under
data/ (CWEC and PCIC weather files, NIR 2026 Annex 7 extract, ECCC normals extract, NRCan PV values), and the
earlier independent stages checks/input_check/input_check.json and checks/place_check/place_check.json.
The tool's JavaScript, index.html, checks/golden, checks/place_check/screen, checks/agents and *_tabs.json were NOT read.

Method follows checks/place_check/build_place_check.py: lots across the columns, one row block per register part,
every computed cell a live Excel formula, constants on the Constants sheet (named ranges) with the register part
and source id they come from. The workbook is recalculated with LibreOffice and the cached values are read back
to write emit_check.json, so the JSON is the workbook's answer.

Size-stage values (cooling type, plant schedule, generator capacity, floor area, massing, server count) are not
available. Where a part can only be bounded, the bound is labelled; where it cannot, the row says
"not computable (needs size stage: ...)".

Run:  python3 build_emit_check.py          (well under a minute)
"""
import json, os, subprocess, shutil, tempfile, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
D = os.path.join(ROOT, 'data')
REG = json.load(open(os.path.join(ROOT, 'register', 'formula_register.json'), encoding='utf-8'))
Q = {x['id']: x for x in REG['quantities'] if x.get('stage') == 'characterise'}
IC = json.load(open(os.path.join(ROOT, 'checks', 'input_check', 'input_check.json'), encoding='utf-8'))['lots']
PC = json.load(open(os.path.join(ROOT, 'checks', 'place_check', 'place_check.json'), encoding='utf-8'))['lots']

LOTS = [  # key, label, city, parcel id, IT load MW (same list as build_place_check.py)
    ('V09', '275 1st Ave E', 'van', '63819363', 2.5),
    ('E05', '8803 Stadium Rd NW', 'edm', '51053', 10.2),
    ('CLARK', '1425 Clark Dr', 'van', '23060741', 5.2),
    ('E123', '12603 123 Street NW', 'edm', '883345', 22.2),
    ('V339', '339 1st Ave E', 'van', '63819485', 3.8),
    ('E06A', '74 Airport Rd NW', 'edm', '176767', 10.2),
    ('E06B', '64 Airport Rd NW', 'edm', '146402', 10.2),
    ('HAST', '555 W Hastings', 'van', '59013407', 1.1),
]
KEYS = [l[0] for l in LOTS]

# ===================================================================== 1. raw data read here
def epw(path):
    T, Td = [], []
    for i, line in enumerate(open(path, encoding='latin-1')):
        if i < 8: continue
        f = line.split(',')
        T.append(float(f[6])); Td.append(float(f[7]))   # EPW field 7 dry bulb, field 8 dew point (°C)
    assert len(T) == 8760, path
    return T, Td
WX = {
    'van_now': epw(os.path.join(D, '51376785-CAN_BC_VANCOUVER-INTL-A_1108395_CWEC.epw')),
    'edm_now': epw(os.path.join(D, 'effbc917-CAN_AB_EDMONTON-INTL-A_3012216_CWEC.epw')),
    'van_245': epw(os.path.join(D, 'raw', 'weather_future', 'MORPHED_SSP245_2050s_CAN_BC_VANCOUVER-INTL-A_CWEC2020.epw')),
    'van_585': epw(os.path.join(D, 'raw', 'weather_future', 'MORPHED_SSP585_2050s_CAN_BC_VANCOUVER-INTL-A_CWEC2020.epw')),
    'edm_245': epw(os.path.join(D, 'raw', 'weather_future', 'MORPHED_SSP245_2050s_CAN_AB_EDMONTON-INTL-A_CWEC2020.epw')),
    'edm_585': epw(os.path.join(D, 'raw', 'weather_future', 'MORPHED_SSP585_2050s_CAN_AB_EDMONTON-INTL-A_CWEC2020.epw')),
}
WXCOL = {'van_now_T': 'B', 'van_now_Td': 'C', 'edm_now_T': 'D', 'edm_now_Td': 'E',
         'van_245_T': 'F', 'van_585_T': 'G', 'edm_245_T': 'H', 'edm_585_T': 'I'}
def wxr(key): c = WXCOL[key]; return f'Weather!${c}$2:${c}$8761'

NIR = json.load(open(os.path.join(D, 'nir2026_annex7_intensity_extract.json'), encoding='utf-8'))
def nir(prov):
    t = [v for v in NIR.values() if prov in v.get('title', '')][0]
    return [v for k, v in t.items() if k.startswith('Consumption Intensity')][0]['2024a']
EF_BC, EF_AB = nir('British Columbia'), nir('Alberta')
NORM = json.load(open(os.path.join(D, 'normals_1991_2020_extract.json'), encoding='utf-8'))
def norm(st, k): return float(NORM[st][k]['year'])
PV = json.load(open(os.path.join(D, 'nrcan_pv_annual.json'), encoding='utf-8'))

# ===================================================================== 2. constants
CONST = [  # name, value, unit, register part, source id(s) and note, trust
    ('lf', 0.6, '—', 'load_factor', 'fan2007 (≈0.6; one 2007 Google site — check 2 says label as such)', 80),
    ('hours', 8760, 'h/yr', 'load_factor', 'fan2007 / load_factor coef "8,760 h/yr"', 100),
    ('pue_lo', 1.4, '—', 'pue', 'lbnl2024 (US average 2023), register "use 1.4–1.55"', 80),
    ('pue_hi', 1.55, '—', 'pue', 'register "use 1.4–1.55" (uptime2025 INDUSTRY SURVEY 1.54)', 80),
    ('ht_air_lo', 25, '°C', 'heat_temp', 'formula "air-cooled 25–35" (lorenz → wahlroos2018; unverified)', 70),
    ('ht_air_hi', 35, '°C', 'heat_temp', 'same', 70),
    ('ht_liq_lo', 50, '°C', 'heat_temp', 'formula "liquid-cooled 50–60" — coef: unverified, secondary only', 50),
    ('ht_liq_hi', 60, '°C', 'heat_temp', 'same', 50),
    ('ht_nrel_lo', 35, '°C', 'heat_temp', 'coef 19 Oct: NREL ESIF measured design return 35–40 °C', 70),
    ('ht_nrel_hi', 40, '°C', 'heat_temp', 'nrelESIF2019', 70),
    ('cap_cp_lo', 0.5, '—', 'liquid_capture', 'formula 0.5–0.8 cold plate — role: "not stated in a source" (reading of 48–75 % / 94 %)', 50),
    ('cap_cp_hi', 0.8, '—', 'liquid_capture', 'same', 50),
    ('cap_imm', 1.0, '—', 'liquid_capture', 'formula "1.0 immersion" — role: no primary named', 40),
    ('eta_lo', 0.40, '—', 'cop', 'formula "η 0.40–0.50 for built systems" (calibrated 13 Oct, calibrate.py)', 70),
    ('eta_hi', 0.50, '—', 'cop', 'same', 70),
    ('ode_src_in', 27, '°C', 'cop (check)', 'coef: Odense source 27→15 °C (fjfAars2020 / dfKatalog2022)', 80),
    ('ode_src_out', 15, '°C', 'cop (check)', 'same', 80),
    ('ode_snk_ret', 40, '°C', 'cop (check)', 'coef: Odense sink 40→75 °C', 80),
    ('ode_snk_sup', 75, '°C', 'cop (check)', 'same', 80),
    ('ode_cop_reg', 4.56, '—', 'cop (check)', 'coef: "full Lorenz with η 0.50 gives 4.56 for Odense"', 80),
    ('kelvin0', 273.15, 'K', 'cop', 'physical constant (°C → K)', 100),
    ('cop_lo', 3.0, '—', 'cop', 'coef "tool range 3.0–4.5" (receiver_temp needed for a lot-specific COP)', 70),
    ('cop_hi', 4.5, '—', 'cop', 'same', 70),
    ('cp', 4.18, 'kJ/kg·K', 'heat_in_water', 'heatPhys (old source)', 100),
    ('hiw_reg', 1.16, 'kWh/m³·K', 'heat_in_water', 'coef "1.16 kWh per m³ per °C" (greenGridWP29 via heatPhys; unreachable)', 100),
    ('fc_in_lo', 18, '°C inlet', 'free_cooling_hours', 'ashrae (TC 9.9 recommended 18–27 °C)', 90),
    ('fc_in_mid1', 21, '°C inlet', 'free_cooling_hours', 'coef lists 18 / 21 / 24 / 27', 90),
    ('fc_in_mid2', 24, '°C inlet', 'free_cooling_hours', 'same', 90),
    ('fc_in_hi', 27, '°C inlet', 'free_cooling_hours', 'same', 90),
    ('fc_app_air', 1.5, 'K', 'free_cooling_hours', 'formula "approach +1.5 air-side" — coef: assumed', 60),
    ('fc_loop_lo', 30, '°C loop', 'free_cooling_hours', 'coef "dry cooler to 30 °C liquid loop … to 40 °C"', 60),
    ('fc_loop_hi', 40, '°C loop', 'free_cooling_hours', 'same', 60),
    ('fc_app_dc', 12, 'K', 'free_cooling_hours', 'formula "+12 dry cooler" — coef: assumed', 60),
    ('fc_full', 8, '°C outdoor', 'free_cooling_hours', 'coef: "full free cooling for an air-cooled hall … only at ≤8 °C"', 60),
    ('fc50_hot', 18, '°C', 'free_cooling_hours_2050s', 'coef "hours > 18 °C"', 80),
    ('fc50_vhot', 30, '°C', 'free_cooling_hours_2050s', 'coef "hours ≥ 30 °C"', 80),
    ('dew_lim', 15, '°C dew point', 'condensate', 'ashraeTC99card (dew point −9 to 15 °C; unreachable)', 90),
    ('two_n', 0.5, 'share of installed', 'spare_capacity', 'formula "2N holds half" — heatPhys internal note; 2N not defined in uptimeTiers text', 50),
    ('ups_min_lo', 1, 'min', 'backup_hours', 'coef: hyperscale 1–2 min (schneiderUPS2020, vendor — OPERATOR CLAIM)', 50),
    ('ups_min_hi', 15, 'min', 'backup_hours', 'coef: financial 10–15 min (same claim)', 50),
    ('fuel_h_lo', 12, 'h', 'backup_hours / fuel_on_site', 'uptimeHeslin2014: 12 h fuel minimum (industry certifier; unreachable in check 2)', 60),
    ('fuel_h_hi', 48, 'h', 'backup_hours / fuel_on_site', 'coef "24–48 h typical (operator claims)"', 50),
    ('flex_lo_sh', 0.10, 'share', 'flexible_load', 'cfd: Call for Demand curtailment ≥10 %, 4 h', 40),
    ('flex_lo_h', 4, 'h/event', 'flexible_load', 'cfd', 40),
    ('flex_hi_sh', 0.25, 'share', 'flexible_load', 'emeraldPhoenix2025: 25 % for 3 h (company-authored preprint — CLAIM)', 40),
    ('flex_hi_h', 3, 'h/event', 'flexible_load', 'emeraldPhoenix2025 (claim)', 40),
    ('h_fg', 2400, 'kJ/kg', 'evaporation', 'formula "h_fg ≈ 2,400 kJ/kg" (femp; our derivation per femp check 2)', 100),
    ('C_lo', 3, 'cycles', 'blowdown', 'coef gives no C; femp text: "3→6 cycles cuts make-up ~20 %" (example values)', 60),
    ('C_hi', 6, 'cycles', 'blowdown', 'same', 60),
    ('femp_cut', 0.20, 'share', 'blowdown (check)', 'femp: 3→6 cycles cuts make-up ~20 %', 60),
    ('wue_cap', 2, 'L/kWh', 'wue', 'coef "DCRE cap 2 L/kWh" (DCRE not in the part\'s sources)', 50),
    ('w_thermo', 1.25, 'L/kWh', 'water_behind_electricity', 'coef "US thermoelectric ~1.25 L/kWh" (role femp; not in femp text)', 50),
    ('w_hydro', 16.8, 'L/kWh', 'water_behind_electricity', 'coef "hydro ~16.8 L/kWh" (same)', 50),
    ('rain_van', norm('VANCOUVER INTERNATIONAL (AIRPORT)', 'Precipitation | Rainfall (mm)'), 'mm/yr', 'rainwater', 'eccc: data/normals_1991_2020_extract.json, Vancouver Intl rainfall', 80),
    ('prec_van', norm('VANCOUVER INTERNATIONAL (AIRPORT)', 'Precipitation | Precipitation (mm)'), 'mm/yr', 'rainwater', 'eccc: same file, Vancouver Intl precipitation (role: R 1,160 mm)', 80),
    ('rain_edm', norm('EDMONTON (CITY)', 'Precipitation | Rainfall (mm)'), 'mm/yr', 'rainwater', 'eccc: same file, Edmonton City rainfall', 80),
    ('prec_edm', norm('EDMONTON (CITY)', 'Precipitation | Precipitation (mm)'), 'mm/yr', 'rainwater', 'eccc: same file, Edmonton City precipitation (role: R 423 mm)', 80),
    ('rw_C', 0.8, '—', 'rainwater', 'coef "C = 0.8 assumed (CSA B805 not read)"', 50),
    ('ef_bc', EF_BC, 'g CO2e/kWh', 'grid_co2', 'nir2026: data/nir2026_annex7_intensity_extract.json Table A7-11, 2024a consumption intensity (2024 flagged "a")', 80),
    ('ef_ab', EF_AB, 'g CO2e/kWh', 'grid_co2', 'nir2026: same, Table A7-10', 80),
    ('egdf_26', 0.4303, 't CO2e/MWh', 'grid_co2', 'abEGDF v3.3 Table 1, 2026 (displacement by renewables; offset policy number)', 90),
    ('egdf_26_ll', 0.4588, 't CO2e/MWh', 'grid_co2', 'abEGDF v3.3 Table 1, 2026 with line loss', 90),
    ('co2_l', 2.681, 'kg CO2/L diesel', 'generator_co2', 'formula; role credits eccc (Climate Normals entry — no such factor in it) ?', 40),
    ('fuel_kwh', 0.263, 'L/kWh', 'fuel_on_site', 'cat3516c (one Cat model, MANUFACTURER CLAIM; unreachable in check 2)', 50),
    ('fuel_chk', 6.3, 'm³ per MW per 24 h', 'fuel_on_site (check)', 'coef "≈6.3 m³ per MW per 24 h"', 60),
    ('t_full', 4, 'h/yr per unit at full load', 'genset_test', 'auc_27938 (derived: full load up to 2 h twice a year)', 65),
    ('t_all', 14.5, 'h/yr per unit, all tests', 'genset_test', 'auc_27938 (derived); used here at full load as an upper bound', 65),
    ('t_manor', 564, 'h/yr (basis not stated)', 'generator_co2', 'coef "testing hours (Manor Farm 564 h/yr)" — no source in the part list', 30),
    ('ef_co2_g', 706, 'g CO2/kWh', 'genset_test', 'ap42_34 Table 3.4-1 (upper bound for modern engines; > 600 hp)', 65),
    ('ef_nox_g', 14.6, 'g NOx/kWh', 'genset_test / genset_emission_rate', 'ap42_34', 65),
    ('ef_pm_g', 0.43, 'g PM/kWh', 'genset_test', 'ap42_34', 65),
    ('nox_lbhp', 0.024, 'lb/hp·h', 'generator_nox_pm', 'coef "NOx 0.024 lb/hp·h uncontrolled" (AP-42 named in formula; ap42_34 not in this part\'s sources)', 65),
    ('kg_lb', 0.45359237, 'kg/lb', 'generator_nox_pm', 'unit conversion', 100),
    ('kw_hp', 0.745699872, 'kW/hp', 'generator_nox_pm', 'unit conversion (mechanical hp)', 100),
    ('ger_chk_kw', 2000, 'kWe', 'genset_emission_rate (check)', 'coef "2,000 kWe at full load → 8.1 g/s"', 65),
    ('ger_chk_gs', 8.1, 'g/s', 'genset_emission_rate (check)', 'same', 65),
    ('nfc_1st', 200000, 'L per room', 'fuel_on_site', 'nfc2020_4313 Table 4.3.13.4.-B (first storey; Class II)', 60),
    ('nfc_up', 20000, 'L per room', 'fuel_on_site', 'nfc2020_4313 (basement / above first storey; table layout unclear in text copy ?)', 60),
    ('lw_low', 98, 'dBA per unit', 'noise_level', 'traneRTAF size 280 (1,010 kW) low-noise — MANUFACTURER CLAIM; coef: "one unit per MW"', 50),
    ('lw_std', 101, 'dBA per unit', 'noise_level', 'traneRTAF size 280 standard — MANUFACTURER CLAIM', 50),
    ('hemi', 8, 'dB', 'noise_level', 'iso9613 (old): "Lp ≈ Lw − 20·log10(r) − 8 (our simplification)"', 70),
    ('ark_lo', 50, 'jobs', 'jobs', 'coef: Ark UP4 applicant "between 50 and 100" (APPLICANT CLAIM) — source not in the part\'s list', 50),
    ('ark_hi', 100, 'jobs', 'jobs', 'same', 50),
    ('ark_mw', 28, 'MW(IT)', 'jobs', 'same', 50),
    ('ark_m2', 15410, 'm² GIA', 'jobs (check)', 'same', 50),
    ('jl_m2', 465, 'm² per job', 'jobs', 'jlarc2024 (one per 5,000 ft²; operator claim relayed)', 60),
    ('e_grey_lo', 0.75, 'kWh/m³', 'treatment_energy', 'arden2021 (unreachable; register entry states no figures)', 40),
    ('e_grey_hi', 1.1, 'kWh/m³', 'treatment_energy', 'same', 40),
    ('e_mix_lo', 0.94, 'kWh/m³', 'treatment_energy', 'same (mixed wastewater)', 40),
    ('e_mix_hi', 1.28, 'kWh/m³', 'treatment_energy', 'same', 40),
    ('br_mu', 0.84, '× make-up', 'blowdown_recovery', 'gsa_gpg052 (one US test; make-up −16 %)', 50),
    ('br_bd', 0.47, '× blowdown', 'blowdown_recovery', 'gsa_gpg052 (blowdown −53 %)', 50),
    ('br_kwh', 3541, 'kWh/yr per skid', 'blowdown_recovery', 'formula (skid count not stated)', 50),
    ('br_c0', 2.8, 'cycles', 'blowdown_recovery (check)', 'coef "cycles 2.8 → 4.9"', 50),
    ('br_c1', 4.9, 'cycles', 'blowdown_recovery (check)', 'same', 50),
    ('kwp_lo', 0.16, 'kWp/m²', 'pv_roof / pv_facade', 'nrel (2014 median)', 70),
    ('kwp_hi', 0.20, 'kWp/m²', 'pv_roof / pv_facade', 'nrel (current)', 70),
    ('y_van_h', PV['Vancouver']['horizontal'], 'kWh/kWp', 'pv_roof', 'nrcanPVmunic; data/nrcan_pv_annual.json', 80),
    ('y_van_l', PV['Vancouver']['latitude'], 'kWh/kWp', 'pv_roof', 'same', 80),
    ('y_van_v', PV['Vancouver']['vertical'], 'kWh/kWp', 'pv_facade', 'same', 80),
    ('y_edm_h', PV['Edmonton']['horizontal'], 'kWh/kWp', 'pv_roof', 'same', 80),
    ('y_edm_l', PV['Edmonton']['latitude'], 'kWh/kWp', 'pv_roof', 'same', 80),
    ('y_edm_v', PV['Edmonton']['vertical'], 'kWh/kWp', 'pv_facade', 'same', 80),
    ('allen_van', 119, 'W/m² (units assumed)', 'heat_flux_site', 'allen2011 (peak grid cell, not a city average — check 1)', 70),
    ('sv_mm', 15, 'mm', 'stormwater_volume', 'vbbl2025_2425 stage 1', 90),
    ('sv_rel1', 5, 'L/s/ha', 'stormwater_volume', 'vbbl2025_2425', 90),
    ('sv_small', 1000, 'm² (small-site limit)', 'stormwater_volume', 'vbbl2025_2425 (≤1,000 m² and FSR ≤ 1.0)', 90),
    ('se_100', 375, 'm³/ha', 'stormwater_volume', 'epcorStorm Exhibit 2, 100-year, C = 0.90 (guideline)', 80),
    ('se_200', 507, 'm³/ha', 'stormwater_volume', 'epcorStorm 200-year', 80),
    ('se_rel', 35, 'L/s/ha', 'stormwater_volume', 'epcorStorm (0.035 m³/s per ha)', 80),
    ('se_min_ha', 0.16, 'ha', 'stormwater_volume', 'epcorStorm (control required above 0.16 ha)', 80),
    ('se_max_ha', 3.0, 'ha', 'stormwater_volume', 'epcorStorm (method up to 3.0 ha)', 80),
    ('emb_srv', 4177, 'kg CO2e per server (life cycle, US use)', 'embodied_it', 'dellR760 — CLAIM (vendor LCA)', 50),
    ('emb_sh_lo', 0.10, 'share', 'embodied_it', 'coef "manufacturing ≈10–15 %" (?)', 40),
    ('emb_sh_hi', 0.15, 'share', 'embodied_it', 'same', 40),
    ('emb_life', 4, 'yr', 'embodied_it', 'dellR760 text: 4-year life (LCA assumption)', 50),
    ('ref_leak', 0.03, '/yr', 'refrigerant_co2e', 'formula 3 %/yr (IPCC 2–15 %; role lists eccc)', 60),
]
CONST_NAMES = {c[0] for c in CONST}

# ===================================================================== 3. workbook
def build():
    wb = Workbook()
    c = wb.active; c.title = 'Constants'
    c.append(['name', 'value', 'unit', 'register part', 'source (register new_sources / old_sources id) and note', 'trust'])
    for row in CONST: c.append(list(row))
    for i, row in enumerate(CONST, start=2):
        wb.defined_names[row[0]] = DefinedName(row[0], attr_text=f"Constants!$B${i}")
    for col, w in zip('ABCDEF', (14, 12, 30, 34, 100, 7)): c.column_dimensions[col].width = w
    for cell in c[1]: cell.font = Font(bold=True)

    w = wb.create_sheet('Weather')
    w.append(['hour', 'Van CWEC dry bulb °C', 'Van CWEC dew point °C', 'Edm CWEC dry bulb °C', 'Edm CWEC dew point °C',
              'Van 2050s SSP2-4.5 dry bulb', 'Van 2050s SSP5-8.5 dry bulb', 'Edm 2050s SSP2-4.5 dry bulb', 'Edm 2050s SSP5-8.5 dry bulb'])
    for h in range(8760):
        w.append([h + 1, WX['van_now'][0][h], WX['van_now'][1][h], WX['edm_now'][0][h], WX['edm_now'][1][h],
                  WX['van_245'][0][h], WX['van_585'][0][h], WX['edm_245'][0][h], WX['edm_585'][0][h]])
    w['K1'] = ('Sources: cwec (data/51376785-…VANCOUVER-INTL-A…CWEC.epw, data/effbc917-…EDMONTON-INTL-A…CWEC.epw); '
               'pcicWx3 (data/raw/weather_future/MORPHED_SSP245/585_2050s_…CWEC2020.epw). EPW fields 7 (dry bulb) and 8 (dew point).')

    s = wb.create_sheet('Emit', 0)
    s['A1'] = 'What it gives off — independent hand calculation (live formulas)'; s['A1'].font = Font(bold=True, size=13)
    s['A2'] = ('Blue = input (given, input_check, place_check); black = formula; grey italic = text (not computable / not applicable). '
               '_lo pairs PUE 1.4 with the low coefficients, _hi pairs PUE 1.55 with the high ones. All arithmetic is ours.')
    s.cell(3, 1, 'lot').font = Font(bold=True)
    for j, (k, l, cty, p, it) in enumerate(LOTS):
        s.cell(3, 2 + j, k).font = Font(bold=True); s.cell(4, 2 + j, l)
    s.cell(4, 1, 'address')
    rows = {}
    blue = Font(color='1F4E9E'); grey = Font(color='808080', italic=True)
    def row(key, label, vals=None, f=None, inp=False):
        r = s.max_row + 1; rows[key] = r; s.cell(r, 1, label)
        for j, (k, l, cty, p, it) in enumerate(LOTS):
            if inp:
                v = vals[j]; cell = s.cell(r, 2 + j, v)
                cell.font = grey if isinstance(v, str) and not v.startswith('=') else blue
            else:
                e = f(L(2 + j), cty, k)
                if isinstance(e, str) and e.startswith('='): s.cell(r, 2 + j, e)
                else: s.cell(r, 2 + j, e).font = grey
        return r
    R = lambda key, col: f'{col}{rows[key]}'
    def head(t): r = s.max_row + 2; s.cell(r, 1, t).font = Font(bold=True)
    city = lambda c, a, b: f'=IF({R("van",c)}=1,{a},{b})'
    NC = lambda why: (lambda c, t, k: why)

    head('INPUTS')
    row('van', 'Vancouver (1) or Edmonton (0)', [1 if x[2] == 'van' else 0 for x in LOTS], inp=True)
    row('it', 'IT load, MW of servers (given; same as place_check)', [x[4] for x in LOTS], inp=True)
    row('lot', 'Lot area m² (input_check 1)', [IC[k]['1_lot_area_m2']['value'] for k in KEYS], inp=True)
    row('fp', 'Existing buildings on lot, footprint m² (input_check 3)', [IC[k]['3_buildings']['value']['total_footprint_m2'] for k in KEYS], inp=True)
    row('rec', 'Nearest residential parcel, m, lot edge to parcel edge (input_check 8)', [IC[k]['8_nearest_receptor_m']['value'] for k in KEYS], inp=True)
    row('gw200', 'Greywater from residents within 200 m, ML/yr (place_check greywater_neighbours low)', [PC[k]['greywater_neighbours']['low'] for k in KEYS], inp=True)
    row('gw500', 'Greywater from residents within 500 m, ML/yr (place_check greywater_neighbours high)', [PC[k]['greywater_neighbours']['high'] for k in KEYS], inp=True)

    head('LOAD (size-stage inputs that the register fixes: IT given, load factor 0.6, PUE 1.4–1.55)')
    row('itavg', 'it_load_actual: average IT MW = IT × 0.6', f=lambda c, t, k: f'={R("it",c)}*lf')
    row('pk_lo', 'Facility peak MW, low = IT × 1.4', f=lambda c, t, k: f'={R("it",c)}*pue_lo')
    row('pk_hi', 'Facility peak MW, high = IT × 1.55', f=lambda c, t, k: f'={R("it",c)}*pue_hi')
    row('mwh_lo', 'Facility electricity MWh/yr, low = IT × 0.6 × 1.4 × 8,760', f=lambda c, t, k: f'={R("itavg",c)}*pue_lo*hours')
    row('mwh_hi', 'Facility electricity MWh/yr, high (PUE 1.55)', f=lambda c, t, k: f'={R("itavg",c)}*pue_hi*hours')
    row('itmwh', 'E_IT MWh/yr = IT × 0.6 × 8,760', f=lambda c, t, k: f'={R("itavg",c)}*hours')

    head('load_factor / pue')
    row('lf_v', 'load_factor (average ÷ sum of server maxima)', f=lambda c, t, k: '=lf')
    row('pue_l', 'PUE low', f=lambda c, t, k: '=pue_lo')
    row('pue_h', 'PUE high', f=lambda c, t, k: '=pue_hi')

    head('heat_rejected — MWh/yr (heat ≈ electricity drawn, Barroso)')
    row('hr_lo', 'Heat rejected, low = facility MWh at PUE 1.4', f=lambda c, t, k: f'={R("mwh_lo",c)}')
    row('hr_hi', 'Heat rejected, high = facility MWh at PUE 1.55', f=lambda c, t, k: f'={R("mwh_hi",c)}')
    row('hr_pk_lo', 'Peak heat rejected MW, low (= facility peak)', f=lambda c, t, k: f'={R("pk_lo",c)}')
    row('hr_pk_hi', 'Peak heat rejected MW, high', f=lambda c, t, k: f'={R("pk_hi",c)}')

    head('heat_temp — °C (two runs: cooling type is not lot data)')
    row('ht_a_lo', 'Air-cooled, low', f=lambda c, t, k: '=ht_air_lo')
    row('ht_a_hi', 'Air-cooled, high', f=lambda c, t, k: '=ht_air_hi')
    row('ht_l_lo', 'Liquid-cooled, low (unverified per coef)', f=lambda c, t, k: '=ht_liq_lo')
    row('ht_l_hi', 'Liquid-cooled, high (unverified per coef)', f=lambda c, t, k: '=ht_liq_hi')
    row('ht_n', 'Liquid-cooled measured return, NREL ESIF (high end)', f=lambda c, t, k: '=ht_nrel_hi')

    head('liquid_capture — share (air-cooled: no figure in the register)')
    row('lc_lo', 'Cold plate, low', f=lambda c, t, k: '=cap_cp_lo')
    row('lc_hi', 'Cold plate, high', f=lambda c, t, k: '=cap_cp_hi')
    row('lc_im', 'Immersion', f=lambda c, t, k: '=cap_imm')
    row('lc_air', 'Air-cooled', f=NC('not stated (?): coef "no capture share for a whole air-cooled hall"'))

    head('cop — Lorenz; lot COP needs receiver_temp (match stage), so the register tool range is carried downstream')
    tlm = lambda a, b: f'(({a})-({b}))/LN((({a})+kelvin0)/(({b})+kelvin0))'
    snk = tlm('ode_snk_sup', 'ode_snk_ret'); src = tlm('ode_src_in', 'ode_src_out')
    row('cop_ode', 'Check: Lorenz COP for Odense (27→15 / 40→75 °C), η 0.50 — register says 4.56', f=lambda c, t, k: f'=eta_hi*({snk})/(({snk})-({src}))')
    row('cop_ode40', 'Check: same at η 0.40', f=lambda c, t, k: f'=eta_lo*({snk})/(({snk})-({src}))')
    row('cop_lot', 'COP for this lot', f=NC('not computable (needs receiver_temp: set by receiver type, not lot data)'))
    row('cop_l', 'COP carried downstream, low (register tool range)', f=lambda c, t, k: '=cop_lo')
    row('cop_h', 'COP carried downstream, high', f=lambda c, t, k: '=cop_hi')

    head('heat_captured — E_reuse MWh/yr = heat_rejected × liquid_capture (liquid-cooled run)')
    row('hc_lo', 'E_reuse, low = HR low × 0.5', f=lambda c, t, k: f'={R("hr_lo",c)}*cap_cp_lo')
    row('hc_hi', 'E_reuse, high = HR high × 0.8', f=lambda c, t, k: f'={R("hr_hi",c)}*cap_cp_hi')
    row('hc_im', 'E_reuse, immersion = HR high × 1.0', f=lambda c, t, k: f'={R("hr_hi",c)}*cap_imm')
    row('hc_it', 'Cross-check: capture share on IT heat only, high (E_IT × 0.8)', f=lambda c, t, k: f'={R("itmwh",c)}*cap_cp_hi')
    row('hc_air', 'Air-cooled run', f=NC('not computable (no air-cooled capture share in the register)'))

    head('heat_delivered — MWh/yr available after the heat pump = E_reuse × COP ÷ (COP − 1)')
    row('hd_lo', 'Low = E_reuse low × 4.5 ÷ 3.5', f=lambda c, t, k: f'={R("hc_lo",c)}*cop_hi/(cop_hi-1)')
    row('hd_hi', 'High = E_reuse high × 3.0 ÷ 2.0', f=lambda c, t, k: f'={R("hc_hi",c)}*cop_lo/(cop_lo-1)')

    head('hp_electricity — MWh/yr = Q_out ÷ COP')
    row('hp_lo', 'Low = Q_out low ÷ 4.5', f=lambda c, t, k: f'={R("hd_lo",c)}/cop_hi')
    row('hp_hi', 'High = Q_out high ÷ 3.0', f=lambda c, t, k: f'={R("hd_hi",c)}/cop_lo')

    head('erf — E_reuse ÷ E_DC')
    row('erf_lo', 'ERF low = E_reuse low ÷ HR low', f=lambda c, t, k: f'={R("hc_lo",c)}/{R("hr_lo",c)}')
    row('erf_hi', 'ERF high = E_reuse high ÷ HR high', f=lambda c, t, k: f'={R("hc_hi",c)}/{R("hr_hi",c)}')

    head('heat_in_water — kWh per m³ per K (no lot volume or ΔT named)')
    row('hiw', 'Check: 4.18 × 1,000 ÷ 3,600', f=lambda c, t, k: '=cp*1000/3600')
    row('hiw_r', 'Register coefficient', f=lambda c, t, k: '=hiw_reg')

    head('free_cooling_hours — h/yr in the CWEC year (≤ threshold; the register counts match ≤, the formula text says <)')
    def cnt(op, thr, now='now'):
        return lambda c, t, k: city(c, f'COUNTIF({wxr("van_"+now+"_T")},"{op}"&({thr}))', f'COUNTIF({wxr("edm_"+now+"_T")},"{op}"&({thr}))')
    row('fc8', 'Full free cooling, air-cooled hall: hours ≤ 8 °C', f=cnt('<=', 'fc_full'))
    row('fc18', 'Air-side at 18 °C inlet: hours ≤ 18 − 1.5', f=cnt('<=', 'fc_in_lo-fc_app_air'))
    row('fc21', 'Air-side at 21 °C inlet', f=cnt('<=', 'fc_in_mid1-fc_app_air'))
    row('fc24', 'Air-side at 24 °C inlet', f=cnt('<=', 'fc_in_mid2-fc_app_air'))
    row('fc27', 'Air-side at 27 °C inlet', f=cnt('<=', 'fc_in_hi-fc_app_air'))
    row('fcd30', 'Dry cooler to 30 °C loop: hours ≤ 30 − 12', f=cnt('<=', 'fc_loop_lo-fc_app_dc'))
    row('fcd40', 'Dry cooler to 40 °C loop: hours ≤ 40 − 12', f=cnt('<=', 'fc_loop_hi-fc_app_dc'))
    row('fc18s', 'Formula as written (strict <): air-side at 18 °C inlet', f=cnt('<', 'fc_in_lo-fc_app_air'))

    head('free_cooling_hours_2050s — PCIC morphed files (SSP2-4.5 / SSP5-8.5)')
    for sc, nm in (('245', 'SSP2-4.5'), ('585', 'SSP5-8.5')):
        row(f'f50_8_{sc}', f'{nm}: hours ≤ 8 °C', f=cnt('<=', 'fc_full', sc))
        row(f'f50_18_{sc}', f'{nm}: air-side at 18 °C inlet (≤ 16.5)', f=cnt('<=', 'fc_in_lo-fc_app_air', sc))
        row(f'f50_27_{sc}', f'{nm}: air-side at 27 °C inlet (≤ 25.5)', f=cnt('<=', 'fc_in_hi-fc_app_air', sc))
        row(f'f50_hot_{sc}', f'{nm}: hours > 18 °C', f=cnt('>', 'fc50_hot', sc))
        row(f'f50_vh_{sc}', f'{nm}: hours ≥ 30 °C', f=cnt('>=', 'fc50_vhot', sc))
        row(f'f50_max_{sc}', f'{nm}: peak dry bulb °C', f=lambda c, t, k, sc=sc: city(c, f'MAX({wxr("van_"+sc+"_T")})', f'MAX({wxr("edm_"+sc+"_T")})'))
    row('f50_hot_now', 'Today (CWEC): hours > 18 °C', f=cnt('>', 'fc50_hot'))
    row('f50_max_now', 'Today (CWEC): peak dry bulb °C', f=lambda c, t, k: city(c, f'MAX({wxr("van_now_T")})', f'MAX({wxr("edm_now_T")})'))

    head('condensate — hours a year with outdoor dew point > 15 °C (CWEC)')
    row('cond', 'Hours with dew point > 15 °C', f=lambda c, t, k: city(c, f'COUNTIF({wxr("van_now_Td")},">"&dew_lim)', f'COUNTIF({wxr("edm_now_Td")},">"&dew_lim)'))

    head('spare_capacity — MW (installed capacity and topology are size stage)')
    row('sp', 'Installed − peak', f=NC('not computable (needs size stage: installed electrical capacity)'))
    row('sp_2n_lo', 'IF 2N: spare = half of installed = peak, low (conditional, our arithmetic)', f=lambda c, t, k: f'={R("pk_lo",c)}*two_n/(1-two_n)')
    row('sp_2n_hi', 'IF 2N: same, high', f=lambda c, t, k: f'={R("pk_hi",c)}*two_n/(1-two_n)')

    head('backup_hours — h (register-stated ranges; battery kWh and fuel stored are size stage)')
    row('bk_lo', 'Fuel runtime, low (Uptime 12 h minimum)', f=lambda c, t, k: '=fuel_h_lo')
    row('bk_hi', 'Fuel runtime, high (operator claims 48 h)', f=lambda c, t, k: '=fuel_h_hi')
    row('ups_lo', 'UPS runtime, h, low (1 min, vendor claim)', f=lambda c, t, k: '=ups_min_lo/60')
    row('ups_hi', 'UPS runtime, h, high (15 min, vendor claim)', f=lambda c, t, k: '=ups_min_hi/60')

    head('flexible_load — MW·h per event = flexible MW × hours (on it_load_actual)')
    row('fl_lo', 'Low: 10 % × 4 h (Call for Demand)', f=lambda c, t, k: f'={R("itavg",c)}*flex_lo_sh*flex_lo_h')
    row('fl_hi', 'High: 25 % × 3 h (Emerald preprint, claim)', f=lambda c, t, k: f'={R("itavg",c)}*flex_hi_sh*flex_hi_h')

    head('evaporation — L/yr, upper bound E ≤ Q ÷ h_fg (all heat evaporated; zero for dry cooling)')
    row('ev_lo', 'E low = HR low × 3.6e6 ÷ 2,400', f=lambda c, t, k: f'={R("hr_lo",c)}*1000*3600/h_fg')
    row('ev_hi', 'E high', f=lambda c, t, k: f'={R("hr_hi",c)}*1000*3600/h_fg')
    row('ev_lkwh', 'Check: L per kWh of heat = 3,600 ÷ 2,400 (register "≤ ~1.5")', f=lambda c, t, k: '=3600/h_fg')

    head('blowdown — L/yr: B = E ÷ (C − 1); make-up = E + B')
    row('bd_lo', 'Blowdown low (C 6)', f=lambda c, t, k: f'={R("ev_lo",c)}/(C_hi-1)')
    row('bd_hi', 'Blowdown high (C 3)', f=lambda c, t, k: f'={R("ev_hi",c)}/(C_lo-1)')
    row('mu_lo', 'Make-up low = E + B', f=lambda c, t, k: f'={R("ev_lo",c)}+{R("bd_lo",c)}')
    row('mu_hi', 'Make-up high', f=lambda c, t, k: f'={R("ev_hi",c)}+{R("bd_hi",c)}')
    row('bd_chk', 'Check femp "3→6 cycles cuts make-up ~20 %": 1 − (6/5)/(3/2)', f=lambda c, t, k: '=1-(C_hi/(C_hi-1))/(C_lo/(C_lo-1))')

    head('wue — L/kWh = site water ÷ E_IT (upper bound: evaporative, all heat evaporated)')
    row('wue_lo', 'WUE low = make-up low ÷ E_IT', f=lambda c, t, k: f'={R("mu_lo",c)}/({R("itmwh",c)}*1000)')
    row('wue_hi', 'WUE high', f=lambda c, t, k: f'={R("mu_hi",c)}/({R("itmwh",c)}*1000)')
    row('wue_cap', 'WUE low ÷ DCRE cap 2 L/kWh', f=lambda c, t, k: f'={R("wue_lo",c)}/wue_cap')

    head('water_behind_electricity — L/yr = E × w_grid (E = facility electricity)')
    row('wb_lo', 'Low = MWh low × 1.25 L/kWh (thermoelectric)', f=lambda c, t, k: f'={R("mwh_lo",c)}*1000*w_thermo')
    row('wb_hi', 'High: Van MWh high × 16.8 (hydro; "BC as a range"); Edm × 1.25 (no Alberta value in register)',
        f=lambda c, t, k: city(c, f'{R("mwh_hi",c)}*1000*w_hydro', f'{R("mwh_hi",c)}*1000*w_thermo'))

    head('rainwater — m³/yr = A × R × C (existing building\'s roofs; a new footprint is size stage)')
    row('rw_A', 'Roof area A, m² (existing footprint)', f=lambda c, t, k: f'=IF({R("fp",c)}>0,{R("fp",c)},"not computable (no building on the lot; new footprint is size stage)")')
    row('rw_lo', 'V low = A × rainfall × 0.8', f=lambda c, t, k: f'=IF(ISNUMBER({R("rw_A",c)}),{R("rw_A",c)}*IF({R("van",c)}=1,rain_van,rain_edm)/1000*rw_C,{R("rw_A",c)})')
    row('rw_hi', 'V high = A × precipitation × 0.8', f=lambda c, t, k: f'=IF(ISNUMBER({R("rw_A",c)}),{R("rw_A",c)}*IF({R("van",c)}=1,prec_van,prec_edm)/1000*rw_C,{R("rw_A",c)})')

    head('grid_co2 — t CO2e/yr = E × EF (average consumption intensity, NIR 2024; Alberta displacement factor beside it)')
    ef = lambda c: f'IF({R("van",c)}=1,ef_bc,ef_ab)'
    row('gc_lo', 'Average, low = MWh low × EF ÷ 1,000', f=lambda c, t, k: f'={R("mwh_lo",c)}*{ef(c)}/1000')
    row('gc_hi', 'Average, high', f=lambda c, t, k: f'={R("mwh_hi",c)}*{ef(c)}/1000')
    row('gc_m_lo', 'Alberta displacement factor 2026 × MWh low', f=lambda c, t, k: city(c, '"n/a: BC publishes no marginal factor"', f'{R("mwh_lo",c)}*egdf_26'))
    row('gc_m_hi', 'Alberta displacement factor 2026 with line loss × MWh high', f=lambda c, t, k: city(c, '"n/a: BC publishes no marginal factor"', f'{R("mwh_hi",c)}*egdf_26_ll'))
    row('gc_hp_lo', 'Heat-pump electricity carbon, low (hp low × EF)', f=lambda c, t, k: f'={R("hp_lo",c)}*{ef(c)}/1000')
    row('gc_hp_hi', 'Heat-pump electricity carbon, high', f=lambda c, t, k: f'={R("hp_hi",c)}*{ef(c)}/1000')

    head('GENERATORS — capacity is size stage; facility peak taken as a lower bound (capacity ≥ load; spares not counted)')
    row('gk_lo', 'Generator kW, lower bound low = facility peak low', f=lambda c, t, k: f'={R("pk_lo",c)}*1000')
    row('gk_hi', 'Generator kW, lower bound high', f=lambda c, t, k: f'={R("pk_hi",c)}*1000')
    row('gt_e_lo', 'Test energy kWh/yr, low = kW × 4 h full load', f=lambda c, t, k: f'={R("gk_lo",c)}*t_full')
    row('gt_e_hi', 'Test energy kWh/yr, high = kW × 14.5 h (all test hours at full load)', f=lambda c, t, k: f'={R("gk_hi",c)}*t_all')
    row('gt_chk', 'Check: 14.5 h = 26 × 0.25 h + 4 × 1 h (zero load) + 2 × 2 h (full load), from auc_27938 text', f=lambda c, t, k: '=26*0.25+4*1+2*2')

    head('generator_co2 — t CO2/yr = litres × 2.681')
    row('g_l_lo', 'Litres low = test kWh × 0.263', f=lambda c, t, k: f'={R("gt_e_lo",c)}*fuel_kwh')
    row('g_l_hi', 'Litres high', f=lambda c, t, k: f'={R("gt_e_hi",c)}*fuel_kwh')
    row('gco2_lo', 't CO2 low', f=lambda c, t, k: f'={R("g_l_lo",c)}*co2_l/1000')
    row('gco2_hi', 't CO2 high', f=lambda c, t, k: f'={R("g_l_hi",c)}*co2_l/1000')
    row('gco2_chk', 'Check: 0.263 L/kWh × 2.681 kg/L in g/kWh (AP-42 says 706)', f=lambda c, t, k: '=fuel_kwh*co2_l*1000')
    row('gco2_man', 'Manor Farm 564 h/yr at the low kW (basis not stated — shown only)', f=lambda c, t, k: f'={R("gk_lo",c)}*t_manor*fuel_kwh*co2_l/1000')

    head('generator_nox_pm — kg/yr = AP-42 factor × rated hp × hours')
    row('nox_lo', 'NOx low = kW ÷ 0.7457 × 4 h × 0.024 lb × 0.4536', f=lambda c, t, k: f'={R("gk_lo",c)}/kw_hp*t_full*nox_lbhp*kg_lb')
    row('nox_hi', 'NOx high (14.5 h)', f=lambda c, t, k: f'={R("gk_hi",c)}/kw_hp*t_all*nox_lbhp*kg_lb')
    row('pm_lo', 'PM low = test kWh × 0.43 g ÷ 1,000 (no lb/hp·h PM factor in this part)', f=lambda c, t, k: f'={R("gt_e_lo",c)}*ef_pm_g/1000')
    row('pm_hi', 'PM high', f=lambda c, t, k: f'={R("gt_e_hi",c)}*ef_pm_g/1000')
    row('nox_chk', 'Check: 0.024 lb/hp·h in g/kWh (genset_test says 14.6)', f=lambda c, t, k: '=nox_lbhp*kg_lb*1000/kw_hp')

    head('genset_test — t CO2, kg NOx, kg PM per year = test kWh × AP-42 factors')
    row('gtc_lo', 'CO2 t low = kWh × 706 g', f=lambda c, t, k: f'={R("gt_e_lo",c)}*ef_co2_g/1000000')
    row('gtc_hi', 'CO2 t high', f=lambda c, t, k: f'={R("gt_e_hi",c)}*ef_co2_g/1000000')
    row('gtn_lo', 'NOx kg low = kWh × 14.6 g', f=lambda c, t, k: f'={R("gt_e_lo",c)}*ef_nox_g/1000')
    row('gtn_hi', 'NOx kg high', f=lambda c, t, k: f'={R("gt_e_hi",c)}*ef_nox_g/1000')
    row('gtp_lo', 'PM kg low', f=lambda c, t, k: f'={R("pm_lo",c)}')
    row('gtp_hi', 'PM kg high', f=lambda c, t, k: f'={R("pm_hi",c)}')

    head('genset_emission_rate — g/s NOx = 14.6 × kWe × load ÷ 3,600 (all sets at full load together)')
    row('ger_lo', 'Low', f=lambda c, t, k: f'=ef_nox_g*{R("gk_lo",c)}/3600')
    row('ger_hi', 'High', f=lambda c, t, k: f'=ef_nox_g*{R("gk_hi",c)}/3600')
    row('ger_chk', 'Check: 2,000 kWe → g/s (register 8.1)', f=lambda c, t, k: '=ef_nox_g*ger_chk_kw/3600')

    head('air_screen_1h / no2_at_receptor / hvo_swap')
    row('air', 'C(1 h) µg/m³', f=NC('not computable: σy, σz (urban coefficients) and stack height H not stated in the register (?); u = 1 m/s given'))
    row('no2', 'NO2 µg/m³', f=NC('not computable: needs air_screen_1h; ratios 1.0 (Tier 1), ARM2 0.9–0.2, in-stack 0.083 stated'))
    row('hvo', 'Δ lifecycle CO2e', f=NC('not computable: fossil diesel reference CI "to retrieve" (not stated); RD −7 to 131 g CO2e/MJ'))

    head('fuel_on_site — L = generator kW × runtime h × 0.263 L/kWh (lower bound: kW = facility peak)')
    row('fu_lo', 'Fuel L, low = peak low × 12 h × 0.263', f=lambda c, t, k: f'={R("gk_lo",c)}*fuel_h_lo*fuel_kwh')
    row('fu_hi', 'Fuel L, high = peak high × 48 h × 0.263', f=lambda c, t, k: f'={R("gk_hi",c)}*fuel_h_hi*fuel_kwh')
    row('fu_rooms_lo', 'First-storey rooms at 200,000 L each, low', f=lambda c, t, k: f'=ROUNDUP({R("fu_lo",c)}/nfc_1st,0)')
    row('fu_rooms_hi', 'First-storey rooms, high', f=lambda c, t, k: f'=ROUNDUP({R("fu_hi",c)}/nfc_1st,0)')
    row('fu_up_hi', 'Rooms if basement / upper storey at 20,000 L, high', f=lambda c, t, k: f'=ROUNDUP({R("fu_hi",c)}/nfc_up,0)')
    row('fu_chk', 'Check: m³ per MW per 24 h = 1,000 × 24 × 0.263 ÷ 1,000 (register ≈6.3)', f=lambda c, t, k: '=24*fuel_kwh')

    head('noise_level — dBA screen: Lp = Lw + 10·log10(n) − 20·log10(r) − 8 (n = 1 unit per MW; r = lot edge to nearest residential parcel)')
    row('nz_n_lo', 'Units, low = ROUNDUP(IT MW)', f=lambda c, t, k: f'=ROUNDUP({R("it",c)},0)')
    row('nz_n_hi', 'Units, high = ROUNDUP(facility peak MW at 1.55)', f=lambda c, t, k: f'=ROUNDUP({R("pk_hi",c)},0)')
    row('nz_lw_lo', 'Combined Lw, low (98 dBA low-noise units)', f=lambda c, t, k: f'=lw_low+10*LOG10({R("nz_n_lo",c)})')
    row('nz_lw_hi', 'Combined Lw, high (101 dBA standard units)', f=lambda c, t, k: f'=lw_std+10*LOG10({R("nz_n_hi",c)})')
    row('nz_lo', 'Lp at receptor, low', f=lambda c, t, k: f'={R("nz_lw_lo",c)}-20*LOG10({R("rec",c)})-hemi')
    row('nz_hi', 'Lp at receptor, high', f=lambda c, t, k: f'={R("nz_lw_hi",c)}-20*LOG10({R("rec",c)})-hemi')

    head('heat_flux_site — W/m² = peak heat (W) ÷ lot area')
    row('hf_lo', 'Low = peak low × 1e6 ÷ lot', f=lambda c, t, k: f'={R("pk_lo",c)}*1000000/{R("lot",c)}')
    row('hf_hi', 'High', f=lambda c, t, k: f'={R("pk_hi",c)}*1000000/{R("lot",c)}')
    row('hf_rat', 'High ÷ Vancouver grid-cell 119 W/m² (Allen 2011, peak cell)', f=lambda c, t, k: f'={R("hf_hi",c)}/allen_van')
    row('hf_chk', 'Check: 10 MW over 1 ha (register 1,000 W/m²)', f=lambda c, t, k: '=10*1000000/10000')

    head('heat_plume_dT')
    row('plume', 'ΔT °C', f=NC('not computable: airflow (plant schedule, manufacturer) and plume factor (air_screen_1h) not available'))

    head('jobs — FTE (floor-area route needs size stage; per-MW route from the Ark UP4 applicant figure)')
    row('jb_lo', 'Low = IT × 50 ÷ 28 (applicant claim)', f=lambda c, t, k: f'={R("it",c)}*ark_lo/ark_mw')
    row('jb_hi', 'High = IT × 100 ÷ 28 (applicant claim)', f=lambda c, t, k: f'={R("it",c)}*ark_hi/ark_mw')
    row('jb_chk', 'Check: m² per job at Ark UP4, 100 jobs (register 154)', f=lambda c, t, k: '=ark_m2/ark_hi')
    row('jb_chk2', 'Check: m² per job at 50 jobs (register 308)', f=lambda c, t, k: '=ark_m2/ark_lo')

    head('treatment_energy — kWh/yr = m³ treated × e (greywater route: min(greywater nearby, make-up); recovery share is a register gap)')
    row('tr_lo', 'Low = MIN(V_gw 200 m, make-up low) × 0.75', f=lambda c, t, k: f'=MIN({R("gw200",c)}*1000,{R("mu_lo",c)}/1000)*e_grey_lo')
    row('tr_hi', 'High = MIN(V_gw 500 m, make-up high) × 1.1', f=lambda c, t, k: f'=MIN({R("gw500",c)}*1000,{R("mu_hi",c)}/1000)*e_grey_hi')
    row('tr_mix', 'If all make-up high came from mixed wastewater × 1.28', f=lambda c, t, k: f'={R("mu_hi",c)}/1000*e_mix_hi')

    head('blowdown_recovery — L/yr recovered = blowdown × (1 − 0.47); make-up after = × 0.84')
    row('br_lo', 'Recovered, low', f=lambda c, t, k: f'={R("bd_lo",c)}*(1-br_bd)')
    row('br_hi', 'Recovered, high', f=lambda c, t, k: f'={R("bd_hi",c)}*(1-br_bd)')
    row('br_mu_lo', 'Make-up after, low', f=lambda c, t, k: f'={R("mu_lo",c)}*br_mu')
    row('br_mu_hi', 'Make-up after, high', f=lambda c, t, k: f'={R("mu_hi",c)}*br_mu')
    row('br_kwh', 'Electricity per skid, kWh/yr (skid count not stated)', f=lambda c, t, k: '=br_kwh')
    row('br_chk', 'Check: make-up cut implied by cycles 2.8 → 4.9 (register says 16 %)', f=lambda c, t, k: '=1-(br_c1/(br_c1-1))/(br_c0/(br_c0-1))')
    row('br_chk2', 'Check: blowdown cut implied by cycles 2.8 → 4.9 (register says 53 %)', f=lambda c, t, k: '=1-(1/(br_c1-1))/(1/(br_c0-1))')

    head('pv_roof / pv_facade — kWh/yr per m² (lot area of panels / facade is size stage)')
    row('pvr', 'pv_roof MWh/yr', f=NC('not computable (needs size stage: roof outline, plant footprint and clearances)'))
    row('pvr_lo', 'Per m² of panel, kWh/yr, low = 0.16 × horizontal yield', f=lambda c, t, k: city(c, 'kwp_lo*y_van_h', 'kwp_lo*y_edm_h'))
    row('pvr_hi', 'Per m² of panel, high = 0.20 × latitude-tilt yield', f=lambda c, t, k: city(c, 'kwp_hi*y_van_l', 'kwp_hi*y_edm_l'))
    row('pvf', 'pv_facade MWh/yr', f=NC('not computable (needs size stage: south facade area)'))
    row('pvf_lo', 'Per m² of south facade, kWh/yr, low = 0.16 × vertical yield', f=lambda c, t, k: city(c, 'kwp_lo*y_van_v', 'kwp_lo*y_edm_v'))
    row('pvf_hi', 'Per m² of south facade, high = 0.20 × vertical yield', f=lambda c, t, k: city(c, 'kwp_hi*y_van_v', 'kwp_hi*y_edm_v'))

    head('stormwater_volume — m³ held (lot area as non-landscaped area: upper bound)')
    row('sv_lo', 'Van: 15 mm × lot; Edm: 375 m³/ha (100-year)', f=lambda c, t, k: city(c, f'IF({R("lot",c)}<=sv_small,"small site: 3,400–7,200 L tank",{R("lot",c)}*sv_mm/1000)',
                                                                                          f'IF({R("lot",c)}/10000>se_min_ha,IF({R("lot",c)}/10000<=se_max_ha,{R("lot",c)}/10000*se_100,"over 3 ha: site-specific"),"no control below 0.16 ha")'))
    row('sv_hi', 'Van: same (stage 1 only; full engineered volume not computed); Edm: 507 m³/ha (200-year)',
        f=lambda c, t, k: city(c, f'{R("sv_lo",c)}', f'IF(ISNUMBER({R("sv_lo",c)}),{R("lot",c)}/10000*se_200,{R("sv_lo",c)})'))
    row('sv_rel', 'Release rate L/s (Van stage 1 at 5 L/s/ha; Edm 35 L/s/ha)', f=lambda c, t, k: city(c, f'{R("lot",c)}/10000*sv_rel1', f'{R("lot",c)}/10000*se_rel'))

    head('embodied_it / embodied_building / refrigerant_co2e')
    row('emi', 'embodied_it t CO2e/yr', f=NC('not computable (needs size stage: server count)'))
    row('emi_lo', 'Per server, kg CO2e/yr, low = 4,177 × 10 % ÷ 4 (claim; share "?")', f=lambda c, t, k: '=emb_srv*emb_sh_lo/emb_life')
    row('emi_hi', 'Per server, kg CO2e/yr, high = 4,177 × 15 % ÷ 4', f=lambda c, t, k: '=emb_srv*emb_sh_hi/emb_life')
    row('emb', 'embodied_building t CO2e', f=NC('not computable (needs size stage: new floor area); 400 kg/m² (Vancouver benchmark) to 461 (CLF median, unverified)'))
    row('ref', 'refrigerant_co2e t CO2e/yr', f=NC('not computable (needs size stage: chiller count and charge 50–600 kg; refrigerant GWP not stated)'))

    head('Other parts that need the size stage or site plans')
    row('lf_flag', 'lf_flag', f=NC('not computable (needs dBC and octave spectra of the plant schedule)'))
    row('light', 'light_at_window lx', f=NC('not computable (needs lighting plan); benchmark 10 / 2 lx (CIE 150 via ILP, unverified)'))
    row('shadow', 'shadow m²·h', f=NC('not computable (needs size stage: massing)'))
    row('ground', 'ground', f=NC('not computable per lot: GWELLS export not in data/ (gwells check 2); Edmonton AWWID reports are by ATS section, no ATS grid in data/'))

    for col, w_ in (('A', 82),): s.column_dimensions[col].width = w_
    for j in range(len(LOTS)): s.column_dimensions[L(2 + j)].width = 17
    s.freeze_panes = 'B5'
    return wb, rows

# ===================================================================== 4. parts: rows read back, status, notes
QUAL = {  # part -> what the register says (no number)
    'visibility': 'Ledger per plant element: seen from public space / closed to view / known only from documents / not checked (fieldJB site fact; vanPermits for the documents column). Field observation, not lot data.',
    'watched_area': 'Visible cameras and signs; indicative watched area past the lot line; first-person account labelled as one person\'s. fieldJB notes support counts and description, not metres. Never a guide to coverage gaps.',
    'vibration_rank': 'Qualitative rank, slab on grade → long-span upper floor. VBBL 2019 / BCBC 2018 ASTC 47 (STC 55 by elevator shafts) are airborne only; isolator deflection rises with floor span (VA spec citing ASHRAE). No dB values.',
    'workload_control': 'Flexibility is offered only by whoever runs the workloads: single-owner or hyperscale operator can shift its own computing; in colocation the building operator controls power and cooling, not tenants\' computing. Crozier & Liska 2025: no single flexible share.',
    'computing_served': 'Up to four statements, no score: federal Protected B and above stored in Canada (gcDataResidency, archived — history); BC FOIPPA s. 33.1 assessment for public bodies (context); Alberta no rule found; Vancouver exchanges 158 participants across three (78 + 64 + 16, isocPulseIXP), Edmonton YEGIX 9; residency ≠ sovereignty.',
    'security_layers': 'One site visit (VAN3, tourVAN3, conf 40): 4 controlled doors passed (gate intercom, buzzed front door, vestibule-to-lobby door, single-person portal), then each tenant\'s enclosure (5 to the servers, guide\'s account). Cologix lists ISO 27001 and SOC 2 Type 2 (claims). Not lot data.',
}
PARTS = {  # part -> (low row, high row, unit, how, status)
    'heat_rejected': ('hr_lo', 'hr_hi', 'MWh/yr', 'facility electricity = IT × 0.6 × PUE 1.4–1.55 × 8,760 (heat ≈ electricity drawn)', 'computed'),
    'heat_temp': ('ht_a_lo', 'ht_l_hi', '°C', 'two runs: air-cooled 25–35, liquid-cooled 50–60 (liquid unverified; NREL measured return 35–40)', 'computed (register lookup; cooling type is size stage)'),
    'liquid_capture': ('lc_lo', 'lc_hi', 'share', 'cold plate 0.5–0.8 (no primary for the band); immersion 1.0; air-cooled not stated (?)', 'computed (register lookup, liquid run only)'),
    'cop': ('cop_l', 'cop_h', '—', 'register tool range 3.0–4.5 carried downstream; Lorenz check for Odense in extra', 'not computable (needs receiver_temp: set by receiver type, match stage); register range shown'),
    'heat_captured': ('hc_lo', 'hc_hi', 'MWh/yr', 'heat_rejected × cold-plate share 0.5–0.8 (liquid-cooled run); air-cooled not computable', 'computed (liquid-cooled run)'),
    'heat_delivered': ('hd_lo', 'hd_hi', 'MWh/yr available', 'E_reuse × COP ÷ (COP − 1), COP 4.5 (low) to 3.0 (high); receiver demand not applied', 'computed (COP from register range)'),
    'hp_electricity': ('hp_lo', 'hp_hi', 'MWh/yr', 'Q_out ÷ COP (4.5 low, 3.0 high)', 'computed (COP from register range)'),
    'erf': ('erf_lo', 'erf_hi', '—', 'E_reuse ÷ E_DC (= heat_rejected); collapses to the capture share by construction', 'computed'),
    'heat_in_water': ('hiw', 'hiw_r', 'kWh per m³·K', 'coefficient only: 4.18 × 1,000 ÷ 3,600 (low) vs register 1.16 (high); no lot volume or ΔT named', 'computed (coefficient check; no lot input)'),
    'free_cooling_hours': ('fc8', 'fc27', 'h/yr', 'CWEC hours ≤ 8 °C (full free cooling, air-cooled hall) to air-side at 27 °C inlet (≤ 25.5 °C); ≤ counts match the register, the formula text says <', 'computed'),
    'pue': ('pue_l', 'pue_h', '—', 'register range 1.4–1.55 (estimate + industry survey)', 'computed (register values)'),
    'spare_capacity': ('sp_2n_lo', 'sp_2n_hi', 'MW', 'only the conditional "if 2N, spare = peak" (our arithmetic); installed capacity is size stage', 'not computable (needs size stage: installed capacity and redundancy); conditional 2N value shown'),
    'backup_hours': ('bk_lo', 'bk_hi', 'h (fuel)', 'register ranges: fuel 12 h (Uptime minimum) to 48 h (operator claims); UPS 1–15 min (vendor claims) in extra', 'computed (register values only; battery kWh and fuel stored are size stage)'),
    'flexible_load': ('fl_lo', 'fl_hi', 'MW·h per event', 'it_load_actual × 10 % × 4 h (Call for Demand) to × 25 % × 3 h (Emerald preprint, claim)', 'computed'),
    'load_factor': ('lf_v', 'lf_v', '—', 'Fan et al. 2007 (one 2007 site)', 'computed (register value)'),
    'evaporation': ('ev_lo', 'ev_hi', 'L/yr', 'upper bound: heat_rejected ÷ 2,400 kJ/kg (all heat evaporated; dry cooling gives 0)', 'computed (upper bound)'),
    'blowdown': ('mu_lo', 'mu_hi', 'L/yr make-up', 'make-up = E + E ÷ (C − 1), C 6 (low) to 3 (high) from the femp example; blowdown in extra', 'computed (upper bound; C from femp example)'),
    'wue': ('wue_lo', 'wue_hi', 'L/kWh', 'make-up ÷ E_IT (evaporative upper bound)', 'computed (upper bound)'),
    'water_behind_electricity': ('wb_lo', 'wb_hi', 'L/yr', 'facility MWh × 1.25 L/kWh (thermo) to × 16.8 (hydro, Vancouver only); Alberta 1.25 both ends', 'computed'),
    'rainwater': ('rw_lo', 'rw_hi', 'm³/yr', 'existing roofs (footprint) × rainfall–precipitation (ECCC 1991–2020) × 0.8', 'computed (existing roofs)'),
    'grid_co2': ('gc_lo', 'gc_hi', 't CO2e/yr', 'facility MWh × NIR 2024 consumption intensity (BC 18.1, AB 346.1 g/kWh); Alberta displacement factor and heat-pump share in extra', 'computed'),
    'generator_co2': ('gco2_lo', 'gco2_hi', 't CO2/yr', 'test kWh (facility peak kW × 4–14.5 h) × 0.263 L/kWh × 2.681 kg/L', 'computed (lower bound on capacity: facility peak)'),
    'generator_nox_pm': ('nox_lo', 'nox_hi', 'kg NOx/yr', 'kW ÷ 0.7457 × hours × 0.024 lb/hp·h; PM in extra at 0.43 g/kWh', 'computed (lower bound on capacity: facility peak)'),
    'refrigerant_co2e': (None, None, 't CO2e/yr', 'charge × 3 %/yr × GWP', 'not computable (needs size stage: chiller count and charge; refrigerant GWP not stated)'),
    'embodied_it': ('emi_lo', 'emi_hi', 'kg CO2e/yr per server', 'per server only: 4,177 kg (claim) × 10–15 % ÷ 4 yr', 'not computable (needs size stage: server count); per-server value shown'),
    'embodied_building': (None, None, 't CO2e', 'floor area × 400–461 kg/m²', 'not computable (needs size stage: floor area)'),
    'noise_level': ('nz_lo', 'nz_hi', 'dBA at nearest residential parcel', 'Lw 98–101 dBA per unit (Trane, claim), one unit per MW (IT to facility peak), Lw + 10·log10 n − 20·log10 r − 8; r = lot edge to parcel edge', 'computed (screening; plant schedule is size stage)'),
    'lf_flag': (None, None, '—', 'dBC − dBA ≥ 20 dB and tone 20–250 Hz → +5 dBA (benchmark)', 'not computable (needs plant spectra)'),
    'light_at_window': (None, None, 'lx', 'vertical illuminance vs 10 / 2 lx (benchmark, unverified)', 'not computable (needs lighting plan)'),
    'shadow': (None, None, 'm²·h', 'fall-equinox shadows 10:00–16:00 (benchmark)', 'not computable (needs size stage: massing)'),
    'ground': (None, None, 'ft, USgpm, m²/d', 'nearest N wells: register gives Vancouver area stats only (333 wells, depth median 59 ft)', 'not computable (GWELLS export not in data/; Edmonton reports by ATS section)'),
    'visibility': (None, None, 'ledger', '', 'qualitative'),
    'watched_area': (None, None, 'm; count', '', 'qualitative'),
    'jobs': ('jb_lo', 'jb_hi', 'FTE', 'IT MW × 50–100 jobs ÷ 28 MW (Ark UP4 applicant claim); floor-area route needs size stage', 'computed (per-MW claim route)'),
    'treatment_energy': ('tr_lo', 'tr_hi', 'kWh/yr', 'MIN(greywater within 200–500 m, make-up) × 0.75–1.1 kWh/m³ (Arden 2021; register has no page/table)', 'computed (upper bound: recovery share is a register gap)'),
    'blowdown_recovery': ('br_lo', 'br_hi', 'L/yr recovered', 'blowdown × 0.53; make-up × 0.84 and 3,541 kWh per skid in extra', 'computed (single US test)'),
    'pv_roof': ('pvr_lo', 'pvr_hi', 'kWh/yr per m² of panel', 'per m² only: 0.16–0.20 kWp/m² × NRCan horizontal–latitude yield', 'not computable (needs size stage: roof layout); per-m² value shown'),
    'pv_facade': ('pvf_lo', 'pvf_hi', 'kWh/yr per m² of facade', 'per m² only: 0.16–0.20 kWp/m² × NRCan vertical yield', 'not computable (needs size stage: south facade area); per-m² value shown'),
    'genset_test': ('gtc_lo', 'gtc_hi', 't CO2/yr', 'test kWh × 706 g CO2 / 14.6 g NOx / 0.43 g PM per kWh (AP-42 upper bounds)', 'computed (lower bound on capacity: facility peak)'),
    'hvo_swap': (None, None, 'g CO2e/MJ', 'Δ = fuel energy × (CI diesel − CI RD)', 'not computable (fossil diesel reference CI not stated)'),
    'fuel_on_site': ('fu_lo', 'fu_hi', 'L', 'facility peak kW × 12–48 h × 0.263 L/kWh (claim); vs NFC 200,000 L per first-storey room', 'computed (lower bound on capacity: facility peak)'),
    'vibration_rank': (None, None, 'rank', '', 'qualitative'),
    'workload_control': (None, None, 'operator · tenants · mixed', '', 'qualitative'),
    'computing_served': (None, None, 'text', '', 'qualitative'),
    'heat_flux_site': ('hf_lo', 'hf_hi', 'W/m²', 'facility peak W ÷ lot area', 'computed'),
    'heat_plume_dT': (None, None, '°C', 'needs airflow and plume factor', 'not computable (needs plant airflow and air_screen_1h)'),
    'genset_emission_rate': ('ger_lo', 'ger_hi', 'g/s NOx', '14.6 g/kWh × facility peak kW ÷ 3,600, all sets at full load', 'computed (lower bound on capacity: facility peak)'),
    'air_screen_1h': (None, None, 'µg/m³', 'Gaussian screen', 'not computable (σy, σz and stack height not stated)'),
    'no2_at_receptor': (None, None, 'µg/m³', 'NOx × 1.0 / 0.9–0.2', 'not computable (needs air_screen_1h)'),
    'free_cooling_hours_2050s': ('f50_8_585', 'f50_27_245', 'h/yr', 'PCIC 2050s: hours ≤ 8 °C under SSP5-8.5 (low) to air-side at 27 °C inlet under SSP2-4.5 (high)', 'computed'),
    'condensate': ('cond', 'cond', 'h/yr (dew point > 15 °C)', 'CWEC hours with dew point > 15 °C; condensate yield not sourced', 'computed (hours only)'),
    'stormwater_volume': ('sv_lo', 'sv_hi', 'm³', 'Van 15 mm × lot (stage 1); Edm 375–507 m³/ha', 'computed (upper bound: whole lot non-landscaped)'),
    'security_layers': (None, None, 'count', '', 'qualitative'),
}
EXTRA = {  # part -> {name: [rows]}
    'heat_rejected': {'peak_MW': ['hr_pk_lo', 'hr_pk_hi'], 'it_energy_MWh': ['itmwh']},
    'heat_temp': {'air_C': ['ht_a_lo', 'ht_a_hi'], 'liquid_C_unverified': ['ht_l_lo', 'ht_l_hi'], 'liquid_measured_return_hi_C': ['ht_n']},
    'liquid_capture': {'immersion': ['lc_im'], 'air': ['lc_air']},
    'cop': {'lorenz_odense_eta050': ['cop_ode'], 'lorenz_odense_eta040': ['cop_ode40']},
    'heat_captured': {'immersion_MWh': ['hc_im'], 'it_heat_only_hi_MWh': ['hc_it']},
    'heat_in_water': {},
    'free_cooling_hours': {'air_18_21_24_27': ['fc18', 'fc21', 'fc24', 'fc27'], 'dry_cooler_30_40': ['fcd30', 'fcd40'], 'strict_lt_air18': ['fc18s']},
    'free_cooling_hours_2050s': {'le8_now_245_585': ['fc8', 'f50_8_245', 'f50_8_585'], 'air18_245_585': ['f50_18_245', 'f50_18_585'],
                                 'gt18_now_245_585': ['f50_hot_now', 'f50_hot_245', 'f50_hot_585'], 'ge30_245_585': ['f50_vh_245', 'f50_vh_585'],
                                 'peak_C_now_245_585': ['f50_max_now', 'f50_max_245', 'f50_max_585']},
    'backup_hours': {'ups_h': ['ups_lo', 'ups_hi']},
    'blowdown': {'blowdown_L': ['bd_lo', 'bd_hi'], 'evaporation_L': ['ev_lo', 'ev_hi'], 'femp_check_cut': ['bd_chk']},
    'wue': {'low_over_DCRE_cap': ['wue_cap']},
    'rainwater': {'roof_m2': ['rw_A']},
    'grid_co2': {'ab_displacement_t': ['gc_m_lo', 'gc_m_hi'], 'heat_pump_t': ['gc_hp_lo', 'gc_hp_hi']},
    'generator_co2': {'litres': ['g_l_lo', 'g_l_hi'], 'g_per_kWh_check': ['gco2_chk'], 'manor_farm_564h_t': ['gco2_man'], 'generator_kW_lower_bound': ['gk_lo', 'gk_hi']},
    'generator_nox_pm': {'pm_kg': ['pm_lo', 'pm_hi'], 'nox_g_per_kWh_check': ['nox_chk']},
    'genset_test': {'nox_kg': ['gtn_lo', 'gtn_hi'], 'pm_kg': ['gtp_lo', 'gtp_hi'], 'test_kWh': ['gt_e_lo', 'gt_e_hi'], 'test_hours_check': ['gt_chk']},
    'genset_emission_rate': {'check_2000kWe': ['ger_chk']},
    'fuel_on_site': {'first_storey_rooms': ['fu_rooms_lo', 'fu_rooms_hi'], 'upper_storey_rooms_hi': ['fu_up_hi'], 'm3_per_MW_24h_check': ['fu_chk']},
    'noise_level': {'units': ['nz_n_lo', 'nz_n_hi'], 'combined_Lw': ['nz_lw_lo', 'nz_lw_hi'], 'receptor_m': ['rec']},
    'heat_flux_site': {'high_over_allen119': ['hf_rat']},
    'jobs': {'m2_per_job_check': ['jb_chk', 'jb_chk2']},
    'treatment_energy': {'all_makeup_mixed_hi_kWh': ['tr_mix']},
    'blowdown_recovery': {'makeup_after_L': ['br_mu_lo', 'br_mu_hi'], 'kWh_per_skid': ['br_kwh'], 'implied_cuts_2.8_to_4.9': ['br_chk', 'br_chk2']},
    'stormwater_volume': {'release_L_per_s': ['sv_rel']},
    'spare_capacity': {}, 'embodied_it': {},
}

NOTES = [  # part, kind, note (register internal consistency / usability)
    ('heat_captured', 'basis', 'Capture shares (LBNL 48–75 %, 0.5–0.8 band) are shares of SERVER heat, but the formula multiplies heat_rejected (facility energy incl. PUE overhead). E_reuse is overstated by the PUE factor (1.4–1.55): compare hc_hi with the "IT heat only" cross-check row.'),
    ('erf', 'construction', 'Because heat_captured = heat_rejected × liquid_capture and E_DC = heat_rejected, ERF equals liquid_capture (0.5–0.8) for every lot; the quantity carries no lot information.'),
    ('heat_rejected', 'source gap', '"idle floor 36% of max" is credited to lbnl2024 (role verified), but the lbnl2024 register text has no 36 % figure (it gives a 50 % capacity utilisation); lbnlCapture mentions 35 % capture at idle — a different quantity.'),
    ('liquid_capture', 'source gap', '0.5–0.8 cold-plate band and immersion 1.0 have no primary (role says so); no air-cooled share → air run not computable.'),
    ('cop', 'input', 'Lot COP needs receiver_temp (not lot data). Lorenz formula reproduces the register Odense check (η 0.50 → 4.56). Downstream uses the register "tool range 3.0–4.5".'),
    ('pv_roof / pv_facade', 'contradiction', 'pv_roof coef P-16 quotes Vancouver 892–1,006 (vertical 718) and Edmonton 947–1,243 (vertical 990); the same coef, pv_facade (716 / 989), nrcanPVmunic and data/nrcan_pv_annual.json give 1,007 / 716 and 944 / 1,244 / 989. P-16 values match no source in the register.'),
    ('refrigerant_co2e', 'contradiction', 'sources = [ipcc2019v3ch7] but source_roles credits eccc (Climate Normals) with "all coefficients"; refrigerant and GWP not named; chiller count is size stage.'),
    ('generator_co2', 'source gap', '2.681 kg CO2/L is credited to eccc, whose register entry is the Climate Normals (no emission factor). Arithmetic is consistent: 0.263 L/kWh × 2.681 = 705 g/kWh vs AP-42 706. "Manor Farm 564 h/yr" has no source in the list and no per-unit / per-site basis.'),
    ('generator_nox_pm', 'source gap', 'NOx 0.024 lb/hp·h: eccc role relabelled "climate input only"; ap42_34 (the AP-42 source) is not in this part\'s list. Value equals 14.6 g/kWh (consistent with genset_test). No PM lb/hp·h factor; PM uses genset_test 0.43 g/kWh.'),
    ('generators (all)', 'size stage', 'Generator capacity, unit count and redundancy are size stage; facility peak kW is used as a lower bound (capacity ≥ load). genset_test 14.5 h reconstructs from auc_27938 text: 26 × 15 min + 4 × 1 h at zero load + 2 × 2 h full load.'),
    ('blowdown', 'unstated input', 'coef gives no value of C; femp text example 3→6 cycles used. femp "~20 % cut" checks (1 − 1.2/1.5 = 20 %). Evaporation is an upper bound (cooling type is size stage).'),
    ('blowdown_recovery', 'contradiction', 'Factors −53 % blowdown / −16 % make-up are not consistent with the stated cycles 2.8 → 4.9 at constant evaporation (implies −54 % blowdown and −19 % make-up); applying 0.84 and 0.47 together over-determines the water balance. femp_ct covers blowdown reduction, not recovery (register open item). Skid count not stated.'),
    ('wue', 'scale', 'At the evaporative upper bound WUE = 1.5 × PUE × C/(C−1) = 2.5–3.5 L/kWh for every lot, above the "DCRE cap 2 L/kWh" (DCRE not in the part\'s sources).'),
    ('water_behind_electricity', 'source gap', '"BC as a range" gives no numbers (1.25–16.8 used); no Alberta value; femp role "all coefficients" but the femp text (BMP #10) has neither 1.25 nor 16.8. Input names it_load_actual; facility electricity used here.'),
    ('grid_co2', 'basis', 'NIR 2024 is flagged "a" (preliminary). Input names it_load_actual; facility electricity used (heat ≈ electricity). Alberta displacement factor is an offset-policy number, shown separately.'),
    ('free_cooling_hours', 'wording', 'Formula says "T_outdoor < threshold"; the register\'s counts (7,109 / 7,614 at 18 °C; 3,284 / 5,522 at 8 °C; 7,675 / 7,846 dry cooler 30 °C) match ≤, not <. Both computed.'),
    ('noise_level', 'basis', '"one unit per MW" does not say MW of IT or of facility load: both used (low / high). r is lot edge to nearest residential parcel (input_check 8), so Lp is a closest-case screen; −8 is a hemispherical free-field simplification (iso9613 old: "our simplification").'),
    ('jobs', 'size stage', 'Formula is floor-area based (size stage); per-MW route uses the Ark UP4 applicant figure (claim), which is not in the part\'s source list (jlarc2024, tourVAN3).'),
    ('treatment_energy', 'source gap', 'arden2021 register entry states no figures (check 2: unreachable) — the 0.75–1.28 kWh/m³ values cannot be checked against the register.'),
    ('heat_flux_site', 'comparability', 'Allen 2011 119 W/m² is a peak grid cell (check 1), units assumed; comparison is indicative.'),
    ('air_screen_1h', 'unusable', 'σy, σz coefficients and stack height H not given ("table to be checked against the final ISC3 edition"); downstream no2_at_receptor and heat_plume_dT (neighbour) therefore not computable.'),
    ('hvo_swap', 'unusable', 'Fossil diesel reference CI "to retrieve" — formula cannot be evaluated.'),
    ('ground', 'contradiction', 'status says "Vancouver in hand (government export)", but gwells check 2 found no export files locally; Edmonton reports are indexed by ATS legal land description.'),
    ('condensate', 'unit', 'Named "condensate" but unit is hours; no litres-per-kW yield exists in the register (EPA office rule of thumb not for data halls).'),
    ('spare_capacity', 'definition', '"2N holds half" — 2N is not defined in the uptimeTiers text (role note).'),
    ('embodied_it', 'basis', '4,177 kg is full life cycle including US use phase (claim); 10–15 % manufacturing share marked "?".'),
    ('stormwater_volume', 'scope', 'Lot area used as non-landscaped area (upper bound, per register); Vancouver full engineered volume not computed; all lots between 0.16 and 3.0 ha.'),
    ('heat_delivered / hp_electricity', 'gap', 'hp_electricity coef "allocation rule stated; show two rules" — the two rules are not stated.'),
]

def recalc(path):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice: return False
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    shutil.move(os.path.join(tmp, os.path.basename(path)), path)
    return True

def main():
    assert set(PARTS) == set(Q), (set(Q) - set(PARTS), set(PARTS) - set(Q))
    wb, rows = build()
    q = wb.create_sheet('Qualitative')
    q.append(['part', 'register statement (no number)']); [setattr(x, 'font', Font(bold=True)) for x in q[1]]
    for p, t in QUAL.items(): q.append([p, t])
    n = wb.create_sheet('Register notes')
    n.append(['part', 'kind', 'note']); [setattr(x, 'font', Font(bold=True)) for x in n[1]]
    for p in NOTES: n.append(list(p))
    st = wb.create_sheet('Parts')
    st.append(['part', 'status', 'unit', 'how', 'register formula']); [setattr(x, 'font', Font(bold=True)) for x in st[1]]
    for p, (lo, hi, unit, how, status) in PARTS.items(): st.append([p, status, unit, how or QUAL.get(p, ''), Q[p]['formula']])
    for ws, widths in ((q, (22, 160)), (n, (24, 16, 160)), (st, (26, 50, 26, 90, 90))):
        for i, w_ in enumerate(widths): ws.column_dimensions[L(i + 1)].width = w_
        for row_ in ws.iter_rows(min_row=2):
            for cell in row_: cell.alignment = Alignment(wrap_text=True, vertical='top')
    out = os.path.join(HERE, 'emit_check.xlsx'); wb.save(out); print('wrote', out)
    ok = recalc(out); print('recalculated with LibreOffice' if ok else 'LibreOffice not found: values not cached')

    vb = load_workbook(out, data_only=True)['Emit']
    val = lambda key, j: vb.cell(rows[key], 2 + j).value
    def rd(x): return round(x, 4) if isinstance(x, float) else x
    res = {}; errors = []
    for j, (k, l, cty, p, it) in enumerate(LOTS):
        o = {'_lot': dict(label=l, city=cty, parcel_id=p, it_mw=it)}
        for part, (lo, hi, unit, how, status) in PARTS.items():
            if status == 'qualitative':
                o[part] = dict(status='qualitative', unit=unit, register_says=QUAL[part]); continue
            d = dict(status=status, unit=unit, how=how)
            if lo:
                a, b = rd(val(lo, j)), rd(val(hi, j))
                if isinstance(a, str) or isinstance(b, str):
                    txt = a if isinstance(a, str) else b
                    if txt.startswith(('#', 'Err')): errors.append((k, part, txt))
                    d['status'] = txt if txt.startswith('not computable') else f'not computable ({txt})'
                    d['low'] = d['high'] = None
                elif lo == hi: d['value'] = a
                else: d['low'], d['high'] = a, b
            ex = {}
            for nm, keys in EXTRA.get(part, {}).items():
                vs = [rd(val(x, j)) for x in keys]; ex[nm] = vs[0] if len(vs) == 1 else vs
                errors += [(k, part, v) for v in vs if isinstance(v, str) and v.startswith(('#', 'Err'))]
            if d['status'].startswith('not computable') and any(d.get(x) is not None for x in ('low', 'high', 'value')):
                # a not-computable part keeps the reference values it shows, but not as the part's answer
                ex['reference_shown'] = d.pop('value') if 'value' in d else [d.get('low'), d.get('high')]
                d['low'] = d['high'] = None
            if ex: d['extra'] = ex
            o[part] = d
        res[k] = o
    if errors: print('FORMULA ERRORS:', errors[:20])
    def bucket(s): return 'qualitative' if s == 'qualitative' else 'not computable' if s.startswith('not computable') else 'computed'
    counts = {}
    for part in PARTS:
        b = sorted({bucket(res[k][part]['status']) for k in KEYS}); counts[part] = '/'.join(b)
    summary = collections_count(counts)
    meta = dict(built=str(datetime.date.today()), script='checks/emit_check/build_emit_check.py',
                inputs='register (stage characterise, 53 parts) + new/old_sources; data/ (CWEC, PCIC 2050s EPW, NIR 2026 Annex 7 extract, ECCC normals extract, nrcan_pv_annual.json); checks/input_check/input_check.json; checks/place_check/place_check.json (greywater); IT load per lot as given',
                not_read='tool/*.js, tool/index.html, checks/golden, checks/place_check/screen, checks/agents, *_tabs.json',
                assumptions=['load: IT × 0.6 (load_factor) × PUE 1.4–1.55; low pairs PUE 1.4 with low coefficients, high pairs 1.55 with high',
                             'generator capacity (size stage) taken as facility peak kW: a lower bound',
                             'cooling type unknown: evaporation / make-up / WUE are evaporative upper bounds; heat capture uses the liquid (cold-plate) run',
                             'COP for downstream parts: register tool range 3.0–4.5 (lot COP needs receiver_temp)'],
                part_status=counts, status_summary=summary, formula_errors=errors,
                register_notes=[dict(part=a, kind=b, note=c) for a, b, c in NOTES])
    json.dump(dict(meta=meta, lots=res), open(os.path.join(HERE, 'emit_check.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('wrote emit_check.json'); print(summary)

def collections_count(counts):
    out = {}
    for v in counts.values(): out[v] = out.get(v, 0) + 1
    return out

if __name__ == '__main__':
    main()
