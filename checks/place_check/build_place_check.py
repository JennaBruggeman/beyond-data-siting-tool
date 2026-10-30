"""Builds checks/place_check/place_check.xlsx and place_check.json: an INDEPENDENT hand calculation of the
"What the site offers" (place) stage of Beyond Data, for eight lots.

Written from the register only (register/formula_register.json quantities + new_sources texts) and from raw
data under data/ (plus tool/*.json data files where named). The tool's JavaScript, the golden results and
checks/place_check/screen were NOT read. Lot-level inputs come from checks/input_check/input_check.json.

Every computed cell in the workbook is a live Excel formula; constants sit on the Constants sheet with the
register part and source id they come from. The script then recalculates the workbook with LibreOffice
(if present) and reads the cached values back to write place_check.json, so the JSON is the workbook's answer.

Run:  python3 build_place_check.py          (about 2–4 minutes: Vancouver footprints and Edmonton rooflines are large)
"""
import csv, json, math, os, glob, collections, subprocess, shutil, tempfile, datetime
import pyproj, fiona, shapefile
from shapely.geometry import shape, Polygon, MultiPolygon, Point, box
from shapely.ops import transform as shp_transform, unary_union
from shapely.strtree import STRtree
from shapely import make_valid
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
D = os.path.join(ROOT, 'data'); TOOL = os.path.join(ROOT, 'tool')
csv.field_size_limit(10 ** 9)
REG = json.load(open(os.path.join(ROOT, 'register', 'formula_register.json'), encoding='utf-8'))
Q = {x['id']: x for x in REG['quantities']}
IC = json.load(open(os.path.join(ROOT, 'checks', 'input_check', 'input_check.json'), encoding='utf-8'))['lots']

LOTS = [  # key, label, city, parcel id, IT load MW (given by the caller)
    ('V09', '275 1st Ave E', 'van', '63819363', 2.5),
    ('E05', '8803 Stadium Rd NW', 'edm', '51053', 10.2),
    ('CLARK', '1425 Clark Dr', 'van', '23060741', 5.2),
    ('E123', '12603 123 Street NW', 'edm', '883345', 22.2),
    ('V339', '339 1st Ave E', 'van', '63819485', 3.8),
    ('E06A', '74 Airport Rd NW', 'edm', '176767', 10.2),
    ('E06B', '64 Airport Rd NW', 'edm', '146402', 10.2),
    ('HAST', '555 W Hastings', 'van', '59013407', 1.1),
]
KEYS = [l[0] for l in LOTS]
EPSG = {'van': 26910, 'edm': 3776}
FWD = {c: pyproj.Transformer.from_crs(4326, e, always_xy=True).transform for c, e in EPSG.items()}
def proj(g, c): return shp_transform(FWD[c], g)
def valid(g): return g if g.is_valid else make_valid(g)
def num(v):
    try: return float(str(v).replace(',', '')) if v not in (None, '') else None
    except ValueError: return None
RADII = (200, 500)   # people_affected radii (register: "summed within 200 m, 500 m and the sound distance"); rooftop_solar names no radius

# ===================================================================== 1. geometry-derived inputs
def derive():
    G = {k: {} for k in KEYS}
    # ---- lots (same rule as input_check: Vancouver parcels by tax_coord; Edmonton Altalis title lot holding the assessment point)
    van_ids = {p: k for k, l, c, p, it in LOTS if c == 'van'}
    vp = json.load(open(os.path.join(D, 'vancouver_parcels_2026.geojson'), encoding='utf-8'))
    parts = collections.defaultdict(list); by_tc = collections.defaultdict(list)
    for f in vp['features']:
        tc = (f['properties'].get('tax_coord') or '').lstrip('0')
        by_tc[tc].append(f['geometry'])
    LOT = {}
    for tc, k in van_ids.items():
        LOT[k] = proj(unary_union([valid(shape(g)) for g in by_tc[tc]]), 'van')
    edm_ids = {p: k for k, l, c, p, it in LOTS if c == 'edm'}
    erec = collections.defaultdict(list)
    for r in csv.DictReader(open(os.path.join(D, 'edmonton_properties_2026.csv'), encoding='utf-8')):
        if r['parcel_point_id'] in edm_ids: erec[r['parcel_point_id']].append(r)
    raw = json.load(open(os.path.join(D, 'raw', 'edm', 'lots', 'edmonton_lots_epsb.json')))
    def dec(ring):
        x = y = 0; out = []
        for i in range(0, len(ring), 2): x += ring[i]; y += ring[i + 1]; out.append((x / 1e6, y / 1e6))
        return out
    lots_ll = []
    for pid, oid, rings in raw['lots']:
        rs = [dec(r) for r in rings]
        if len(rs[0]) < 4: continue
        lots_ll.append(valid(Polygon(rs[0], [r for r in rs[1:] if len(r) >= 4])))
    LT = STRtree(lots_ll)
    for ppid, k in edm_ids.items():
        r = erec[ppid][0]; pt = Point(float(r['lon']), float(r['lat']))
        hits = [lots_ll[i] for i in LT.query(pt, predicate='intersects')]
        LOT[k] = proj(unary_union(hits), 'edm')
        yb = sorted({x['year_built'] for x in erec[ppid] if x['year_built']})
        G[k]['year_built'] = ', '.join(yb) if yb else 'none recorded'
    # Vancouver year built (2026 tax report)
    for r in csv.DictReader(open(os.path.join(D, 'raw', 'van', 'property-tax-report.csv'), encoding='utf-8-sig'), delimiter=';'):
        tc = r['LAND_COORDINATE'].lstrip('0')
        if tc in van_ids:
            k = van_ids[tc]; s = set(filter(None, G[k].get('_yb', '').split(', '))) | ({r['YEAR_BUILT']} if r['YEAR_BUILT'] else set())
            G[k]['_yb'] = ', '.join(sorted(s))
    for k in KEYS:
        if '_yb' in G[k]: G[k]['year_built'] = G[k].pop('_yb') or 'none recorded'
        elif 'year_built' not in G[k]: G[k]['year_built'] = 'none recorded'
    city = {k: c for k, l, c, p, it in LOTS}

    # ---- neighbour roofs (footprints not on the lot: < 50 % inside, as in input_check), clipped to each radius
    print('roofs: Vancouver 2009 footprints + LiDAR 2022 changes')
    LF, LN = {}, []
    for fn in glob.glob(os.path.join(D, 'lidar2022', '*.json')):
        r = json.load(open(fn)); LF.update({int(a): v for a, v in r['fp'].items()}); LN += r['new']
    vroofs = []
    for f in json.load(open(os.path.join(D, 'vancouver_footprints_2009.geojson'), encoding='utf-8'))['features']:
        lv = LF.get(f['properties']['id'])
        if lv is not None and lv[2] < 0.1: continue        # demolished by 2022 (LiDAR share still built < 0.1)
        vroofs.append(proj(valid(shape(f['geometry'])), 'van'))
    for p90, hmax, area, ring in LN:
        vroofs.append(proj(valid(Polygon(ring)), 'van'))
    print('roofs: Edmonton rooflines 2019 (large)')
    eroofs = [proj(valid(shape(f['geometry'])), 'edm')
              for f in json.load(open(os.path.join(D, 'edmonton_rooflines_est.geojson'), encoding='utf-8'))['features']]
    RT = {'van': (STRtree(vroofs), vroofs), 'edm': (STRtree(eroofs), eroofs)}
    for k in KEYS:
        tree, arr = RT[city[k]]; lot = LOT[k]
        for R in RADII:
            buf = lot.buffer(R); tot = 0.0
            for i in tree.query(buf, predicate='intersects'):
                g = arr[i]
                if g.area and g.intersection(lot).area / g.area >= 0.5: continue   # the lot's own buildings
                tot += g.intersection(buf).area
            G[k][f'roof_{R}'] = round(tot, 0)

    # ---- residents (2021 census DA population × share of DA area within the radius)
    print('census')
    pop = {}
    for r in csv.DictReader(open(os.path.join(D, 'raw', 'census', 'census2021_DA_selected_Vancouver_Edmonton_regions.csv'), encoding='utf-8')):
        pop[r['dauid']] = num(r['pop2021']) or 0
    das = {'van': [], 'edm': []}
    with fiona.open('zip://' + os.path.join(D, 'raw', 'census', 'lda_000b21a_e.zip') + '!lda_000b21a_e.shp') as src:
        to_ll = pyproj.Transformer.from_crs(src.crs_wkt, 4326, always_xy=True).transform
        for f in src:
            da = f['properties']['DAUID']
            c = 'van' if da.startswith('5915') else 'edm' if da.startswith('4811') else None
            if c is None or da not in pop: continue
            g = proj(shp_transform(to_ll, valid(shape(f['geometry']))), c)
            das[c].append((g, pop[da]))
    for k in KEYS:
        lot = LOT[k]; arr = das[city[k]]
        for R in RADII + (3000,):
            buf = lot.buffer(R)
            G[k][f'res_{R}'] = round(sum(p * g.intersection(buf).area / g.area for g, p in arr if g.area and g.intersects(buf)), 0)

    # ---- nearest transformer station (NRCan CanVec 50K Res_MGT, raw shapefiles)
    for c, P in (('van', 'BC'), ('edm', 'AB')):
        st = []
        for kind in ('0', '2'):
            r = shapefile.Reader(os.path.join(D, 'raw', 'canvec', P, 'transformer_station_' + kind))
            for s in r.shapes():
                g = Point(s.points[0]) if kind == '0' else valid(Polygon(s.points))
                st.append(proj(g, c))
        for k in KEYS:
            if city[k] == c: G[k]['d_substation'] = round(min(LOT[k].distance(g) for g in st), 0)

    # ---- distance to water: Vancouver ocean / Fraser River (CanVec extract), Edmonton North Saskatchewan River (EPSB polygon)
    wv = json.load(open(os.path.join(D, 'raw', 'hydro', 'canvec_van_water_extract.json')))
    sea = [proj(valid(Polygon(w['r'][0], w['r'][1:])), 'van') for w in wv if w['t'] == 'Ocean']
    fraser = [proj(valid(Polygon(w['r'][0], w['r'][1:])), 'van') for w in wv if w['n'] == 'Fraser River']
    nsr = proj(valid(shape(json.load(open(os.path.join(D, 'raw', 'hydro', 'nsr_epsb.geojson')))['features'][0]['geometry'])), 'edm')
    for k in KEYS:
        if city[k] == 'van':
            G[k]['d_sea'] = round(min(LOT[k].distance(g) for g in sea), 0)
            G[k]['d_river'] = round(min(LOT[k].distance(g) for g in fraser), 0) if fraser else None
        else:
            G[k]['d_sea'] = None
            G[k]['d_river'] = round(LOT[k].distance(nsr), 0)

    # ---- VANIX connection points named in register source vanixFacilities (operator claim), located by civic address in the tax report
    want = {('1050', 'PENDER ST W'): 'Cologix VAN2, 1050 W Pender', ('2828', 'NATAL ST'): 'Cologix VAN3, 2828 Natal St',
            ('175', 'CORDOVA ST W'): 'Cologix VAN4, 175 W Cordova', ('555', 'HASTINGS ST W'): 'Harbour Centre MMR, 555 W Hastings'}
    ixp = {}
    for r in csv.DictReader(open(os.path.join(D, 'raw', 'van', 'property-tax-report.csv'), encoding='utf-8-sig'), delimiter=';'):
        kk = (r['TO_CIVIC_NUMBER'], r['STREET_NAME'])
        if kk in want: ixp[want[kk]] = r['LAND_COORDINATE'].lstrip('0')
    ixg = {n: proj(unary_union([valid(shape(g)) for g in by_tc[tc]]), 'van') for n, tc in ixp.items() if by_tc.get(tc)}
    for k in KEYS:
        if city[k] != 'van': G[k]['d_ixp'] = None; G[k]['ixp_name'] = ''; continue
        n, g = min(ixg.items(), key=lambda t: LOT[k].distance(t[1]))
        G[k]['d_ixp'] = round(LOT[k].distance(g), 0); G[k]['ixp_name'] = n
    G['_ixp_found'] = sorted(ixg)

    # ---- seismic: register points (nrcSeismic2020) — which lot holds each, and how far
    pts = {'V09': ('van', -123.097, 49.269, 1.08), 'E05': ('edm', -113.477, 53.555, 0.152), 'E06': ('edm', -113.522, 53.569, 0.154)}
    G['_seis'] = {}
    for n, (c, lo, la, v) in pts.items():
        p = proj(Point(lo, la), c)
        G['_seis'][n] = {k: round(LOT[k].distance(p), 0) for k in KEYS if city[k] == c}
    return G

# ===================================================================== 2. workbook
CONST = [  # name, value, unit, register part, source id(s), trust
    ('pue_lo', 1.4, '—', 'pue', 'lbnl2024 (US average 2023)', 80),
    ('pue_hi', 1.55, '—', 'pue', 'register "use 1.4–1.55" (uptime2025 survey 1.54)', 80),
    ('lf', 0.6, '—', 'load_factor / it_load_actual', 'fan2007 (≈0.6; tool uses IT × 8,760 × 0.6)', 80),
    ('hours', 8760, 'h/yr', 'load_factor', 'fan2007', 100),
    ('cp', 4.18, 'kJ/kg·K', 'sewer_sink / ates formula', 'heatPhys', 100),
    ('sw_dT_lo', 5, 'K', 'surface_water_sink', '"flow 24–49 L/s per MW at a 10–5 °C rise (derived)"', 60),
    ('sw_dT_hi', 10, 'K', 'surface_water_sink', 'same', 60),
    ('bc_rent_cool', 1.30, '$/1,000 m³', 'surface_water_sink', 'bcWSRfees (purpose "cooling")', 90),
    ('bc_rent_hx', 2.25, '$/1,000 m³', 'surface_water_sink', 'bcWSRfees (purpose "heat exchange"; the coef uses this one)', 90),
    ('bc_rent_min', 200, '$/yr', 'surface_water_sink', 'bcWSRfees', 90),
    ('bc_app_fee', 1000, '$', 'surface_water_sink', 'bcWSRfees', 90),
    ('pv_usable', 0.81, '—', 'rooftop_solar', 'nrcanPV2024', 70),
    ('pv_cover', 0.70, '—', 'rooftop_solar', 'nrcanPV2024', 70),
    ('kwp_lo', 0.16, 'kWp/m²', 'pv_roof (rooftop_solar names kWp/m² but gives no value)', 'NREL via pv_roof', 70),
    ('kwp_hi', 0.20, 'kWp/m²', 'pv_roof', 'NREL via pv_roof', 70),
    ('yld_van_lo', 892, 'kWh/kWp', 'pv_roof', 'NRCan municipal, horizontal (data/nrcan_pv_annual.json agrees)', 80),
    ('yld_van_hi', 1007, 'kWh/kWp', 'pv_roof', 'NRCan municipal, latitude tilt', 80),
    ('yld_edm_lo', 944, 'kWh/kWp', 'pv_roof', 'NRCan municipal, horizontal', 80),
    ('yld_edm_hi', 1244, 'kWh/kWp', 'pv_roof', 'NRCan municipal, latitude tilt', 80),
    ('pv_reg_van', 128, 'kWh/m² footprint/yr', 'rooftop_solar coef', 'nrcanPV2024 ("our derivation")', 60),
    ('pv_reg_edm', 159, 'kWh/m² footprint/yr', 'rooftop_solar coef', 'nrcanPV2024 ("our derivation")', 60),
    ('rs2289_kw', 100, 'kW per phase', 'self_gen_rules', 'bchTariff2026 / bchRS2289', 90),
    ('rs2289_price', 0.10, '$/kWh', 'self_gen_rules', 'bchTariff2026', 90),
    ('rs2290_mw', 2, 'MW net injection', 'self_gen_rules', 'bchSelfGen2026', 90),
    ('ab_micro_mw', 5, 'MW', 'self_gen_rules', 'abMicrogen / aucMicro2026', 90),
    ('ab_retail_kw', 150, 'kW', 'self_gen_rules', 'abMicrogen', 90),
    ('lgs_kw', 13.83, '$/kW billing demand (period not stated)', 'electricity_price', 'bchTariff2026 RS 1600–1611', 90),
    ('lgs_kwh', 0.0679, '$/kWh', 'electricity_price', 'bchTariff2026', 90),
    ('lgs_day', 0.2999, '$/day', 'electricity_price', 'bchTariff2026', 90),
    ('t1830_kva', 12.178, '$/kVA (period not stated)', 'electricity_price', 'bchTariff2026 RS 1830', 90),
    ('t1830_kwh', 0.04914, '$/kWh', 'electricity_price', 'bchTariff2026 RS 1830', 90),
    ('pool_lo', 43.68, '$/MWh', 'electricity_price', 'aesoStats2025 — NOTE: the source text covers operating reserves only; pool price not in it', 50),
    ('pool_hi', 62.78, '$/MWh', 'electricity_price', 'aesoStats2025 (2024) — same note', 50),
    ('sas_kwh', 0.0025, '$/kWh', 'electricity_price', 'epcorDASSAS SAS-TOUP', 90),
    ('sas_or', 0.0799, '× pool', 'electricity_price', 'epcorDASSAS SAS-TOUP operating reserve', 90),
    ('das_day', 67.45678, '$/site/day', 'electricity_price', 'epcorDASSAS DAS-TOUP', 90),
    ('das_kwday', 0.13822, '$/kW/day', 'electricity_price', 'epcorDASSAS DAS-TOUP', 90),
    ('sas_cap', 0.13497, '$/kW/day', 'electricity_price', 'epcorDASSAS SAS-TOUP capacity', 90),
    ('sas_dem', 0.28126, '$/kW/day', 'electricity_price', 'epcorDASSAS SAS-TOUP demand', 90),
    ('van_w_lo', 4.091, '$/unit', 'water_supply', 'vanWater2026 (16 Oct–30 Apr; proposed rates)', 90),
    ('van_w_hi', 5.128, '$/unit', 'water_supply', 'vanWater2026 (1 May–15 Oct)', 90),
    ('van_sewer', 7.960, '$/unit', 'water_supply', 'vanWater2026 / vanRates2026 (marked ?)', 80),
    ('van_unit', 2.8316, 'm³/unit', 'water_supply', 'vanRates2026 (2,831.6 L)', 90),
    ('edm_w_lo', 0.9472, '$/m³', 'water_supply', 'epcorBylaw19626 (2022–23, lowest tier)', 70),
    ('edm_w_hi', 1.6122, '$/m³', 'water_supply', 'epcorBylaw19626 (highest tier)', 70),
    ('edm_ww_lo', 0.4636, '$/m³', 'water_supply', 'epcorWW2024app (proposed 2025)', 60),
    ('edm_ww_hi', 1.1486, '$/m³', 'water_supply', 'epcorWW2024app (proposed 2025)', 60),
    ('edm_san', 1.1212, '$/m³', 'water_supply', 'epcorWW2024app (proposed 2025)', 60),
    ('epcor_rec_ML', 4226, 'ML/yr (2024, all industry)', 'water_supply', 'epcorReclaimed', 80),
    ('h_fg', 2400, 'kJ/kg', 'evaporation (input to reclaimed_makeup)', 'register evaporation formula', 100),
    ('C_lo', 2, 'cycles', 'reclaimed_makeup', 'femp_ct (2–4 typical, up to 6)', 60),
    ('C_hi', 6, 'cycles', 'reclaimed_makeup', 'femp_ct', 60),
    ('lbnl_w_lo', 0.32, 'L/kWh (denominator to confirm)', 'reclaimed_makeup', 'lbnl2024 — role relabelled "not supported (?)"', 40),
    ('lbnl_w_hi', 0.36, 'L/kWh', 'reclaimed_makeup', 'same', 40),
    ('s_lo', 0.74, 'share', 'reclaimed_makeup', 'msft_water2026 (operator claim)', 50),
    ('s_hi', 0.99, 'share', 'reclaimed_makeup', 'msft_water2026 (operator claim)', 50),
    ('gw_lpd', 98, 'L/person/day', 'greywater_neighbours', 'nasem2016 (US data)', 45),
    ('q_ith_lo', 12, 'W/m', 'borefield', 'zurmuhl2018 (base design, 1,248 m)', 40),
    ('q_ith_hi', 19, 'W/m', 'borefield', 'zurmuhl2018 (economic optimum, ~760 m)', 40),
    ('k_ith', 3.27, 'W/m·K', 'borefield', 'zurmuhl2018 (Ithaca)', 40),
    ('k_edm', 1.9, 'W/m·K', 'borefield', 'ags_edmonton', 60),
    ('bh_len', 150, 'm/borehole', 'borefield', 'register formula', 40),
    ('bh_area', 36, 'm²/borehole', 'borefield', 'register formula', 40),
    ('ates_dT', 5, 'K', 'ates', 'register coef ("assumption")', 40),
    ('ates_rec_lo', 0.4, '—', 'ates', 'barth2026', 45),
    ('ates_rec_hi', 0.85, '—', 'ates', 'barth2026', 45),
    ('ates_eur_kw', 300, '€/kW (at 2 MW)', 'ates', 'barth2026', 45),
    ('geo_lo', 2.1, 'MWth per well pair', 'deep_geo', 'majorowicz2021', 55),
    ('geo_hi', 5.6, 'MWth per well pair', 'deep_geo', 'majorowicz2021', 55),
    ('geo_m_lo', 30, 'kg/s', 'deep_geo', 'majorowicz2021', 55),
    ('geo_m_hi', 80, 'kg/s', 'deep_geo', 'majorowicz2021', 55),
    ('geo_T', 90, '°C at 2.5 km', 'deep_geo', 'majorowicz2021 / vanDeepGeo', 55),
    ('geo_cost_lo', 4.7, 'M$ per well', 'deep_geo', 'majorowicz2021', 55),
    ('geo_cost_hi', 8, 'M$ per well', 'deep_geo', 'majorowicz2021', 55),
    ('fcl', 4.6, 'm (GVRD datum)', 'flood_exposure', 'van_f014 / vanFCL2014', 70),
    ('bch_2f_12', 6500, 'kVA', 'service_class', 'bch_primary (12.5 kV areas)', 75),
    ('bch_2f_25', 13000, 'kVA', 'service_class', 'bch_primary (25 kV areas)', 75),
    ('epcor_prim', 2500, 'kVA', 'service_class', 'epcor_ccg', 75),
    ('ab_reg117', 75, 'MW max demand from transmission', 'self_gen_rules', 'altaReg117 s. 2(2)', 90),
    ('fibre_ms_km', 0.005, 'ms per km one way', 'interconnection', 'vanixPoints note: no primary named', 40),
    ('sa_v09', 1.08, 'g', 'seismic_sa02', 'nrcSeismic2020 (49.269, −123.097)', 70),
    ('sa_e05', 0.152, 'g', 'seismic_sa02', 'nrcSeismic2020 (53.555, −113.477)', 70),
    ('sa_e06', 0.154, 'g', 'seismic_sa02', 'nrcSeismic2020 (53.569, −113.522)', 70),
]

def build(G):
    wb = Workbook()
    c = wb.active; c.title = 'Constants'
    c.append(['name', 'value', 'unit', 'register part', 'source (register new_sources id) and note', 'trust'])
    for row in CONST: c.append(list(row))
    for i, row in enumerate(CONST, start=2):
        wb.defined_names[row[0]] = DefinedName(row[0], attr_text=f"Constants!$B${i}")
    for col, w in zip('ABCDEF', (14, 10, 30, 34, 90, 7)): c.column_dimensions[col].width = w
    for cell in c[1]: cell.font = Font(bold=True)

    s = wb.create_sheet('Place')
    s['A1'] = 'What the site offers — independent hand calculation (live formulas)'; s['A1'].font = Font(bold=True, size=13)
    s['A2'] = 'Blue = input read from data or given; black = formula; grey italic = not derivable / text. Ranges: _lo uses the low coefficients, _hi the high.'
    s.cell(3, 1, 'lot').font = Font(bold=True)
    for j, (k, l, cty, p, it) in enumerate(LOTS):
        s.cell(3, 2 + j, k).font = Font(bold=True); s.cell(4, 2 + j, l)
    s.cell(4, 1, 'address')
    rows = {}; HOW = {}
    blue = Font(color='1F4E9E'); grey = Font(color='808080', italic=True)
    def row(key, label, vals=None, f=None, inp=False, how=''):
        r = s.max_row + 1; rows[key] = r; s.cell(r, 1, label); HOW[key] = how
        for j, (k, l, cty, p, it) in enumerate(LOTS):
            col = L(2 + j)
            if inp:
                v = vals[j]
                cell = s.cell(r, 2 + j, v); cell.font = grey if isinstance(v, str) and v.startswith(('not ', 'n/a')) else blue
            else:
                e = f(col, cty, k)
                if e is None: s.cell(r, 2 + j, 'n/a').font = grey
                elif isinstance(e, str) and e.startswith('='): s.cell(r, 2 + j, e)
                else: s.cell(r, 2 + j, e).font = grey
        return r
    R = lambda key, col: f'{col}{rows[key]}'
    def head(t): r = s.max_row + 2; s.cell(r, 1, t).font = Font(bold=True)
    V = lambda k: [x for x in LOTS if x[0] == k][0]
    ic = lambda k, item: IC[k][item]['value']

    head('INPUTS')
    row('van', 'Vancouver (1) or Edmonton (0)', [1 if x[2] == 'van' else 0 for x in LOTS], inp=True)
    row('it', 'IT load, MW of servers (given)', [x[4] for x in LOTS], inp=True)
    row('lot', 'Lot area m² (input_check 1)', [ic(k, '1_lot_area_m2') for k in KEYS], inp=True)
    row('fp', 'Buildings on lot, footprint m² (input_check 3)', [ic(k, '3_buildings')['total_footprint_m2'] for k in KEYS], inp=True)
    row('gross', 'Assessed gross floor area m² (Edmonton record; Vancouver not in open data)',
        [IC[k]['9_coverage_and_plot_ratio']['value'].get('gross_area_m2', 'not derivable (no floor area in Vancouver open data; BCA licensed)') for k in KEYS], inp=True)
    row('yb', 'Year built (tax report / assessment record)', [G[k]['year_built'] for k in KEYS], inp=True)
    row('zone', 'Zone (input_check 2)', [ic(k, '2_zone') for k in KEYS], inp=True)
    def fshare(k):
        h = IC[k]['4_floodplain']['how']
        if 'Designated Floodplain (FCL 4.6m)' not in h: return 'not derivable (no Edmonton floodplain layer in data/)'
        return float(h.split("'Designated Floodplain (FCL 4.6m)': ")[1].split('}')[0])
    row('flood', 'Share of lot in designated floodplain (FCL 4.6 m layer)', [fshare(k) for k in KEYS], inp=True)
    for R_ in RADII:
        row(f'roof{R_}', f'Neighbour roof footprint within {R_} m of the lot, m² (own buildings excluded)', [G[k][f'roof_{R_}'] for k in KEYS], inp=True,
            how='Vancouver 2009 footprints less those LiDAR 2022 shows gone, plus LiDAR 2022 new outlines; Edmonton 2019 rooflines; clipped to the buffer (EPSG:26910 / 3776).')
    for R_ in RADII + (3000,):
        row(f'res{R_}', f'Residents within {R_} m (2021 census DA pop × area share)', [G[k][f'res_{R_}'] for k in KEYS], inp=True,
            how='Census 2021 DA population (data/raw/census) × share of each DA polygon inside the lot buffer.')
    row('d_sub', 'Distance to nearest transformer station, m (CanVec 50K, 1971–2019 maps)', [G[k]['d_substation'] for k in KEYS], inp=True)
    row('d_sea', 'Distance to sea (CanVec Ocean polygons), m', [G[k]['d_sea'] if G[k]['d_sea'] is not None else 'n/a' for k in KEYS], inp=True)
    row('d_river', 'Distance to Fraser River (Van) / North Saskatchewan (Edm), m', [G[k]['d_river'] if G[k]['d_river'] is not None else 'not derivable' for k in KEYS], inp=True)
    row('d_ixp', 'Distance to nearest located VANIX point, m (operator claim)', [G[k]['d_ixp'] if G[k]['d_ixp'] is not None else 'not derivable (YEGIX: no address in register)' for k in KEYS], inp=True)
    row('ixp_name', 'Nearest VANIX point', [G[k]['ixp_name'] or 'n/a' for k in KEYS], inp=True)
    def sa(k):
        for n, nm in (('V09', 'sa_v09'), ('E05', 'sa_e05'), ('E06', 'sa_e06')):
            d = G['_seis'][n].get(k)
            if d is not None and d == 0: return '=' + nm
        return 'not derivable (no register point on this lot)'
    row('sa', 'Sa(0.2) register value whose point lies on this lot, g', [sa(k) for k in KEYS], inp=True)

    head('LOAD (from IT, PUE, load factor)')
    row('pk_lo', 'Facility peak MW, low = IT × PUE 1.4', f=lambda c, t, k: f'={R("it",c)}*pue_lo')
    row('pk_hi', 'Facility peak MW, high = IT × PUE 1.55', f=lambda c, t, k: f'={R("it",c)}*pue_hi')
    row('av_lo', 'Facility average MW, low = IT × 0.6 × 1.4', f=lambda c, t, k: f'={R("it",c)}*lf*pue_lo')
    row('av_hi', 'Facility average MW, high = IT × 0.6 × 1.55', f=lambda c, t, k: f'={R("it",c)}*lf*pue_hi')
    row('mwh_lo', 'Electricity = heat rejected, MWh/yr, low', f=lambda c, t, k: f'={R("av_lo",c)}*hours')
    row('mwh_hi', 'Electricity = heat rejected, MWh/yr, high', f=lambda c, t, k: f'={R("av_hi",c)}*hours')
    row('itmwh', 'IT energy MWh/yr = IT × 0.6 × 8,760', f=lambda c, t, k: f'={R("it",c)}*lf*hours')

    head('sewer_sink — MW at the pipe')
    row('sewer', 'Q = flow × 4.18 × ΔT', f=lambda c, t, k: 'not derivable: no sewer flow at any parcel (input_check 5); ΔT not in register')

    head('surface_water_sink — flow and BC rental (capacity itself unmodelled per register)')
    row('sw_perMW_lo', 'Flow per MW, L/s, at ΔT 10 K = 1000 ÷ (4.18 × 10)', f=lambda c, t, k: '=1000/(cp*sw_dT_hi)')
    row('sw_perMW_hi', 'Flow per MW, L/s, at ΔT 5 K (register says 49)', f=lambda c, t, k: '=1000/(cp*sw_dT_lo)')
    row('sw_lo', 'Flow for average heat, L/s, low = av_lo × perMW_lo', f=lambda c, t, k: f'={R("av_lo",c)}*{R("sw_perMW_lo",c)}')
    row('sw_hi', 'Flow for average heat, L/s, high = av_hi × perMW_hi', f=lambda c, t, k: f'={R("av_hi",c)}*{R("sw_perMW_hi",c)}')
    row('sw_vol_lo', 'Volume m³/yr, low = L/s × 31,536', f=lambda c, t, k: f'={R("sw_lo",c)}*hours*3.6')
    row('sw_vol_hi', 'Volume m³/yr, high', f=lambda c, t, k: f'={R("sw_hi",c)}*hours*3.6')
    row('sw_rent_lo', 'BC annual rental $/yr, low ("cooling" $1.30, min $200)', f=lambda c, t, k: f'=IF({R("van",c)}=1,MAX(bc_rent_min,{R("sw_vol_lo",c)}/1000*bc_rent_cool),"n/a (Alberta)")')
    row('sw_rent_hi', 'BC annual rental $/yr, high ("heat exchange" $2.25)', f=lambda c, t, k: f'=IF({R("van",c)}=1,MAX(bc_rent_min,{R("sw_vol_hi",c)}/1000*bc_rent_hx),"n/a (Alberta)")')
    row('sw_dist', 'Distance to the nearest usable water, m (sea for Van, else river)', f=lambda c, t, k: f'=IF({R("van",c)}=1,MIN({R("d_sea",c)},{R("d_river",c)}),{R("d_river",c)})')

    head('rooftop_solar — neighbour roofs, MWh/yr')
    for R_ in RADII:
        row(f'pvreg{R_}', f'{R_} m: register coefficient (128 Van / 159 Edm kWh/m²)', f=lambda c, t, k, R_=R_: f'={R(f"roof{R_}",c)}*IF({R("van",c)}=1,pv_reg_van,pv_reg_edm)/1000')
        row(f'pvlo{R_}', f'{R_} m: rebuilt from register factors, low (0.81 × 0.70 × 0.16 × horizontal yield)',
            f=lambda c, t, k, R_=R_: f'={R(f"roof{R_}",c)}*pv_usable*pv_cover*kwp_lo*IF({R("van",c)}=1,yld_van_lo,yld_edm_lo)/1000')
        row(f'pvhi{R_}', f'{R_} m: rebuilt from register factors, high (0.20 kWp/m², latitude-tilt yield)',
            f=lambda c, t, k, R_=R_: f'={R(f"roof{R_}",c)}*pv_usable*pv_cover*kwp_hi*IF({R("van",c)}=1,yld_van_hi,yld_edm_hi)/1000')
    row('pv_fac_lo', 'Factor check: register factors give kWh/m², low', f=lambda c, t, k: f'=pv_usable*pv_cover*kwp_lo*IF({R("van",c)}=1,yld_van_lo,yld_edm_lo)')
    row('pv_fac_hi', 'Factor check: register factors give kWh/m², high', f=lambda c, t, k: f'=pv_usable*pv_cover*kwp_hi*IF({R("van",c)}=1,yld_van_hi,yld_edm_hi)')
    row('pv_kwp_impl', 'kWp/m² implied by 128/159 at latitude-tilt yield', f=lambda c, t, k: f'=IF({R("van",c)}=1,pv_reg_van/(pv_usable*pv_cover*yld_van_hi),pv_reg_edm/(pv_usable*pv_cover*yld_edm_hi))')
    row('pv_share_lo', 'Neighbour solar (500 m, register coef) ÷ facility electricity, low', f=lambda c, t, k: f'={R("pvreg500",c)}/{R("mwh_hi",c)}')
    row('pv_share_hi', 'same, high', f=lambda c, t, k: f'={R("pvreg500",c)}/{R("mwh_lo",c)}')

    head('electricity_price')
    row('el_lgs_e', 'BC LGS energy charge $/yr, low (6.79 ¢/kWh × MWh)', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("mwh_lo",c)}*1000*lgs_kwh,"n/a")')
    row('el_lgs_e_hi', 'BC LGS energy charge $/yr, high', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("mwh_hi",c)}*1000*lgs_kwh,"n/a")')
    row('el_1830_e', 'BC RS 1830 energy charge $/yr, low (4.914 ¢/kWh)', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("mwh_lo",c)}*1000*t1830_kwh,"n/a")')
    row('el_1830_e_hi', 'BC RS 1830 energy charge $/yr, high', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("mwh_hi",c)}*1000*t1830_kwh,"n/a")')
    row('el_lgs_d', 'BC LGS demand charge per billing period $, low–high peak (period not in register)', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("pk_lo",c)}*1000*lgs_kw,"n/a")')
    row('el_lgs_d_hi', 'BC LGS demand charge per billing period $, high', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("pk_hi",c)}*1000*lgs_kw,"n/a")')
    row('el_1830_d', 'BC RS 1830 demand per period $, at least (kVA ≥ kW)', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("pk_lo",c)}*1000*t1830_kva,"n/a")')
    row('el_ab_rate_lo', 'Alberta energy $/MWh, low = pool × (1 + 7.99 %) + $2.50', f=lambda c, t, k: f'=IF({R("van",c)}=0,pool_lo*(1+sas_or)+sas_kwh*1000,"n/a")')
    row('el_ab_rate_hi', 'Alberta energy $/MWh, high', f=lambda c, t, k: f'=IF({R("van",c)}=0,pool_hi*(1+sas_or)+sas_kwh*1000,"n/a")')
    row('el_ab_fix_lo', 'EPCOR DAS + SAS daily charges $/yr, low (peak_lo kW)', f=lambda c, t, k: f'=IF({R("van",c)}=0,(das_day+{R("pk_lo",c)}*1000*(das_kwday+sas_cap+sas_dem))*365,"n/a")')
    row('el_ab_fix_hi', 'EPCOR DAS + SAS daily charges $/yr, high', f=lambda c, t, k: f'=IF({R("van",c)}=0,(das_day+{R("pk_hi",c)}*1000*(das_kwday+sas_cap+sas_dem))*365,"n/a")')
    row('el_ab_tot_lo', 'Alberta total $/yr, low (riders excluded)', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("mwh_lo",c)}*{R("el_ab_rate_lo",c)}+{R("el_ab_fix_lo",c)},"n/a")')
    row('el_ab_tot_hi', 'Alberta total $/yr, high', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("mwh_hi",c)}*{R("el_ab_rate_hi",c)}+{R("el_ab_fix_hi",c)},"n/a")')
    row('el_ab_eff_lo', 'Alberta all-in $/MWh, low', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("el_ab_tot_lo",c)}/{R("mwh_lo",c)},"n/a")')
    row('el_ab_eff_hi', 'Alberta all-in $/MWh, high', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("el_ab_tot_hi",c)}/{R("mwh_hi",c)},"n/a")')

    head('water_supply — $/m³')
    row('wp_lo', 'Potable price $/m³, low (Van water + sewer ÷ 2.8316; Edm lowest tiers)', f=lambda c, t, k: f'=IF({R("van",c)}=1,(van_w_lo+van_sewer)/van_unit,edm_w_lo+edm_ww_lo+edm_san)')
    row('wp_hi', 'Potable price $/m³, high', f=lambda c, t, k: f'=IF({R("van",c)}=1,(van_w_hi+van_sewer)/van_unit,edm_w_hi+edm_ww_hi+edm_san)')

    head('reclaimed_makeup — evaporation upper bound (all heat evaporated; cooling type unknown)')
    row('ev_lo', 'Evaporation m³/yr, low = MWh × 3.6e6 ÷ 2,400 ÷ 1,000', f=lambda c, t, k: f'={R("mwh_lo",c)}*3600000/h_fg/1000')
    row('ev_hi', 'Evaporation m³/yr, high', f=lambda c, t, k: f'={R("mwh_hi",c)}*3600000/h_fg/1000')
    row('mu_lo', 'Make-up m³/yr, low = E × C/(C−1) at C 6', f=lambda c, t, k: f'={R("ev_lo",c)}*C_hi/(C_hi-1)')
    row('mu_hi', 'Make-up m³/yr, high at C 2', f=lambda c, t, k: f'={R("ev_hi",c)}*C_lo/(C_lo-1)')
    row('rec_lo', 'V_rec ML/yr, low (s 0.74, operator claim)', f=lambda c, t, k: f'=s_lo*{R("mu_lo",c)}/1000')
    row('rec_hi', 'V_rec ML/yr, high (s 0.99)', f=lambda c, t, k: f'=s_hi*{R("mu_hi",c)}/1000')
    row('rec_epcor', 'Edmonton: V_rec high ÷ EPCOR 2024 industrial reclaimed supply', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("rec_hi",c)}/epcor_rec_ML,"n/a")')
    row('lbnl_lo', 'Cross-check: LBNL site water m³/yr on IT kWh, low (0.32 L/kWh)', f=lambda c, t, k: f'={R("itmwh",c)}*lbnl_w_lo')
    row('lbnl_hi', 'Cross-check: LBNL on facility kWh, high (0.36 L/kWh)', f=lambda c, t, k: f'={R("mwh_hi",c)}*lbnl_w_hi')
    row('mu_cost_lo', 'Potable cost of make-up $/yr, low', f=lambda c, t, k: f'={R("mu_lo",c)}*{R("wp_lo",c)}')
    row('mu_cost_hi', 'Potable cost of make-up $/yr, high', f=lambda c, t, k: f'={R("mu_hi",c)}*{R("wp_hi",c)}')

    head('greywater_neighbours — ML/yr (recovery share is a register gap, so "used" is not derivable)')
    for R_ in RADII:
        row(f'gw{R_}', f'V_gw within {R_} m, ML/yr = residents × 98 × 365 ÷ 1e6', f=lambda c, t, k, R_=R_: f'={R(f"res{R_}",c)}*gw_lpd*365/1000000')
    row('gw_ratio_lo', 'V_gw (500 m) ÷ make-up high', f=lambda c, t, k: f'={R("gw500",c)}*1000/{R("mu_hi",c)}')
    row('gw_ratio_hi', 'V_gw (500 m) ÷ make-up low', f=lambda c, t, k: f'={R("gw500",c)}*1000/{R("mu_lo",c)}')

    head('existing_building — m² (suitability by use and age has no numeric rule in the register)')
    row('eb_fp', 'Footprint of buildings on the lot, m²', f=lambda c, t, k: f'={R("fp",c)}')
    row('eb_gross', 'Floor area, m²', f=lambda c, t, k: f'=IF(ISNUMBER({R("gross",c)}),{R("gross",c)},"not derivable")')

    head('borefield — heat to ground = heat rejected (register warns q was derived from cooling load)')
    row('q_lo', 'q low, W/m (Van: Ithaca transferred; Edm: × 1.9/3.27)', f=lambda c, t, k: f'=IF({R("van",c)}=1,q_ith_lo,q_ith_lo*k_edm/k_ith)')
    row('q_hi', 'q high, W/m', f=lambda c, t, k: f'=IF({R("van",c)}=1,q_ith_hi,q_ith_hi*k_edm/k_ith)')
    row('bf_L_lo', 'Borehole length m, low = average heat (PUE 1.4) ÷ q high', f=lambda c, t, k: f'={R("av_lo",c)}*1000000/{R("q_hi",c)}')
    row('bf_L_hi', 'Borehole length m, high = peak heat (PUE 1.55) ÷ q low', f=lambda c, t, k: f'={R("pk_hi",c)}*1000000/{R("q_lo",c)}')
    row('bf_n_lo', 'Boreholes, low = L ÷ 150', f=lambda c, t, k: f'=ROUNDUP({R("bf_L_lo",c)}/bh_len,0)')
    row('bf_n_hi', 'Boreholes, high', f=lambda c, t, k: f'=ROUNDUP({R("bf_L_hi",c)}/bh_len,0)')
    row('bf_A_lo', 'Borefield area m², low = boreholes × 36', f=lambda c, t, k: f'={R("bf_n_lo",c)}*bh_area')
    row('bf_A_hi', 'Borefield area m², high', f=lambda c, t, k: f'={R("bf_n_hi",c)}*bh_area')
    row('bf_lots_lo', 'Borefield ÷ lot area, low', f=lambda c, t, k: f'={R("bf_A_lo",c)}/{R("lot",c)}')
    row('bf_lots_hi', 'Borefield ÷ lot area, high', f=lambda c, t, k: f'={R("bf_A_hi",c)}/{R("lot",c)}')
    row('bf_haMW', 'Check: ha per MW at q low (register "≈2 ha per MW")', f=lambda c, t, k: f'=1000000/{R("q_lo",c)}/bh_len*bh_area/10000')

    head('ates — open-loop flow and heat returned (heat in = heat rejected; heat_captured share is another stage)')
    row('at_fl_lo', 'Flow m³/h, low = average MW × 3,600 ÷ (4.18 × 5)', f=lambda c, t, k: f'={R("av_lo",c)}*3600/(cp*ates_dT)')
    row('at_fl_hi', 'Flow m³/h, high = peak MW × 3,600 ÷ (4.18 × 5)', f=lambda c, t, k: f'={R("pk_hi",c)}*3600/(cp*ates_dT)')
    row('at_fl_asw', 'Register formula AS WRITTEN (MW × 3.6 ÷ (4.18 × ΔT)), peak high', f=lambda c, t, k: f'={R("pk_hi",c)}*3.6/(cp*ates_dT)')
    row('at_out_lo', 'Heat returned MWh/yr, low = MWh_lo × 0.4', f=lambda c, t, k: f'={R("mwh_lo",c)}*ates_rec_lo')
    row('at_out_hi', 'Heat returned MWh/yr, high = MWh_hi × 0.85', f=lambda c, t, k: f'={R("mwh_hi",c)}*ates_rec_hi')
    row('at_eur_lo', 'Cost € at €300/kW, low (average kW)', f=lambda c, t, k: f'={R("av_lo",c)}*1000*ates_eur_kw')
    row('at_eur_hi', 'Cost € at €300/kW, high (peak kW)', f=lambda c, t, k: f'={R("pk_hi",c)}*1000*ates_eur_kw')

    head('deep_geo — MWth per well pair (Edmonton only)')
    row('dg_lo', 'MWth per well pair, low', f=lambda c, t, k: f'=IF({R("van",c)}=0,geo_lo,"not plausible / unproven (register)")')
    row('dg_hi', 'MWth per well pair, high', f=lambda c, t, k: f'=IF({R("van",c)}=0,geo_hi,"not plausible / unproven (register)")')
    row('dg_dT', 'Implied ΔT K = MW × 1000 ÷ (flow × 4.18) (register gives no T_return)', f=lambda c, t, k: f'=IF({R("van",c)}=0,geo_lo*1000/(geo_m_lo*cp),"n/a")')
    row('dg_dT2', 'Implied ΔT K at the high end', f=lambda c, t, k: f'=IF({R("van",c)}=0,geo_hi*1000/(geo_m_hi*cp),"n/a")')
    row('dg_pairs_lo', 'Well pairs to equal the average heat rejected, low (av_lo ÷ 5.6)', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("av_lo",c)}/geo_hi,"n/a")')
    row('dg_pairs_hi', 'Well pairs to equal it, high (av_hi ÷ 2.1)', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("av_hi",c)}/geo_lo,"n/a")')

    head('flood_exposure — m above / below FCL')
    row('fl_share', 'Share of lot in the designated floodplain', f=lambda c, t, k: f'={R("flood",c)}' if t == 'van' else 'not derivable (no Edmonton floodplain layer; FPO sets no FCL)')
    row('fl_fcl', 'FCL m (GVRD datum)', f=lambda c, t, k: '=fcl' if t == 'van' else 'n/a')
    row('fl_m', 'Plant floor − FCL, m', f=lambda c, t, k: 'not derivable: no ground/floor elevation input and no datum conversion (register)')

    head('service_class')
    row('sc_kw_lo', 'Peak kW, low', f=lambda c, t, k: f'={R("pk_lo",c)}*1000')
    row('sc_kw_hi', 'Peak kW, high', f=lambda c, t, k: f'={R("pk_hi",c)}*1000')
    row('sc_cls', 'Class (kVA ≥ kW, so "above" is certain; "below" holds only at power factor near 1)',
        f=lambda c, t, k: (f'=IF({R("van",c)}=1,IF({R("sc_kw_lo",c)}>bch_2f_25,"above 13 MVA: two feeders may be required in both 12.5 and 25 kV areas",'
                           f'IF({R("sc_kw_lo",c)}>bch_2f_12,"above 6.5 MVA: two feeders may be required in 12.5 kV areas",'
                           f'IF({R("sc_kw_hi",c)}>bch_2f_12,"straddles 6.5 MVA across PUE 1.4–1.55","below 6.5 MVA in kW (kVA depends on power factor)"))),'
                           f'IF({R("sc_kw_lo",c)}>epcor_prim,"above 2,500 kVA: EPCOR primary service typical","below 2,500 kW"))'))
    row('sc_reg117', 'Alberta: peak MW ÷ 75 MW Reg 117 threshold (high)', f=lambda c, t, k: f'=IF({R("van",c)}=0,{R("pk_hi",c)}/ab_reg117,"n/a")')
    row('sc_sub', 'Distance to nearest transformer station, m', f=lambda c, t, k: f'={R("d_sub",c)}')

    head('interconnection')
    row('ix_km', 'Distance to nearest VANIX point, km', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("d_ixp",c)}/1000,"not derivable")')
    row('ix_ms', 'Added one-way fibre latency, ms (straight line; route longer)', f=lambda c, t, k: f'=IF({R("van",c)}=1,{R("ix_km",c)}*fibre_ms_km,"not derivable")')

    head('seismic_sa02 — g')
    row('sa_g', 'Sa(0.2), 2%/50 yr, Site Class C', f=lambda c, t, k: f'={R("sa",c)}' if not isinstance(sa(k), str) or sa(k).startswith('=') else 'not derivable')

    for col, w in (('A', 74),): s.column_dimensions[col].width = w
    for j in range(len(LOTS)): s.column_dimensions[L(2 + j)].width = 17
    s.freeze_panes = 'B5'
    return wb, rows, HOW

# ===================================================================== 3. rules and register notes
RULES = [  # part, city, stated rule (register coef / formula), status
    ('self_gen_rules', 'van', 'RS 2289 from 1 Jul 2026: export ≤100 kW per phase (or ≤100 kW nameplate) at $0.10/kWh, same or adjacent parcel; distribution voltage only, so a data centre on RS 1830 (≥60 kV) cannot use it; RS 2290 community generation ≤2 MW net injection.'),
    ('self_gen_rules', 'edm', 'Micro-generation ≤5 MW (lesser of 5 MW or service rating), retail credit <150 kW, wholesale above; sized to own use, on owned/leased/adjacent land; data centres paired with dedicated new generation or storage go to the front of the connection line (Alta Reg 117/2026, via government page).'),
    ('electricity_price', 'van', 'BC Hydro LGS (RS 1600–1611): $13.83/kW, 6.79 ¢/kWh, 29.99 ¢/day; RS 1830 (≥60 kV): $12.178/kVA, 4.914 ¢/kWh. >150 MW full-cost rule is context only (J.B. decision 30 Oct).'),
    ('electricity_price', 'edm', 'Pool $43.68/MWh (2025), $62.78 (2024) + EPCOR DAS-TOUP and SAS-TOUP; riders excluded; Rate of Last Resort not available.'),
    ('water_supply', 'van', 'Water $4.091–5.128 + sewer $7.960 per unit (2.8316 m³) → $4.26–4.62/m³; reclaimed case by case (UT-019), no rate.'),
    ('water_supply', 'edm', 'Placeholder $2.53–3.88/m³ (2022–23 water + proposed 2025 wastewater and sanitary); EPCOR supplied 4,226 ML reclaimed to industry in 2024, price not stated.'),
    ('ambient_cover', 'van', 'By-law 6555: quiet→quiet 55/45, activity→quiet 60/55, (source adds intermediate→quiet 60/50), intermediate or activity→other 70/65 dBA day/night; no ambient credit for equipment. Zone→by-law class mapping not in register, so the pair for each lot is not derivable.'),
    ('ambient_cover', 'edm', 'Coef: "Bylaw 14600 limits pending". Source bylaw14600: residential 65 day / 50 night, non-residential 75 / 60 dB(A) at the property line; AUC Rule 012 is a benchmark only.'),
    ('interconnection', 'van', 'VANIX 78 participants, six connection points (operator claim); fibre ≈0.005 ms/km one way; latency targets a few tens of ms, VR <16 ms.'),
    ('interconnection', 'edm', 'YEGIX 9 participants; one site (Wolfpaw, Scotia Place — no address in register).'),
    ('uhi_context', 'van', 'Heat mapping; 30% canopy by 2050; surface UHI summer day 7.25 °C (Duan 2026); no rule on equipment heat.'),
    ('uhi_context', 'edm', 'Cool Edmonton (Action 9); days >30 °C ~1 → 16 by 2050s; surface UHI winter night 2.86 °C (Duan 2026); no rule on equipment heat.'),
    ('restriction_history', 'van', 'Stage 3 2015; Stage 2 2023 (4 Aug–15 Oct); 2026 Stage 3 from 8 Jun.'),
    ('restriction_history', 'edm', 'One mandatory ban 29 Jan–2 Feb 2024 (plant failure, not drought); restrictions not every year.'),
    ('nations_named', 'van', 'xʷməθkʷəy̓əm (Musqueam), Sḵwx̱wú7mesh Úxwumixw (Squamish), səlilwətaɬ (Tsleil-Waututh); unceded territories; links to each Nation\'s referral page.'),
    ('nations_named', 'edm', 'Treaty 6 territory and the Métis homeland; peoples named, not individual Nations (Confederacy of Treaty No. 6 lists 15).'),
    ('sewer_sink', 'van', 'Temperature rules: source ≥11 °C at the plant (KWL); sink <40 °C, non-domestic <65 °C (Bylaw 299); UT-008 first-come, letter of support, user pays; city-wide ~45 MW (claim); NEU 3.2 + 6.6 MW.'),
    ('sewer_sink', 'edm', 'No EPCOR sewer-heat programme; cooling water to sanitary sewer needs EWSI written approval (Bylaw 19627 Sch. 2 s. 5.10); Blatchford sewer heat deferred.'),
    ('existing_building', 'both', 'Office floors 2.4 kPa (4.8 kPa ground/basement), computer rooms excluded (NBC 2020 / BCBC 2024 T 4.1.5.3); ~6.0 kPa US VA (claim) → upper office floors need strengthening.'),
    ('deep_geo', 'van', 'Not plausible for power; direct use unproven.'),
    ('deep_geo', 'edm', '~90 °C at 2.5 km; 30–80 kg/s → 2.1–5.6 MWth per well pair; $4.7–8 M per well; AER Directive 089.'),
    ('flood_exposure', 'van', 'FCL 4.6 m (GVRD datum; 1 m SLR to 2100, 1:500, 0.6 m freeboard); switchgear not below FCL unless protected (VBBL Div. C 2.2.9.3(2)).'),
    ('flood_exposure', 'edm', 'Not yet looked at (status); source edmFPO2250: FPO overlay sets no FCL of its own.'),
    ('service_class', 'van', 'Two-feeder supply may be required above 6.5 MVA (12.5 kV areas) or 13 MVA (25 kV areas); ≥60 kV under transmission rules.'),
    ('service_class', 'edm', 'EPCOR primary service typical above 2,500 kVA; transformer clearances 3.0 / 2.0 m; access road 4.6 m in 6.5 m strip.'),
]
NOTES = [  # part, kind, note
    ('surface_water_sink', 'contradiction', 'Coef repeats the same "19 Oct 2026 (K)" block three times; the first copy says distance to shore is not computed (no shoreline file), the second says Vancouver "not yet (CanVec Hydro to come)", the third says Vancouver is from CanVec. Source canvecHydroBC check says no extract under data/raw/canvec/BC, but data/raw/hydro/canvec_van_water_extract.json exists.'),
    ('surface_water_sink', 'unit/arithmetic', '"24–49 L/s per MW at 10–5 °C": with c = 4.18 kJ/kg·K the values are 23.9–47.8 L/s; 49 needs c ≈ 4.08 (sea water ≈ 3.99 gives 50).'),
    ('surface_water_sink', 'contradiction', 'Coef uses the $2.25 "heat exchange" rental; source bcWSRfees separates "cooling" $1.30 (open question to J.B.). Both computed.'),
    ('surface_water_sink', 'contradiction', 'DFO approach velocity 0.035 vs 0.055 m/s (acknowledged internal conflict in source).'),
    ('rooftop_solar', 'contradiction', '128 / 159 kWh/m² cannot be rebuilt from register factors: 0.81 × 0.70 × 0.16–0.20 kWp/m² × NRCan yield gives 81–114 (Van) and 86–141 (Edm). 128/159 imply ≈0.224 kWp/m² at latitude-tilt yield, above pv_roof\'s 0.16–0.20. The formula names kWp/m² but gives no value. Source check: 128/159 not found in NRCan (our derivation). Radius not stated.'),
    ('self_gen_rules', 'contradiction', 'Coef says the "same or adjacent parcel" clause is not confirmed, but bchTariff2026 (tariff text read 28 Oct) states it. Source bchRS2289 still says "RS 2289 tariff not read".'),
    ('electricity_price', 'source gap', 'Alberta pool $43.68 / $62.78 are cited to aesoStats2025, whose register text covers only operating reserves (Table 13); OPEN item asks to add the pool average. Unverified.'),
    ('electricity_price', 'unit', 'BC demand charges ($/kW, $/kVA) have no billing period in the register, so annual demand cost is not derivable; RS 1830 bills kVA but the tool load is kW (power factor not in register).'),
    ('electricity_price', 'scope', 'EPCOR DAS-CS (>5,000 kVA) is closed to new service; every Edmonton lot here exceeds 5 MW peak, so the applicable DAS rate for such loads is not stated beyond DAS-TOUP.'),
    ('service_class', 'unit', 'Thresholds are in kVA (6.5 / 13 MVA; 2,500 kVA); IT × PUE gives kW. Power factor is not in the register, so only "above" results are certain. bch_primary entry quotes no clause for the thresholds (check 2).'),
    ('ambient_cover', 'contradiction', 'Coef says "Edmonton Bylaw 14600 limits pending", but source bylaw14600 already gives 65/50 (residential) and 75/60 (non-residential) dB(A). Coef omits the intermediate→quiet 60/50 pair listed in van6555zones.'),
    ('interconnection', 'contradiction', 'Coef/status say exchange-point locations are not sourced, yet vanixFacilities lists VANIX addresses (operator claim, unreachable in check 2). 515 W Hastings could not be located in the tax report; eStruxture VAN-2 has no address. The 0.005 ms/km has no primary.'),
    ('reclaimed_makeup', 'contradiction', 'Coef cites LBNL 2024 for 0.32–0.36 L/kWh, but lbnl2024 role is relabelled "not supported (?)" and its text has no water figure. Make-up = E × C/(C − 1) agrees with blowdown (E + E/(C − 1)).'),
    ('reclaimed_makeup', 'scope', 's (share substituted) is user-set 0–1; computed at the 0.74–0.99 operator-claim precedents. Evaporation is an upper bound (all heat evaporated); cooling type is another stage.'),
    ('greywater_neighbours', 'gap', 'Recovery share is a register gap, so "used = min(V_gw × recovery, make-up)" is not derivable; "within reach" radius not defined — 200 and 500 m (people_affected radii) used.'),
    ('borefield', 'unit/basis', 'q (12 / 19 W/m) was derived by dividing the 14.5 kW cooling load by length (register says so), but the formula divides heat to ground, which also includes compressor work → L understated for a heat-pump system. Edmonton 1.9/3.27 scaling gives 6.97–11.04 W/m (register 7–11 ✓). "≈2 ha per MW" holds at 12 W/m only (19 W/m → 1.26 ha; Edmonton 7 W/m → 3.4 ha). No Vancouver ground data: Ithaca q transferred.'),
    ('ates', 'unit', 'Formula "flow = MW × 3.6 ÷ (4.18 × ΔT)" gives 0.172 at 1 MW, 5 K; the coef says 172 m³/h per MW. Factor 1,000: it must be MW × 3,600 ÷ (4.18 × ΔT) (or kW × 3.6). Recovery 0.6 point vs 0.4–0.85 range; €300/kW stated at 2 MW only.'),
    ('deep_geo', 'unstated input', '2.1–5.6 MWth at 30–80 kg/s implies ΔT ≈ 16.7 K (return ≈ 73 °C from 90 °C) with c = 4.18; T_return is not in the register. "$4.7–8 M per well" vs output "per well pair" — wells per pair cost unclear.'),
    ('flood_exposure', 'gap', 'FCL in GVRD datum; datum conversion not found; no plant floor elevation → m above/below not derivable. Edmonton FPO sets no FCL.'),
    ('restriction_history', 'contradiction', 'Coef gives "Stage 3 2015; Stage 2 2023 (4 Aug–15 Oct); 2026 Stage 3 from 8 Jun" as fact, but source mvStageHistory says these dates are not in the cited report and need their own sources (report: Stage 3 "anticipated early June"). Edmonton end date 2 Feb not confirmed (epcorStages ?).'),
    ('seismic_sa02', 'coverage', 'Only three register points (V09, E05, "E06"). "E06" is ambiguous between E06A and E06B; the point is checked against each lot. Source not reopened in check 1 (unreachable).'),
    ('sewer_sink', 'contradiction', 'UT-008 governs regional GVS&DD sewers only (check 2: municipal sewers need no GVS&DD approval) but the coef applies it generally. The "≥11 °C KWL design rule" is not in the kwlSewer source text. Source id edm18100 names a repealed bylaw; text is Bylaw 19627, which governs discharge, not heat extraction.'),
    ('uhi_context', 'comparability', 'Surface UHI values are different seasons and times (Van summer day 7.25 °C, Edm winter night 2.86 °C) — not comparable with each other.'),
    ('pue', 'contradiction', 'Register pue says use 1.4–1.55; the worked example in make_diagram_data.py (EX values) uses 1.4–1.6.'),
    ('existing_building', 'gap', 'No numeric rule for "use and age suit conversion" or for "what policy steers to other uses"; floor area for Vancouver lots is not in open data. Footprint, Edmonton assessed gross area and year built reported.'),
]

# ===================================================================== 4. main
def recalc(path):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice: return False
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    shutil.move(os.path.join(tmp, os.path.basename(path)), path)
    return True

PART_ROWS = {  # part -> (low row, high row, unit, how)
    'sewer_sink': (None, None, 'MW at the pipe', 'not derivable: no sewer flow data at any parcel; ΔT not given in the register'),
    'surface_water_sink': ('sw_lo', 'sw_hi', 'L/s of water for the average heat rejected',
                           'average heat (IT × 0.6 × PUE 1.4–1.55) × 1000 ÷ (4.18 × 10–5 K); capacity at the outfall is unmodelled in the register'),
    'rooftop_solar': ('pvlo200', 'pvreg500', 'MWh/yr', 'neighbour roofs within 200–500 m (radius not stated in the register): low = 200 m with register factors 0.81 × 0.70 × 0.16 kWp/m² × horizontal yield; high = 500 m × register coefficient 128 / 159 kWh/m²'),
    'electricity_price': (None, None, '$/yr', ''),
    'water_supply': ('wp_lo', 'wp_hi', '$/m³', 'potable water (+ sewer, Vancouver) at 2026 / placeholder tariffs; reclaimed price not published'),
    'existing_building': ('eb_fp', 'eb_fp', 'm² footprint on the lot', 'footprint of buildings on the lot (input_check); floor area only for Edmonton (assessed gross); conversion suitability not derivable'),
    'reclaimed_makeup': ('rec_lo', 'rec_hi', 'ML/yr', 's 0.74–0.99 (operator claim) × make-up; make-up = evaporation upper bound (MWh × 1.5 m³/MWh) × C/(C−1), C 6–2'),
    'greywater_neighbours': ('gw200', 'gw500', 'ML/yr', 'residents within 200 m (low) and 500 m (high) × 98 L/person/day × 365; recovery share is a register gap'),
    'borefield': ('bf_A_lo', 'bf_A_hi', 'm² of borefield', 'L = heat ÷ q (low: average heat, PUE 1.4, q high; high: peak heat, PUE 1.55, q low); boreholes = L ÷ 150 m, ROUNDUP; × 36 m²'),
    'ates': ('at_fl_lo', 'at_fl_hi', 'm³/h', 'flow = MW × 3,600 ÷ (4.18 × 5 K) (register formula off by 1,000); low average MW at PUE 1.4, high peak MW at PUE 1.55'),
    'deep_geo': ('dg_lo', 'dg_hi', 'MWth per well pair', 'register Edmonton figure (Majorowicz & Grasby 2021); Vancouver not plausible / unproven'),
    'flood_exposure': (None, None, 'm above / below FCL', 'not derivable: no floor elevation and no datum conversion; floodplain share and FCL reported'),
    'service_class': ('sc_kw_lo', 'sc_kw_hi', 'kW peak (class in "class")', 'IT × PUE 1.4–1.55 × 1000; compared with BC Hydro 6.5 / 13 MVA and EPCOR 2,500 kVA'),
    'interconnection': ('ix_km', 'ix_km', 'km to nearest located VANIX point', 'straight line from the lot edge; latency = km × 0.005 ms'),
    'seismic_sa02': ('sa_g', 'sa_g', 'g', 'register value from the NRCan tool where the register point lies on this lot'),
}

def main():
    G = derive()
    wb, rows, HOW = build(G)
    r = wb.create_sheet('Rules')
    r.append(['part', 'city', 'register stated rule (coef / source text)']); [setattr(x, 'font', Font(bold=True)) for x in r[1]]
    for p in RULES: r.append(list(p))
    r.column_dimensions['A'].width = 22; r.column_dimensions['C'].width = 150
    n = wb.create_sheet('Register notes')
    n.append(['part', 'kind', 'note']); [setattr(x, 'font', Font(bold=True)) for x in n[1]]
    for p in NOTES: n.append(list(p))
    n.column_dimensions['A'].width = 22; n.column_dimensions['B'].width = 18; n.column_dimensions['C'].width = 160
    for ws in (r, n):
        for row_ in ws.iter_rows(min_row=2):
            row_[-1].alignment = Alignment(wrap_text=True, vertical='top')
    out = os.path.join(HERE, 'place_check.xlsx'); wb.save(out); print('wrote', out)
    ok = recalc(out); print('recalculated with LibreOffice' if ok else 'LibreOffice not found: values not cached')

    vb = load_workbook(out, data_only=True)['Place']
    val = lambda key, j: vb.cell(rows[key], 2 + j).value
    def rd(x): return round(x, 3) if isinstance(x, float) else x
    res = {}
    for j, (k, l, cty, p, it) in enumerate(LOTS):
        o = {'_lot': dict(label=l, city=cty, parcel_id=p, it_mw=it)}
        for part, (lo, hi, unit, how) in PART_ROWS.items():
            o[part] = dict(low=rd(val(lo, j)) if lo else None, high=rd(val(hi, j)) if hi else None, unit=unit, how=how)
        o['surface_water_sink']['extra'] = dict(bc_rental_per_yr=[rd(val('sw_rent_lo', j)), rd(val('sw_rent_hi', j))], dist_to_water_m=rd(val('sw_dist', j)),
                                                flow_per_MW_Ls=[rd(val('sw_perMW_lo', j)), rd(val('sw_perMW_hi', j))])
        o['rooftop_solar']['extra'] = {f'{R_}m': dict(roof_m2=val(f'roof{R_}', j), register_coef=rd(val(f'pvreg{R_}', j)),
                                                      rebuilt=[rd(val(f'pvlo{R_}', j)), rd(val(f'pvhi{R_}', j))]) for R_ in RADII}
        if cty == 'van':
            o['electricity_price'] = dict(low=rd(val('el_1830_e', j)), high=rd(val('el_lgs_e_hi', j)), unit='$/yr energy charge only',
                how='RS 1830 (4.914 ¢/kWh) at MWh low to LGS (6.79 ¢/kWh) at MWh high; demand charges per billing period not annualisable (period not in register)',
                extra=dict(lgs_demand_per_period=[rd(val('el_lgs_d', j)), rd(val('el_lgs_d_hi', j))], rs1830_demand_per_period_at_least=rd(val('el_1830_d', j)),
                           mwh_per_yr=[rd(val('mwh_lo', j)), rd(val('mwh_hi', j))]))
        else:
            o['electricity_price'] = dict(low=rd(val('el_ab_tot_lo', j)), high=rd(val('el_ab_tot_hi', j)), unit='$/yr (riders excluded)',
                how='MWh × (pool × 1.0799 + $2.50/MWh) + EPCOR DAS-TOUP + SAS-TOUP daily charges × 365 on peak kW; pool $43.68–62.78 unverified in register',
                extra=dict(all_in_per_MWh=[rd(val('el_ab_eff_lo', j)), rd(val('el_ab_eff_hi', j))], mwh_per_yr=[rd(val('mwh_lo', j)), rd(val('mwh_hi', j))]))
        o['existing_building']['extra'] = dict(gross_floor_m2=val('eb_gross', j), year_built=val('yb', j))
        o['reclaimed_makeup']['extra'] = dict(makeup_m3=[rd(val('mu_lo', j)), rd(val('mu_hi', j))], lbnl_crosscheck_m3=[rd(val('lbnl_lo', j)), rd(val('lbnl_hi', j))],
                                              potable_cost_per_yr=[rd(val('mu_cost_lo', j)), rd(val('mu_cost_hi', j))], share_of_epcor_reclaimed=rd(val('rec_epcor', j)))
        o['greywater_neighbours']['extra'] = dict(residents=[val('res200', j), val('res500', j), val('res3000', j)], gw500_over_makeup=[rd(val('gw_ratio_lo', j)), rd(val('gw_ratio_hi', j))])
        o['borefield']['extra'] = dict(length_m=[rd(val('bf_L_lo', j)), rd(val('bf_L_hi', j))], boreholes=[val('bf_n_lo', j), val('bf_n_hi', j)],
                                       times_lot_area=[rd(val('bf_lots_lo', j)), rd(val('bf_lots_hi', j))], q_W_per_m=[rd(val('q_lo', j)), rd(val('q_hi', j))])
        o['ates']['extra'] = dict(heat_returned_MWh=[rd(val('at_out_lo', j)), rd(val('at_out_hi', j))], cost_eur=[rd(val('at_eur_lo', j)), rd(val('at_eur_hi', j))],
                                  register_formula_as_written=rd(val('at_fl_asw', j)))
        o['deep_geo']['extra'] = dict(implied_dT_K=[rd(val('dg_dT', j)), rd(val('dg_dT2', j))], pairs_for_avg_heat=[rd(val('dg_pairs_lo', j)), rd(val('dg_pairs_hi', j))])
        o['flood_exposure']['extra'] = dict(floodplain_share=val('fl_share', j), fcl_m=val('fl_fcl', j))
        o['service_class']['extra'] = dict(cls=val('sc_cls', j), reg117_share=rd(val('sc_reg117', j)), substation_m=val('sc_sub', j))
        o['interconnection']['extra'] = dict(nearest=val('ixp_name', j), one_way_ms=rd(val('ix_ms', j)))
        if cty == 'edm':
            o['interconnection']['how'] = 'not derivable: YEGIX site named (Scotia Place) but no address or coordinates in the register'
        if o['seismic_sa02']['low'] is None or isinstance(o['seismic_sa02']['low'], str):
            o['seismic_sa02'] = dict(low=None, high=None, unit='g', how='not derivable: no register value for a point on this lot; distances to the register points (m): ' +
                                     json.dumps({nm: d.get(k) for nm, d in G['_seis'].items() if k in d}))
        for part in ('self_gen_rules', 'ambient_cover', 'uhi_context', 'restriction_history', 'nations_named'):
            o[part] = dict(low=None, high=None, unit='rule' if part != 'ambient_cover' else 'dBA allowed (rule)',
                           how=' | '.join(t for pp, cc, t in RULES if pp == part and cc in (cty, 'both')))
        if cty == 'edm':
            o['ambient_cover']['low'], o['ambient_cover']['high'] = 50, 75
            o['ambient_cover']['unit'] = 'dB(A) at property line (Bylaw 14600: residential night 50 … non-residential day 75)'
        for part, v in o.items():   # text in a value cell means "not derivable / not applicable": keep it in how, null the number
            if part != '_lot' and (isinstance(v.get('low'), str) or isinstance(v.get('high'), str)):
                v['how'] = f"{v['low'] if isinstance(v['low'], str) else v['high']} — {v['how']}"; v['low'] = v['high'] = None
        res[k] = o
    meta = dict(built=str(datetime.date.today()), script='checks/place_check/build_place_check.py',
                inputs='checks/input_check/input_check.json; IT load per lot as given; register pue 1.4–1.55, load_factor 0.6',
                not_read='tool/*.js, tool/index.html, checks/golden, checks/place_check/screen',
                vanix_points_located=G['_ixp_found'], seismic_point_to_lot_m=G['_seis'],
                register_notes=[dict(part=a, kind=b, note=c) for a, b, c in NOTES])
    json.dump(dict(meta=meta, lots=res), open(os.path.join(HERE, 'place_check.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('wrote place_check.json')

if __name__ == '__main__':
    main()
