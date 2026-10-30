#!/usr/bin/env python3
"""Independent check of the "What we start with" (input) stage of the Beyond Data tool, 30 Oct 2026.

Derived straight from the raw data under data/ (and, where named, tool/receivers_*.json), without reading the
tool's JavaScript or its computed results. Every value carries how it was found and which file it came from.

Run:  python3 checks/input_check/build_input_check.py      (from the repo root or anywhere)
Out:  checks/input_check/input_check.json, checks/input_check/input_check.xlsx

Projections (stated once here, repeated in each 'how'):
  Vancouver  EPSG:26910  NAD83 / UTM zone 10N  (the projection the City's LiDAR is delivered in)
  Edmonton   EPSG:3776   NAD83 / Alberta 3TM ref. meridian 114 W  (the City of Edmonton's projection)
  Areas are also checked geodesically on the GRS80 ellipsoid (pyproj.Geod) as a CRS-free cross-check.
"""
import csv, json, math, os, glob, collections, datetime
import pyproj
from shapely.geometry import shape, Polygon, MultiPolygon, Point
from shapely.ops import transform as shp_transform, unary_union
from shapely.strtree import STRtree
from shapely import make_valid

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
D = os.path.join(ROOT, 'data'); TOOL = os.path.join(ROOT, 'tool')
csv.field_size_limit(10 ** 9)

LOTS = [  # key, label, city, parcel id as given
    ('V09', '275 1st Ave E', 'van', '63819363'),
    ('E05', '8803 Stadium Rd NW', 'edm', '51053'),
    ('CLARK', '1425 Clark Dr', 'van', '23060741'),
    ('E123', '12603 123 Street NW', 'edm', '883345'),
    ('V339', '339 1st Ave E', 'van', '63819485'),
    ('E06A', '74 Airport Rd NW', 'edm', '176767'),
    ('E06B', '64 Airport Rd NW', 'edm', '146402'),
    ('HAST', '555 W Hastings', 'van', '59013407'),
]
EPSG = {'van': 26910, 'edm': 3776}
FWD = {c: pyproj.Transformer.from_crs(4326, e, always_xy=True).transform for c, e in EPSG.items()}
GEOD = pyproj.Geod(ellps='GRS80')
REACH_M = 3000.0      # register receiver_radius: "search to 3 km in a straight line"
RECEPTOR_AUC_M = 1500.0  # register noise_limit (AUC Rule 012): receptor = any dwelling within 1.5 km of the property boundary

def rel(p): return os.path.relpath(p, ROOT)
def proj(geom, city): return shp_transform(FWD[city], geom)
def valid(g): return g if g.is_valid else make_valid(g)
def num(v):
    try: return float(str(v).replace(',', '')) if v not in (None, '') else None
    except ValueError: return None
def geod_area(geom_ll):
    polys = [geom_ll] if geom_ll.geom_type == 'Polygon' else list(geom_ll.geoms)
    return sum(abs(GEOD.geometry_area_perimeter(p)[0]) for p in polys)
def r1(x): return None if x is None else round(x, 1)

OUT = {k: {'_lot': dict(label=l, city=c, parcel_id=p)} for k, l, c, p in LOTS}
ASSUME = []
def put(k, item, value, how, files):
    OUT[k][item] = dict(value=value, how=how, data_file=files if isinstance(files, str) else '; '.join(files))

# =============================================================== Vancouver
VPAR = os.path.join(D, 'vancouver_parcels_2026.geojson')
VTAX = os.path.join(D, 'raw', 'van', 'property-tax-report.csv')
VFP = os.path.join(D, 'vancouver_footprints_2009.geojson')
VFLOOD = os.path.join(D, 'vancouver_designated_floodplain.csv')
van_ids = {p: k for k, l, c, p in LOTS if c == 'van'}
print('reading Vancouver parcels')
vp = json.load(open(VPAR, encoding='utf-8'))
van_lot_ll = collections.defaultdict(list); van_lot_props = collections.defaultdict(list)
res_parcels, strata_parcels = [], []
RES_CLASSES = {'Residential', 'Residential Inclusive', 'Residential Rental'}
for f in vp['features']:
    pr = f['properties']; tc = (pr.get('tax_coord') or '').lstrip('0')
    g = valid(shape(f['geometry']))
    if tc in van_ids:
        van_lot_ll[tc].append(g); van_lot_props[tc].append(pr)
    if pr.get('zoning_class') in RES_CLASSES: res_parcels.append((g, pr))
    elif pr.get('legal_type') == 'STRATA': strata_parcels.append((g, pr))

tax = collections.defaultdict(list)
for r in csv.DictReader(open(VTAX, encoding='utf-8-sig'), delimiter=';'):
    tc = r['LAND_COORDINATE'].lstrip('0')
    if tc in van_ids: tax[tc].append(r)

VLOT = {}; LLB = {}
for tc, k in van_ids.items():
    gl = unary_union(van_lot_ll[tc]); gp = proj(gl, 'van'); VLOT[k] = gp; LLB[k] = gl.buffer(0.0005).bounds
    attr = [p.get('area_m2') for p in van_lot_props[tc]]
    put(k, '1_lot_area_m2', round(gp.area, 1),
        f'{len(van_lot_ll[tc])} parcel polygon(s) with tax_coord {tc}, unioned, projected to EPSG:26910 (NAD83 UTM 10N), planar area. '
        f'Geodesic (GRS80) check: {geod_area(gl):.1f} m². Attribute area_m2 in the joined parcel file: {attr} (sum {sum(a for a in attr if a):.1f}). '
        'The 2026 tax report itself states no lot area.', [rel(VPAR), rel(VTAX)])
    trs = tax[tc]
    put(k, '2_zone', trs[0]['ZONING_DISTRICT'] if trs else None,
        f'Raw 2026 property-tax report, LAND_COORDINATE {tc}: ZONING_DISTRICT {sorted({r["ZONING_DISTRICT"] for r in trs})}, '
        f'ZONING_CLASSIFICATION {sorted({r["ZONING_CLASSIFICATION"] for r in trs})}, {len(trs)} record(s); joined parcel file says '
        f'{sorted({p.get("zoning") for p in van_lot_props[tc]})}. Independent zoning-map check NOT possible: no Vancouver zoning-district '
        'polygon layer is in data/ (the parcel file\'s zoning is the same tax-report field joined on).', [rel(VTAX), rel(VPAR)])

print('reading Vancouver footprints')
fpj = json.load(open(VFP, encoding='utf-8'))
LF, LN = {}, []
for fn in glob.glob(os.path.join(D, 'lidar2022', '*.json')):
    r = json.load(open(fn)); LF.update({int(a): v for a, v in r['fp'].items()}); LN += r['new']
for k, lot in VLOT.items():
    rows = []; part = []
    for f in fpj['features']:
        g = shape(f['geometry'])
        # cheap lon/lat pre-filter before projecting
        b = g.bounds; lb = LLB[k]
        if b[2] < lb[0] or b[0] > lb[2] or b[3] < lb[1] or b[1] > lb[3]: continue
        gp = proj(valid(g), 'van')
        if not gp.intersects(lot): continue
        share = gp.intersection(lot).area / gp.area if gp.area else 0
        pr = f['properties']; lv = LF.get(pr['id'])
        rec = dict(id=pr['id'], area_m2_attr=pr.get('area_m2'), area_m2_proj=round(gp.area, 1), share_on_lot=round(share, 2),
                   hgt_agl=pr.get('hgt_agl'), max_ht=pr.get('max_ht'),
                   lidar2022=None if lv is None else dict(p90=lv[0], max=lv[1], share_still_built=lv[2]))
        (rows if share >= 0.5 else part).append(rec)
    new = []
    for p90, hmax, area, ring in LN:
        pg = proj(Polygon(ring), 'van')
        if pg.representative_point().within(lot): new.append(dict(p90=p90, max=hmax, area_m2=area))
    hs = [r['hgt_agl'] for r in rows if r['hgt_agl'] is not None]; mx = [r['max_ht'] for r in rows if r['max_ht'] is not None]
    gone = [r['id'] for r in rows if r['lidar2022'] and r['lidar2022']['share_still_built'] < 0.1]
    put(k, '3_buildings', dict(count_2009=len(rows), footprints_m2=[r['area_m2_proj'] for r in rows],
                               total_footprint_m2=round(sum(r['area_m2_proj'] for r in rows), 1),
                               max_hgt_agl_m=max(hs) if hs else None, max_max_ht_m=max(mx) if mx else None,
                               lidar2022_gone_since_2009=gone, lidar2022_new_outlines=new, straddling_lot_line=part, detail=rows),
        'City 2009 footprints projected to EPSG:26910; a footprint counts as on the lot if >= 50% of its area lies inside the lot '
        '(assumption); hgt_agl and max_ht are the file\'s fields (m). LiDAR 2022 cross-check from data/lidar2022 (90th-percentile and max '
        'building height; share_still_built < 0.1 read as demolished). New LiDAR outlines counted if their representative point is on the lot.',
        [rel(VFP), 'data/lidar2022/*.json'])

print('reading floodplain')
fl = collections.defaultdict(list)
for r in csv.DictReader(open(VFLOOD, encoding='utf-8-sig'), delimiter=';'):
    fl[r['Name']].append(proj(valid(shape(json.loads(r['Geom']))), 'van'))
fl = {n: unary_union(v) for n, v in fl.items()}
for k, lot in VLOT.items():
    shares = {n: round(g.intersection(lot).area / lot.area, 3) for n, g in fl.items()}
    d = shares.get('Designated Floodplain (FCL 4.6m)', 0)
    put(k, '4_floodplain', ('yes' if d >= 0.999 else 'partly' if d > 0.001 else 'no'),
        f'Share of the lot (EPSG:26910) inside each layer: {shares}. Answer refers to "Designated Floodplain (FCL 4.6m)"; '
        'the other layers (Still Creek, wave effect zone, Fraser risk today) are shown beside it.', rel(VFLOOD))

# =============================================================== Edmonton
EPROP = os.path.join(D, 'edmonton_properties_2026.csv')
EZON = os.path.join(D, 'edmonton_zoning.geojson')
ELOTS = os.path.join(D, 'raw', 'edm', 'lots', 'edmonton_lots_epsb.json')
EROOF = os.path.join(D, 'edmonton_rooflines_est.geojson')
EGDB = os.path.join(D, 'raw', 'edm', '3d', '3dBuildings.gdb')
edm_ids = {p: k for k, l, c, p in LOTS if c == 'edm'}
print('reading Edmonton properties')
erec = collections.defaultdict(list); res_pts = []
for r in csv.DictReader(open(EPROP, encoding='utf-8')):
    if r['parcel_point_id'] in edm_ids: erec[r['parcel_point_id']].append(r)
    if r['lat'] and any(c in ('RESIDENTIAL', 'OTHER RESIDENTIAL') for c in (r['class1'], r['class2'], r['class3'])):
        res_pts.append((float(r['lon']), float(r['lat']), r['house_number'] + ' ' + r['street'], r['class1']))
print('reading Edmonton title lots')
raw = json.load(open(ELOTS))
def dec(ring):
    x = y = 0; out = []
    for i in range(0, len(ring), 2): x += ring[i]; y += ring[i + 1]; out.append((x / 1e6, y / 1e6))
    return out
lots_ll = []
for pid, oid, rings in raw['lots']:
    rs = [dec(r) for r in rings]
    if len(rs[0]) < 4: continue
    lots_ll.append((pid, oid, valid(Polygon(rs[0], [r for r in rs[1:] if len(r) >= 4]))))
LT = STRtree([g for _, _, g in lots_ll])
ELOT = {}; ELOT_LL = {}
for ppid, k in edm_ids.items():
    rs = erec[ppid]; r = rs[0]; pt = Point(float(r['lon']), float(r['lat']))
    hits = [lots_ll[i] for i in LT.query(pt, predicate='within')] or [lots_ll[i] for i in LT.query(pt, predicate='intersects')]
    hits = [h for h in hits if h[2].contains(pt) or h[2].intersects(pt)]
    if not hits:
        put(k, '1_lot_area_m2', None, 'no title lot contains the assessment point', rel(ELOTS)); continue
    gl = unary_union([h[2] for h in hits]); gp = proj(gl, 'edm'); ELOT[k] = gp; ELOT_LL[k] = gl
    put(k, '1_lot_area_m2', round(gp.area, 1),
        f'Alberta title-mapping lot(s) (EPSB copy, Altalis, current to 1 Sep 2025) containing the City assessment point of parcel_point_id {ppid} '
        f'(lot pid {[h[0] for h in hits]}), projected to EPSG:3776 (3TM 114), planar area. Geodesic (GRS80) check: {geod_area(gl):.1f} m². '
        f'City attributes: lot_size_m2 {[x["lot_size_m2"] for x in rs]}, parcel_area_m2 {[x["parcel_area_m2"] for x in rs]} '
        f'({len(rs)} assessment account(s): {[x["account"] for x in rs]}; legal "{r["legal"]}").', [rel(ELOTS), rel(EPROP)])
print('reading Edmonton zoning')
zj = json.load(open(EZON, encoding='utf-8'))
zpolys = [(f['properties'].get('zone'), valid(shape(f['geometry']))) for f in zj['features']]
ZT = STRtree([g for _, g in zpolys])
for ppid, k in edm_ids.items():
    rs = erec[ppid]; r = rs[0]; pt = Point(float(r['lon']), float(r['lat']))
    zp = sorted({zpolys[i][0] for i in ZT.query(pt, predicate='intersects')})
    ov = {}
    if k in ELOT_LL:
        lot = ELOT[k]
        for i in ZT.query(ELOT_LL[k], predicate='intersects'):
            z, g = zpolys[i]; a = proj(g, 'edm').intersection(lot).area
            if a > 0.5: ov[z] = round(ov.get(z, 0) + a / lot.area, 3)
    put(k, '2_zone', r['zoning'],
        f'Assessment record zoning field {sorted({x["zoning"] for x in rs})} (zoning_source {r["zoning_source"]}, zone_from_map {r["zone_from_map"]}). '
        f'Independently, the City zoning map (12 Oct 2026) polygon holding the assessment point: {zp}; share of the title lot in each zone: {ov}.',
        [rel(EPROP), rel(EZON)])
    put(k, '4_floodplain', 'not derivable',
        'No Edmonton floodplain layer is in data/ (the register names only the City of Vancouver designated-floodplain layer).', '—')

print('reading Edmonton rooflines (large)')
rj = json.load(open(EROOF, encoding='utf-8'))
roofs = []
for f in rj['features']:
    g = shape(f['geometry']); b = g.bounds
    roofs.append((f['properties'], g))
RT = STRtree([g for _, g in roofs])
import fiona
for k, lot in ELOT.items():
    rows = []
    for i in RT.query(ELOT_LL[k], predicate='intersects'):
        pr, g = roofs[i]; gp = proj(valid(g), 'edm')
        share = gp.intersection(lot).area / gp.area if gp.area else 0
        if share < 0.5: continue
        rows.append(dict(id=pr['id'], footprint_attr=pr.get('footprint_m2'), footprint_proj=round(gp.area, 1), share_on_lot=round(share, 2),
                         height_est_m=pr.get('height_est_m'), storeys_est=pr.get('storeys_est'), geom=gp))
    # measured: City 3D buildings, multipatch parts whose vertex centroid lies inside the footprint; height = top - lowest base
    b = lot.bounds
    parts = []
    with fiona.open(EGDB, layer='Building_MP_Merge') as src:
        for f in src.filter(bbox=(b[0] - 5, b[1] - 5, b[2] + 5, b[3] + 5)):
            cs = []
            def walk(c):
                if isinstance(c[0], (int, float)): cs.append(c)
                else:
                    for x in c: walk(x)
            def visit(gm):
                if gm is None: return
                if gm['type'] == 'GeometryCollection':
                    for h in gm['geometries']: visit(h)
                elif gm['coordinates']: walk(gm['coordinates'])
            visit(f['geometry'])
            if not cs: continue
            zs = [c[2] for c in cs]
            parts.append((Point(sum(c[0] for c in cs) / len(cs), sum(c[1] for c in cs) / len(cs)), min(zs), max(zs)))
    for row in rows:
        ins = [(z0, z1) for p, z0, z1 in parts if row['geom'].contains(p) and z0 > 500]
        row['height_3d_m'] = round(max(z1 for _, z1 in ins) - min(z0 for z0, _ in ins), 1) if ins else None
        row['n_3d_parts'] = len(ins)
        del row['geom']
    lot_parts = [(z0, z1) for p, z0, z1 in parts if lot.contains(p) and z0 > 500]
    he = [r['height_est_m'] for r in rows if r['height_est_m']]; hm = [r['height_3d_m'] for r in rows if r['height_3d_m']]
    put(k, '3_buildings', dict(count=len(rows), footprints_m2=[r['footprint_proj'] for r in rows],
                               total_footprint_m2=round(sum(r['footprint_proj'] for r in rows), 1),
                               max_height_3d_measured_m=max(hm) if hm else None, max_height_estimate_m=max(he) if he else None,
                               parts_3d_on_lot=len(lot_parts), detail=rows),
        'City rooflines (2019) projected to EPSG:3776; footprint on the lot if >= 50% inside (assumption). Two heights: height_est_m is the file\'s '
        'ESTIMATE (assessed gross area ÷ footprint × 4 m, not measured); height_3d_m is measured here from the City 3D buildings (EPSG:3776): '
        'highest vertex − lowest vertex of multipatch parts whose vertex centroid falls inside the footprint (parts with base elevation <= 500 m skipped as artefacts).',
        [rel(EROOF), rel(EGDB)])

# =============================================================== all lots: sewer, receivers, receptors, coverage
LOTP = dict(VLOT); LOTP.update(ELOT)
for k, l, c, p in LOTS:
    put(k, '5_sewer_flow_m3_day', 'not derivable',
        'No sewer-main or sewer-flow data file exists in data/ or tool/ (searched for sewer/sanitary/wastewater/main). The register (sewer_flow, conf 20) '
        'says the sewer-mains layer gives type and material only and flow is by request to the City.', '—')

print('receivers')
REC = {c: json.load(open(os.path.join(TOOL, f'receivers_{c}.json'), encoding='utf-8')) for c in ('van', 'edm')}
NETF = {'van': glob.glob(os.path.join(D, 'networks', 'van_*.geojson')), 'edm': glob.glob(os.path.join(D, 'networks', 'edm_*.geojson'))}
for k, l, c, p in LOTS:
    lot = LOTP.get(k)
    if lot is None: continue
    ds = []
    for x in REC[c]:
        pt = proj(Point(x['lon'], x['lat']), c); ds.append((lot.distance(pt), lot.centroid.distance(pt), x))
    inr = [d for d in ds if d[0] <= REACH_M]
    off = [d for d in inr if d[0] > 0]
    near = min(off, key=lambda d: d[0]) if off else None
    on = [dict(t=x['t'], n=x['n']) for d, _, x in inr if d == 0]
    nets = {}
    for fn in NETF[c]:
        g = json.load(open(fn)); nm = os.path.basename(fn)
        if 'cpcn' in nm: continue  # electricity study area, not a heat network
        gp = proj(valid(shape(g['geometry'])), c); nets[nm] = round(lot.distance(gp), 0)
    put(k, '7_receivers_within_3km', dict(count=len(inr), count_off_lot=len(off), by_type=dict(collections.Counter(x['t'] for _, _, x in inr)),
                                          on_lot=on,
                                          nearest_off_lot=None if not near else dict(name=near[2]['n'], type=near[2]['t'], src=near[2]['src'], dist_edge_m=round(near[0], 1), dist_centroid_m=round(near[1], 1)),
                                          count_by_centroid=sum(1 for _, dc, _ in ds if dc <= REACH_M),
                                          heat_network_outlines_dist_m=nets),
        f'Receiver points in tool/receivers_{c}.json ({len(REC[c])} records), projected to EPSG:{EPSG[c]}; straight-line distance from the lot boundary '
        '(0 if on the lot) ≤ 3,000 m, the register\'s receiver_radius ("search to 3 km in a straight line"). Per-type reaches (RVO 1.0–2.1 km) are '
        'retired for searching (19 Oct 2026) and not applied. count_by_centroid uses the lot centroid instead. Network outlines (data/networks) '
        'given as distance only (0 = lot inside).', [f'tool/receivers_{c}.json', 'data/networks/*.geojson'])

print('receptors')
RP_V = STRtree([proj(g, 'van') for g, _ in res_parcels]); RPV = [(proj(g, 'van'), pr) for g, pr in res_parcels]
SP_V = [(proj(g, 'van'), pr) for g, pr in strata_parcels]; SPT = STRtree([g for g, _ in SP_V])
for k, l, c, p in LOTS:
    if c != 'van': continue
    lot = VLOT[k]
    i = RP_V.nearest(lot); g, pr = RPV[i]
    j = SPT.nearest(lot); g2, pr2 = SP_V[j]
    n15 = len(RP_V.query(lot.buffer(RECEPTOR_AUC_M), predicate='intersects'))
    put(k, '8_nearest_receptor_m', round(lot.distance(g), 1),
        f'Nearest parcel zoned residential (zoning_class Residential / Residential Inclusive / Residential Rental) in the 2026 parcel file: '
        f'{pr.get("address")} ({pr.get("zoning")}), edge-to-edge in EPSG:26910. ASSUMPTION: a residentially zoned parcel stands for a dwelling; '
        f'dwellings in commercial, CD-1, DD or industrial zones are missed. Cross-check, nearest STRATA parcel in a non-residential zone (often '
        f'condos, use not stated): {pr2.get("address")} ({pr2.get("zoning")}) at {lot.distance(g2):.1f} m. Residential-zoned parcels within 1.5 km '
        f'(AUC receptor reach): {n15}.', rel(VPAR))
EPTS = [(proj(Point(lo, la), 'edm'), a, cl) for lo, la, a, cl in res_pts]
EPT = STRtree([g for g, _, _ in EPTS])
for k, l, c, p in LOTS:
    if c != 'edm' or k not in ELOT: continue
    lot = ELOT[k]; ll = ELOT_LL[k]
    cand = [EPTS[i] for i in EPT.query(lot.buffer(RECEPTOR_AUC_M))]
    best = None
    # measure to the title lot holding each residential point (edge to edge); fall back to the point
    for g, a, cl in sorted(cand, key=lambda t: lot.distance(t[0]))[:60]:
        pll = Point(*pyproj.Transformer.from_crs(EPSG['edm'], 4326, always_xy=True).transform(g.x, g.y))
        hs = [lots_ll[i][2] for i in LT.query(pll, predicate='within')]
        d = lot.distance(proj(hs[0], 'edm')) if hs else lot.distance(g)
        if best is None or d < best[0]: best = (d, a, cl, bool(hs), lot.distance(g))
    put(k, '8_nearest_receptor_m', None if best is None else round(best[0], 1),
        ('none within 1.5 km' if best is None else
         f'Nearest assessment record with a residential class (RESIDENTIAL / OTHER RESIDENTIAL in class1–3): {best[1]} ({best[2]}); distance from the '
         f'subject title lot to that record\'s title lot, edge to edge, EPSG:3776 ({"lot found" if best[3] else "no lot, point used"}); to the assessment point itself '
         f'{best[4]:.1f} m. ASSUMPTION: a residential assessment class stands for a dwelling. Residential records within 1.5 km (AUC receptor reach): {len(cand)}.'),
        [rel(EPROP), rel(ELOTS)])

# coverage / plot ratio (register input plot_ratio = built area ÷ land area)
for k, l, c, p in LOTS:
    o = OUT[k]; lotA = o['1_lot_area_m2']['value']; b = o.get('3_buildings', {}).get('value')
    if not lotA or not b: continue
    cov = round(b['total_footprint_m2'] / lotA, 3)
    if c == 'edm':
        gross = sum(num(x['gross_area_m2']) or 0 for x in erec[p])
        put(k, '9_coverage_and_plot_ratio', dict(site_coverage=cov, plot_ratio_assessed_gross=round(gross / lotA, 3), gross_area_m2=gross),
            'site coverage = footprints on lot ÷ lot area (both EPSG:3776); plot ratio = assessed gross_area_m2 (City record, all accounts) ÷ lot area.', rel(EPROP))
    else:
        put(k, '9_coverage_and_plot_ratio', dict(site_coverage=cov, plot_ratio='not derivable without a storey count'),
            'site coverage = 2009 footprints on lot ÷ lot area (EPSG:26910). Plot ratio needs floor area: the tax report has none; '
            'footprint × storeys needs an assumed storey height (the register\'s 3.5–4.5 m) — not applied here.', rel(VFP))

# =============================================================== weather
print('weather')
def epw(path):
    L = open(path, encoding='latin-1').read().splitlines()
    hdr = L[0]; rows = [x.split(',') for x in L[8:] if x.strip()]
    t = [float(r[6]) for r in rows]; yrs = collections.Counter(r[0] for r in rows)
    return dict(hours=len(t), le8=sum(1 for v in t if v <= 8.0), lt8=sum(1 for v in t if v < 8.0), mean=round(sum(t) / len(t), 2),
                hdd18=round(sum(max(0, 18 - v) for v in t) / 24, 0), tmin=min(t), tmax=max(t), header=hdr, source_years=dict(sorted(yrs.items())))
W = {'van': os.path.join(D, '51376785-CAN_BC_VANCOUVER-INTL-A_1108395_CWEC.epw'), 'edm': os.path.join(D, 'effbc917-CAN_AB_EDMONTON-INTL-A_3012216_CWEC.epw')}
WV = {c: epw(p) for c, p in W.items()}
NORM = {'van': os.path.join(D, 'en_1991-2020_Normals_BC_VANCOUVER_INTERNATIONAL_AIRPORT.csv'), 'edm': os.path.join(D, 'en_1991-2020_Normals_AB_EDMONTON_AIRPORT.csv')}
for c, pth in NORM.items():
    for r in csv.reader(open(pth, encoding='utf-8-sig')):
        if len(r) > 17 and r[3] == 'Temperature' and r[4].startswith('Daily Average'):
            WV[c]['normal_mean'] = f'{r[0]} 1991–2020 normal annual daily average {r[17]} °C'; break
for k, l, c, p in LOTS:
    w = WV[c]
    put(k, '6_weather', dict(hours_le_8C=w['le8'], hours_lt_8C=w['lt8'], annual_mean_C=w['mean'], hdd18_hourly=w['hdd18'], min_C=w['tmin'], max_C=w['tmax'], hours=w['hours']),
        f'EPW field 7 (dry-bulb °C), {w["hours"]} hourly rows; hours with T <= 8.0 °C (and < 8.0 shown). HDD18 = Σ max(0, 18 − T) ÷ 24. '
        f'Header: "{w["header"]}". Year column in the rows: {w["source_years"]}. Sanity check: {w.get("normal_mean")} '
        '(normals station as named in the file; for Edmonton it is not confirmed to be the same station as the Edmonton Intl CWEC file).', [rel(W[c]), rel(NORM[c])])

# =============================================================== data notes found while checking
DATA_NOTES = []
for k, l, c, p in LOTS:
    if c == 'van':
        attr = sum(x.get('area_m2') or 0 for x in van_lot_props[p]); a = OUT[k]['1_lot_area_m2']['value']
        DATA_NOTES.append(f'{k}: parcel-file area_m2 {attr:.1f} vs EPSG:26910 {a:.1f} (ratio {attr / a:.4f}).')
DATA_NOTES += [
    'Vancouver parcel-file area_m2 runs a steady ~0.7% below the UTM and geodesic areas on all four lots; the ratio matches an equirectangular '
    'projection using 110,540 m per degree of latitude (the equator value; ~111,210 m at 49.25° N). make_tiles.py uses that constant for both '
    'cities, so tool distances and areas measured north–south would be ~0.6% short. Edmonton City lot_size_m2 / parcel_area_m2 agree with the title-lot '
    'geodesic area within 0.2%.',
    'Vancouver 2009 footprint area_m2 is per building (bldgid), not per polygon: multi-part buildings repeat the whole-building area on each part '
    '(e.g. 555 W Hastings parts 162924 / 161714 carry 11,907 m²). Areas here are measured from each polygon.',
    'Edmonton rooflines height_est_m disagrees with the City 3D buildings: 8803 Stadium Rd est. 4.0 m vs measured 16.6 m; 12603 123 St est. 20.0 m vs 9.3 m.',
    'No Vancouver zoning-district polygon layer, no Edmonton floodplain layer, and no sewer-main or sewer-flow file are in data/.',
    'CWEC files carry the header format "CWEC2011" while the register calls them the CWEC 2020 edition; the year column spans 2000–2017, consistent with the 2020 edition.',
    'Noise receptors for downtown lots: the residential-zone proxy puts 555 W Hastings 934 m from a dwelling, but strata parcels in CD-1/DD zones lie ~60 m away; '
    'a zone-based receptor rule understates downtown dwellings.',
]

# =============================================================== asked / setting inputs (not data-derived)
NOTDATA = {'start_case': 'asked of the user (quick-path default: new, whole building)', 'host_building': 'asked', 'cooling_type': 'permits / operator / site visit',
           'plant_schedule': 'permits, noise reports, site visit', 'spare_cooling': 'operator claim or judged', 'lighting_plan': 'site visit',
           'receiver_temp': 'set by receiver type, not by lot data', 'area': 'tool radius setting'}
for k in OUT:
    OUT[k]['10_not_from_lot_data'] = dict(value='not derivable from lot data', how='; '.join(f'{a}: {b}' for a, b in NOTDATA.items()), data_file='—')

ASSUME += [
    'Lot geometry: Vancouver = City parcel polygons sharing the tax coordinate, unioned; Edmonton = Alberta title lot(s) (EPSB copy) containing the City assessment point.',
    'Projections: Vancouver EPSG:26910 (NAD83 UTM 10N); Edmonton EPSG:3776 (NAD83 3TM 114). Geodesic area given as a cross-check.',
    'A building counts as on the lot if at least 50% of its footprint lies inside the lot.',
    'Receiver reach: 3,000 m straight line from the lot boundary (register receiver_radius); no per-type reach applied.',
    'Noise receptor = nearest dwelling: Vancouver residential-zoned parcel (zoning_class), Edmonton residential assessment class; neither file states dwelling use directly.',
    'Weather threshold: dry-bulb <= 8.0 °C inclusive (register: "only at <= 8 °C outdoors").',
    'Floodplain answer uses the "Designated Floodplain (FCL 4.6m)" layer only; other layers reported beside it.',
]
res = dict(meta=dict(built=datetime.date.today().isoformat(), script=rel(os.path.abspath(__file__)), assumptions=ASSUME, data_notes=DATA_NOTES,
                     crs=dict(van='EPSG:26910', edm='EPSG:3776')), lots=OUT)
json.dump(res, open(os.path.join(HERE, 'input_check.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)

# =============================================================== xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
def short(item, v):
    if not isinstance(v, dict): return v
    if item == '3_buildings':
        h = v.get('max_hgt_agl_m', v.get('max_height_3d_measured_m')); h2 = v.get('max_max_ht_m', v.get('max_height_estimate_m'))
        lab = ('max hgt_agl / max_ht' if 'max_hgt_agl_m' in v else 'max 3D measured / file estimate')
        s = f"{v.get('count_2009', v.get('count'))} bldg; footprints {v['footprints_m2']} (Σ {v['total_footprint_m2']}); {lab} {h} / {h2} m"
        if v.get('lidar2022_gone_since_2009'): s += f"; LiDAR 2022: gone {v['lidar2022_gone_since_2009']}"
        if v.get('lidar2022_new_outlines'): s += f"; LiDAR new {len(v['lidar2022_new_outlines'])}"
        return s
    if item == '6_weather': return f"≤8 °C: {v['hours_le_8C']} h; mean {v['annual_mean_C']} °C; HDD18 {v['hdd18_hourly']}"
    if item == '7_receivers_within_3km':
        n = v['nearest_off_lot']
        return f"{v['count']} within 3 km ({len(v['on_lot'])} on lot); nearest {n['name']} ({n['type']}) {n['dist_edge_m']} m" if n else f"{v['count']}"
    return json.dumps(v, ensure_ascii=False)
wb = Workbook(); ws = wb.active; ws.title = 'input check'
keys = [k for k, *_ in LOTS]
ws.append(['item'] + [f'{k} — {OUT[k]["_lot"]["label"]}' for k in keys] + ['how (first lot of each city shown in JSON)', 'data file'])
items = ['1_lot_area_m2', '2_zone', '3_buildings', '4_floodplain', '5_sewer_flow_m3_day', '6_weather', '7_receivers_within_3km', '8_nearest_receptor_m', '9_coverage_and_plot_ratio', '10_not_from_lot_data']
for it in items:
    row = [it] + [short(it, OUT[k].get(it, {}).get('value', 'not derivable')) for k in keys]
    hows = {OUT[k][it]['how'].split('. ')[0] for k in keys if it in OUT[k]}
    files = {OUT[k][it]['data_file'] for k in keys if it in OUT[k]}
    ws.append(row + [' | '.join(sorted(hows)), ' | '.join(sorted(files))])
ws2 = wb.create_sheet('how (full)')
ws2.append(['item', 'lot', 'value (full)', 'how', 'data file'])
for it in items:
    for k in keys:
        if it in OUT[k]:
            e = OUT[k][it]; ws2.append([it, k, json.dumps(e['value'], ensure_ascii=False, default=str)[:32000], e['how'], e['data_file']])
ws3 = wb.create_sheet('assumptions')
ws3.append(['ASSUMPTIONS'])
for a in ASSUME: ws3.append([a])
ws3.append([]); ws3.append(['DATA NOTES'])
for a in DATA_NOTES: ws3.append([a])
for w in (ws, ws2):
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='CC746C')
    for col in w.columns: w.column_dimensions[col[0].column_letter].width = 42
    for r in w.iter_rows():
        for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
ws3.column_dimensions['A'].width = 140
wb.save(os.path.join(HERE, 'input_check.xlsx'))
print('done')
